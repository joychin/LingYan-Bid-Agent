"""发布管线：校验链、包落盘、索引与 task-single upsert 语义（§16 作用域必填）。"""

import json

import pytest

from app import artifact_store, db, publish
from tests.util import init_env

KEY = "tender.directory/tender-response-docs@1"


@pytest.fixture
def env(tmp_path, monkeypatch):
    task, conv = init_env(tmp_path, monkeypatch)
    return {"root": tmp_path, "task": task, "conv": conv}


def _content(name="技术部分"):
    return {
        "response_documents": [
            {"name": name, "scope": "", "directory": [{"目录名称": "目录", "level": 1, "children": []}]}
        ]
    }


def _pub(env, content, **kw):
    """以任务归属发布（2026-08-31：文件归任务，conversation_id 仅 provenance）。"""
    return publish.publish_artifact(
        KEY, content, task_id=env["task"]["id"], conversation_id=env["conv"]["id"], **kw
    )


def test_publish_creates_package_and_index(env):
    m = _pub(env, _content(), source={"skill": "demo-skill", "thread_id": env["conv"]["id"], "run_id": "r_1"})
    aid = m["artifact_id"]

    # meta：稳定身份字段齐备；task_id 恒为所属任务
    assert m["kind"] == "tender.directory"
    assert m["schema"] == {"id": "tender-response-docs", "version": 1}
    assert m["cardinality"] == "task-single"
    assert m["display_name"] == "投标目录"  # 未传 display_name → 契约默认名
    assert m["task_id"] == env["task"]["id"]
    assert m["conversation_id"] == env["conv"]["id"]
    assert artifact_store.read_meta(aid, m) == m

    # content：纯业务内容，JSON 可解析且与提交一致；包落在 work/artifacts/ 下
    content = json.loads(artifact_store.read_content(aid, m))
    assert content["response_documents"][0]["name"] == "技术部分"
    assert artifact_store.artifact_dir(aid, m) == (
        artifact_store.work_artifacts_dir(env["task"]["id"]) / aid
    )

    # 索引：emitted=0（待 run 边界发事件），task_id 同样恒写
    row = db.get_artifact_index(aid)
    assert row["kind"] == "tender.directory"
    assert row["task_id"] == env["task"]["id"]
    assert row["content_seq"] == 1
    assert row["emitted"] == 0
    assert db.pending_emit("r_1") == [row]


def test_publish_requires_task(env):
    """文件归任务：不指定 task 且无 conversation 反查 → 拒绝。"""
    with pytest.raises(publish.PublishError, match="所属任务"):
        publish.publish_artifact(KEY, _content())
    with pytest.raises(publish.PublishError, match="所属任务"):
        publish.publish_artifact(KEY, _content(), conversation_id="c_nope00000000")
    with pytest.raises(publish.PublishError, match="任务.*不存在"):
        publish.publish_artifact(KEY, _content(), task_id="t_nope00000000")


def test_publish_rejects_unknown_contract(env):
    with pytest.raises(publish.PublishError, match="未注册"):
        publish.publish_artifact("no.such/contract@1", _content(), task_id=env["task"]["id"])


def test_publish_rejects_schema_violation(env):
    with pytest.raises(publish.PublishError, match="不符合契约"):
        _pub(env, {"foo": "bar"})


def test_task_single_upsert_reuses_artifact(env):
    m1 = _pub(env, _content("旧版"), source={"skill": "t", "run_id": "r_1"})
    m2 = _pub(env, _content("新版"), source={"skill": "t", "run_id": "r_2"})

    # 复用 artifact_id 与 meta（稳定身份），内容替换
    assert m2["artifact_id"] == m1["artifact_id"]
    content = json.loads(artifact_store.read_content(m1["artifact_id"], m1))
    assert content["response_documents"][0]["name"] == "新版"
    # meta 的 created_at 不因重发布改写
    assert artifact_store.read_meta(m1["artifact_id"], m1)["created_at"] == m1["created_at"]

    # 索引：seq 递增、emitted 复位、last_run 更新
    row = db.get_artifact_index(m1["artifact_id"])
    assert row["content_seq"] == 2
    assert row["emitted"] == 0
    assert row["last_run_id"] == "r_2"
    assert len(db.list_artifact_index()) == 1
    # 旧 run 不再有待发事件，新 run 有
    assert db.pending_emit("r_1") == []
    assert len(db.pending_emit("r_2")) == 1


def test_list_from_disk_skips_conversation_history(env):
    """deepagents 压缩逐出历史的落盘目录（workspace/conversation_history/，压缩触发后
    才出现）不是任务——结构化扫描必须跳过，否则索引里出现幽灵任务。"""
    history = artifact_store.workspace_dir() / "conversation_history"
    history.mkdir(parents=True, exist_ok=True)
    (history / "sess-1.md").write_text("# evicted history", encoding="utf-8")
    assert artifact_store.list_from_disk() == []


def test_publish_identical_content_noop(env):
    """内容未变短路（2026-09-12）：同内容+同名重发布不产生新版本——seq/emitted/
    last_run_id/恢复点全不动（run 收尾 pending_emit 捞不到 → 无 artifact.created →
    聊天产物卡不挪位）；仅改显示名或内容 → 照常发布。"""
    m1 = _pub(env, _content(), source={"skill": "t", "run_id": "r_1"})
    db.mark_emitted(m1["artifact_id"])

    m2 = _pub(env, _content(), source={"skill": "t", "run_id": "r_2"})
    assert m2.get("_unchanged") is True
    row = db.get_artifact_index(m1["artifact_id"])
    assert row["content_seq"] == 1
    assert row["emitted"] == 1
    assert row["last_run_id"] == "r_1"  # 未随 r_2 更新——卡不挪位的关键
    assert db.pending_emit("r_2") == []
    assert artifact_store.latest_restore_point(m1["artifact_id"], m1) is None

    # 仅改显示名（内容相同）→ 不短路、走完整发布路径（保守：输入有差就不省）；
    # 注：existing 路径本就不回写 display_name（重发布改名=既有静默忽略，本批不动）
    m3 = _pub(env, _content(), display_name="改名", source={"skill": "t", "run_id": "r_2"})
    assert "_unchanged" not in m3
    row = db.get_artifact_index(m1["artifact_id"])
    assert row["content_seq"] == 2
    assert row["last_run_id"] == "r_2"

    # 内容变化 → 照常发布
    m4 = _pub(env, _content("改版"), source={"skill": "t", "run_id": "r_3"})
    assert "_unchanged" not in m4
    assert db.get_artifact_index(m1["artifact_id"])["content_seq"] == 3


def test_rebuild_preserves_runtime_state(env):
    """启动重建同步语义（2026-09-12）：幸存行保留 content_seq/updated_at/last_run_id/
    last_thread_id/emitted，不回卷到 meta.source 的创建 run（meta 从不在重发布时回写
    source——旧重建语义让聊天产物卡跳回首次发布回合、编辑器把 content_seq 当版本号
    探测外部更新也被重启归零误报）；库有磁盘无的行删除、磁盘有库无的行默认插入。"""
    import shutil

    to_path = lambda md: str(artifact_store.content_path(md["artifact_id"], md))  # noqa: E731
    m1 = _pub(env, _content("初版"), source={"skill": "t", "thread_id": env["conv"]["id"], "run_id": "r_1"})
    db.mark_emitted(m1["artifact_id"])
    _pub(env, _content("新版"), source={"skill": "t", "thread_id": env["conv"]["id"], "run_id": "r_2"})
    aid = m1["artifact_id"]
    before = db.get_artifact_index(aid)
    assert (before["content_seq"], before["emitted"], before["last_run_id"]) == (2, 0, "r_2")
    # meta.source 是创建时化石（r_1）——重建不得回卷
    assert artifact_store.read_meta(aid, m1)["source"]["run_id"] == "r_1"

    assert db.rebuild_artifact_index(artifact_store.list_from_disk(), to_path) == 1
    after = db.get_artifact_index(aid)
    assert after["content_seq"] == 2
    assert after["emitted"] == 0
    assert after["last_run_id"] == "r_2"
    assert after["last_thread_id"] == before["last_thread_id"]
    assert after["updated_at"] == before["updated_at"]

    # 库无磁盘有（索引被手删/全新 DB）→ 默认插入：seq=1、emitted=1、last_run 从 meta.source 兜底
    db.delete_artifact_index(aid)
    assert db.rebuild_artifact_index(artifact_store.list_from_disk(), to_path) == 1
    fresh = db.get_artifact_index(aid)
    assert (fresh["content_seq"], fresh["emitted"], fresh["last_run_id"]) == (1, 1, "r_1")

    # 磁盘包删除 → 行删除
    shutil.rmtree(artifact_store.artifact_dir(aid, m1))
    assert db.rebuild_artifact_index(artifact_store.list_from_disk(), to_path) == 0
    assert db.get_artifact_index(aid) is None


def test_rebuild_index_from_metas(env):
    m = _pub(env, _content(), source={"skill": "t", "thread_id": env["conv"]["id"], "run_id": "r_1"})
    db.mark_emitted(m["artifact_id"])

    # meta 权威：索引清空后可重建（包位置按 meta scope 派生），幂等
    to_path = lambda md: str(artifact_store.content_path(md["artifact_id"], md))  # noqa: E731
    rebuilt = db.rebuild_artifact_index(artifact_store.list_from_disk(), to_path)
    assert rebuilt == 1
    row = db.get_artifact_index(m["artifact_id"])
    assert row is not None
    assert row["task_id"] == env["task"]["id"]
    assert row["emitted"] == 1  # 启动重建后无待发事件
    # 发布来源随重建恢复（索引是 last_run_id 的唯一宿主；聊天产物卡按发布 run 归位依赖它）
    assert row["last_run_id"] == "r_1"
    assert row["last_thread_id"] == env["conv"]["id"]
    assert db.rebuild_artifact_index(artifact_store.list_from_disk(), to_path) == 1


def test_republish_keeps_restore_point(env):
    """发布即覆盖（文件夹语义）：覆盖前自动留恢复点，用户可恢复上一版。"""
    m1 = _pub(env, _content("初版"))
    _pub(env, _content("新版"))

    # 当前内容为新版
    content = json.loads(artifact_store.read_content(m1["artifact_id"], m1))
    assert content["response_documents"][0]["name"] == "新版"

    # 恢复点留底了初版
    rp = artifact_store.latest_restore_point(m1["artifact_id"], m1)
    assert rp is not None
    assert json.loads(rp.read_text(encoding="utf-8"))["response_documents"][0]["name"] == "初版"


def test_republish_after_package_deleted_creates_new(env):
    """僵尸索引行：包被删后重发布应新建完整包，而不是写进无 meta 的目录。"""
    import shutil

    m1 = _pub(env, _content())
    shutil.rmtree(artifact_store.artifact_dir(m1["artifact_id"], m1))

    m2 = _pub(env, _content("重建"))
    assert m2["artifact_id"] != m1["artifact_id"]
    assert artifact_store.read_meta(m2["artifact_id"], m2) == m2
    content = json.loads(artifact_store.read_content(m2["artifact_id"], m2))
    assert content["response_documents"][0]["name"] == "重建"

    # 僵尸行仍在索引（API 侧被 package_ready 过滤），可用的只有新行
    rows = db.list_artifact_index()
    assert len(rows) == 2
    ready = [r for r in rows if artifact_store.package_ready(r["artifact_id"], r)]
    assert len(ready) == 1
    assert ready[0]["artifact_id"] == m2["artifact_id"]


def test_resolved_content_path_containment(env):
    m = _pub(env, _content())
    aid = m["artifact_id"]

    assert artifact_store.resolved_content_path(aid, m).is_file()
    # 非法 aid（路径穿越）直接拒绝
    assert artifact_store.resolved_content_path("../../etc/passwd", m) is None
    assert artifact_store.resolved_content_path("art_short", m) is None

    # symlink 逃逸：content 指向 workspace 外 → 拒绝
    secret = env["root"] / "secret.json"
    secret.write_text("{}", encoding="utf-8")
    target = artifact_store.content_path(aid, m)
    target.unlink()
    target.symlink_to(secret)
    assert artifact_store.resolved_content_path(aid, m) is None


# ---------- 文件型产物（tender.volume=整本标书 docx，2026-09-13） ----------

VOLUME_KEY = "tender.volume/tender-volume-docx@1"


def _make_volume_docx(env, filename: str, text: str):
    from docx import Document

    p = artifact_store.work_dir(env["task"]["id"]) / "body" / filename
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(text)
    doc.save(p)
    return p


def _pub_volume(env, path, book: str, run_id: str = "r_1"):
    import hashlib

    data = path.read_bytes()
    return publish.publish_file_artifact(
        VOLUME_KEY,
        file_path=path,
        content_meta={
            "filename": path.name,
            "book": book,
            "size": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "merged_sections": 2,
            "images": 0,
            "comments": 0,
        },
        display_name=book,
        source={"skill": "tender-body", "thread_id": env["conv"]["id"], "run_id": run_id},
        task_id=env["task"]["id"],
        conversation_id=env["conv"]["id"],
    )


def test_publish_volume_creates_package_with_docx(env):
    src = _make_volume_docx(env, "整本-技术册.docx", "初版正文")
    m = _pub_volume(env, src, "技术册")
    aid = m["artifact_id"]

    assert m["kind"] == "tender.volume"
    assert m["content_type"].startswith("application/vnd.openxmlformats")
    # 包内唯一 docx 与源字节一致；content.json 是机器元信息（schema 校验对象）
    files = artifact_store.package_files(aid, m)
    assert [f.name for f in files] == ["整本-技术册.docx"]
    assert files[0].read_bytes() == src.read_bytes()
    meta = json.loads(artifact_store.read_content(aid, m))
    assert meta["book"] == "技术册"
    assert meta["filename"] == "整本-技术册.docx"
    # 索引行：display_name=册名（身份键）、content_path 恒指包内 content.json（重启重建同口径，
    # docx 本体不走该列、由 /file 端点按包内唯一 .docx 取）、run 边界可捞事件
    row = db.get_artifact_index(aid)
    assert row["display_name"] == "技术册"
    assert row["content_path"].endswith("content.json")
    assert db.pending_emit("r_1") == [row]


def test_publish_volume_zip_equal_noop_and_change_publishes(env):
    """zip 内容级去重：同内容重保存（zip 时间戳必变、CRC 不变）→ 短路（seq/
    last_run_id/emitted 全不动）；内容变 → 覆盖发布 seq+1；文件型产物明确不留恢复点。"""
    from docx import Document

    src = _make_volume_docx(env, "整本-技术册.docx", "初版正文")
    m1 = _pub_volume(env, src, "技术册")
    db.mark_emitted(m1["artifact_id"])

    doc = Document(str(src))
    doc.save(src)  # 同内容重保存——字节必变（zip 时间戳），CRC 不变
    m2 = _pub_volume(env, src, "技术册", run_id="r_2")
    assert m2.get("_unchanged") is True
    row = db.get_artifact_index(m1["artifact_id"])
    assert (row["content_seq"], row["last_run_id"], row["emitted"]) == (1, "r_1", 1)
    assert db.pending_emit("r_2") == []

    doc = Document(str(src))
    doc.add_paragraph("第二版正文")
    doc.save(src)
    m3 = _pub_volume(env, src, "技术册", run_id="r_3")
    assert m3["artifact_id"] == m1["artifact_id"]
    assert "_unchanged" not in m3
    assert db.get_artifact_index(m1["artifact_id"])["content_seq"] == 2
    assert db.pending_emit("r_3") != []
    assert artifact_store.package_files(m1["artifact_id"], m1)[0].read_bytes() == src.read_bytes()
    assert artifact_store.latest_restore_point(m1["artifact_id"], m1) is None


def test_publish_volume_book_is_identity_key(env):
    """身份复用键=任务+kind+册名：多册各一条；同册重发布（内容未变）复用同一条不增行。"""
    src_t = _make_volume_docx(env, "整本-技术册.docx", "技术内容")
    src_b = _make_volume_docx(env, "整本-商务册.docx", "商务内容")
    m1 = _pub_volume(env, src_t, "技术册")
    m2 = _pub_volume(env, src_b, "商务册")
    assert m1["artifact_id"] != m2["artifact_id"]

    m3 = _pub_volume(env, src_t, "技术册", run_id="r_2")
    assert m3["artifact_id"] == m1["artifact_id"]
    assert m3.get("_unchanged") is True
    vrows = [r for r in db.list_artifact_index() if r["kind"] == "tender.volume"]
    assert len(vrows) == 2


def test_publish_volume_update_failure_keeps_old_state(env, monkeypatch):
    """更新路径中途失败：包保持完整旧态（旧 docx+旧 content.json+旧 seq），
    重跑收敛新态——不再被内容短路把失败半态固化（2026-09-15 审计）。"""
    from docx import Document

    src = _make_volume_docx(env, "整本-技术册.docx", "初版正文")
    m1 = _pub_volume(env, src, "技术册")
    aid = m1["artifact_id"]
    old_bytes = artifact_store.package_files(aid, m1)[0].read_bytes()
    old_content = artifact_store.read_content(aid, m1)

    doc = Document(str(src))
    doc.add_paragraph("第二版正文")
    doc.save(src)

    real = artifact_store.replace_current_content

    def _boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(artifact_store, "replace_current_content", _boom)
    with pytest.raises(OSError):
        _pub_volume(env, src, "技术册", run_id="r_2")
    monkeypatch.setattr(artifact_store, "replace_current_content", real)

    # 包保持完整旧态
    row = db.get_artifact_index(aid)
    assert row["content_seq"] == 1
    assert artifact_store.package_files(aid, row)[0].read_bytes() == old_bytes
    assert artifact_store.read_content(aid, row) == old_content
    assert db.pending_emit("r_2") == []

    # 重跑收敛新态（半态不会命中内容短路）
    m2 = _pub_volume(env, src, "技术册", run_id="r_3")
    assert "_unchanged" not in m2
    assert db.get_artifact_index(aid)["content_seq"] == 2
    assert artifact_store.package_files(aid, m2)[0].read_bytes() == src.read_bytes()
    # 无暂存件残留
    assert list(artifact_store.artifact_dir(aid, m2).glob("incoming-*.part")) == []


def test_publish_volume_create_failure_cleans_package_dir(env, monkeypatch):
    """新建路径中途失败：半包目录清掉（无 meta 的目录读侧本就当不存在，
    不清只会留磁盘孤儿）。"""
    src = _make_volume_docx(env, "整本-技术册.docx", "正文")

    def _boom(*a, **k):
        raise OSError("disk full")

    monkeypatch.setattr(artifact_store, "create_package", _boom)
    with pytest.raises(OSError):
        _pub_volume(env, src, "技术册")
    assert [r for r in db.list_artifact_index() if r["kind"] == "tender.volume"] == []
    arts = artifact_store.work_artifacts_dir(env["task"]["id"])
    assert not arts.is_dir() or list(arts.iterdir()) == []
