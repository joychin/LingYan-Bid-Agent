"""workbench API：work/ 工作台 列表/读取/编辑（409+force+修订标记+恢复点栈）。

工作台不是 Artifact（无索引无事件）；parse/ 只读；写走探测+裁决+恢复点兜底
（restorepoints/ 保留 3 个，旧单一 .bak 收编）。
"""



from app import artifact_store
from tests.util import create_task


def _seed_out(task_id: str, rel: str, text: str) -> None:
    p = artifact_store.work_dir(task_id) / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")


def test_list_filters_and_flags(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_out(tid, "analysis/disqualification.md", "<!-- tender-analysis | 节=disqualification | 生成=2026-08-28T00:00:00+00:00 -->\n| 1 | 逾期 | 高 |")
    _seed_out(tid, "analysis/structure.md", "<!-- tender-analysis | 节=structure | 生成=2026-08-28T00:00:00+00:00 -->\n| 项目名称 | 测试 |")
    _seed_out(tid, "parse/招标文件.docx/招标文件.docx.md", "# 招标文件\n正文")
    _seed_out(tid, "parse/招标文件.docx/招标文件.docx.outline.json", "[]")
    _seed_out(tid, "outline/fragments/.hidden.md", "x")

    r = client.get("/api/workbench", params={"task_id": tid})
    assert r.status_code == 200
    files = {f["path"]: f for f in r.json()["files"]}
    # json/隐藏文件不进列表；md 全进
    assert set(files) == {
        "analysis/disqualification.md",
        "analysis/structure.md",
        "parse/招标文件.docx/招标文件.docx.md",
    }
    assert files["parse/招标文件.docx/招标文件.docx.md"]["editable"] is False
    assert files["analysis/disqualification.md"]["editable"] is True
    assert files["analysis/disqualification.md"]["revised"] is False
    assert "mtime" in files["analysis/structure.md"] and "size" in files["analysis/structure.md"]


def test_read_and_containment(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_out(tid, "analysis/evaluation.md", "<!-- x -->\n内容")
    r = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/evaluation.md"})
    assert r.status_code == 200
    d = r.json()
    assert d["content"].startswith("<!-- x -->")
    assert len(d["hash"]) == 64 and d["editable"] is True and d["has_restore"] is False

    # meta 端点：轻量探测只回哈希与标志，不含全文
    m = client.get("/api/workbench/meta", params={"task_id": tid, "path": "analysis/evaluation.md"}).json()
    assert m["hash"] == d["hash"] and m["editable"] is True and m["has_restore"] is False and "content" not in m

    # 越界（../ 逃出 out/）与非 md 一律 404
    for bad in ("../files/a.md", "analysis/../../x.md", "analysis/x.json", "nope.md"):
        rr = client.get("/api/workbench/content", params={"task_id": tid, "path": bad})
        assert rr.status_code == 404, bad


def test_write_stamps_revised_and_conflict(client):
    task = create_task(client)["task"]
    tid = task["id"]
    # 存量旧文件形态（2026-09-04 前的模型写头）：修订标记仍追加进原注释内——旧文件兼容回归
    _seed_out(tid, "analysis/structure.md", "<!-- tender-analysis | 节=structure | 生成=2026-08-28T00:00:00+00:00 -->\n旧内容\n")

    base = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/structure.md"}).json()

    # base_hash 过期 → 409
    r = client.put(
        "/api/workbench/content",
        json={"task_id": tid, "path": "analysis/structure.md", "content": "用户改的\n", "base_hash": "0" * 64},
    )
    assert r.status_code == 409

    # 正确 base → 成功 + 修订标记盖进首行注释 + .bak 留底
    # （编辑器保存全量内容——旧文件的首行注释随内容保留、标记追加其内）
    edited = "<!-- tender-analysis | 节=structure | 生成=2026-08-28T00:00:00+00:00 -->\n用户改的\n"
    r = client.put(
        "/api/workbench/content",
        json={"task_id": tid, "path": "analysis/structure.md", "content": edited, "base_hash": base["hash"]},
    )
    assert r.status_code == 200
    d = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/structure.md"}).json()
    first = d["content"].splitlines()[0]
    assert first.startswith("<!--") and "修订=用户" in first and "节=structure" in first
    assert d["revised"] is True and d["has_restore"] is True
    # 恢复点目录不进列表（列表只收 .md；.bak 后缀不匹配）
    listed = client.get("/api/workbench", params={"task_id": tid}).json()["files"]
    assert all(not f["path"].endswith(".bak") for f in listed)

    # force 覆盖过期 base 也成功
    r = client.put(
        "/api/workbench/content",
        json={"task_id": tid, "path": "analysis/structure.md", "content": "再改\n", "base_hash": "0" * 64, "force": True},
    )
    assert r.status_code == 200

    # 无注释头的文件：前插一行注释
    _seed_out(tid, "outline/fragments/技术标.md", "- 封面\n- 目录\n")
    r = client.put(
        "/api/workbench/content",
        json={"task_id": tid, "path": "outline/fragments/技术标.md", "content": "- 封面\n", "base_hash": "0" * 64, "force": True},
    )
    assert r.status_code == 200
    d = client.get("/api/workbench/content", params={"task_id": tid, "path": "outline/fragments/技术标.md"}).json()
    assert d["content"].splitlines()[0].startswith("<!-- 工作文件 | 修订=用户")
    assert d["content"].splitlines()[1] == "- 封面"


def test_parse_readonly(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_out(tid, "parse/a.pdf/a.pdf.md", "# 原文")
    r = client.put(
        "/api/workbench/content",
        json={"task_id": tid, "path": "parse/a.pdf/a.pdf.md", "content": "篡改", "base_hash": "0" * 64, "force": True},
    )
    assert r.status_code == 403
    assert "只读" in r.json()["detail"]


def test_restore_roundtrip(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_out(tid, "analysis/evaluation.md", "<!-- x -->\n版本1\n")
    h1 = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/evaluation.md"}).json()["hash"]
    client.put("/api/workbench/content", json={"task_id": tid, "path": "analysis/evaluation.md", "content": "<!-- x -->\n版本2\n", "base_hash": h1})

    r = client.post("/api/workbench/restore", json={"task_id": tid, "path": "analysis/evaluation.md"})
    assert r.status_code == 200
    d = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/evaluation.md"}).json()
    assert "版本1" in d["content"]

    # 恢复可再撤销：再 restore 回版本2（恢复前内容已入栈成为最新恢复点）
    r = client.post("/api/workbench/restore", json={"task_id": tid, "path": "analysis/evaluation.md"})
    d = client.get("/api/workbench/content", params={"task_id": tid, "path": "analysis/evaluation.md"}).json()
    assert "版本2" in d["content"]

    # 无恢复点 → 409
    _seed_out(tid, "analysis/structure.md", "无备份文件\n")
    assert client.post("/api/workbench/restore", json={"task_id": tid, "path": "analysis/structure.md"}).status_code == 409


def test_restore_point_stack_keeps_three(client):
    """恢复点栈：连续写 4 次保留最近 3 个；恢复点目录不出现在列表/本轮文件口径。"""
    from pathlib import Path

    from app.api import workbench as wb

    task = create_task(client)["task"]
    tid = task["id"]
    rel = "analysis/evaluation.md"
    _seed_out(tid, rel, "<!-- x -->\nv0\n")
    target: Path = artifact_store.work_dir(tid) / rel

    for i in range(1, 5):
        base = client.get("/api/workbench/content", params={"task_id": tid, "path": rel}).json()["hash"]
        r = client.put(
            "/api/workbench/content",
            json={"task_id": tid, "path": rel, "content": f"<!-- x -->\nv{i}\n", "base_hash": base},
        )
        assert r.status_code == 200

    points = wb._restore_points(target)
    assert len(points) == 3  # v0 被裁掉；栈内为写 v2/v3/v4 前的内容（即 v1/v2/v3，已盖修订标记）
    assert points[0].read_text(encoding="utf-8").endswith("v1\n")
    assert points[-1].read_text(encoding="utf-8").endswith("v3\n")

    # 恢复点目录不进列表
    listed = client.get("/api/workbench", params={"task_id": tid}).json()["files"]
    assert [f["path"] for f in listed] == [rel]


def test_legacy_bak_adopted(client):
    """旧版单一 .bak：首次 push 或 restore 时收编为栈内最旧一条（不删数据）。"""
    from pathlib import Path

    from app.api import workbench as wb

    task = create_task(client)["task"]
    tid = task["id"]
    rel = "analysis/structure.md"
    _seed_out(tid, rel, "新内容\n")
    target: Path = artifact_store.work_dir(tid) / rel
    legacy = target.with_name(target.name + ".bak")
    legacy.write_text("旧世界的备份\n", encoding="utf-8")

    # has_restore 把 legacy 算在内
    d = client.get("/api/workbench/content", params={"task_id": tid, "path": rel}).json()
    assert d["has_restore"] is True

    # restore 触发收编：恢复到 legacy 内容
    r = client.post("/api/workbench/restore", json={"task_id": tid, "path": rel})
    assert r.status_code == 200 and "旧世界的备份" in r.json()["content"]
    assert not legacy.exists()  # 已收编进栈
    points = wb._restore_points(target)
    assert len(points) == 2  # legacy（最旧）+ 恢复前的当前内容
    assert points[0].read_text(encoding="utf-8") == "旧世界的备份\n"


# ---------- docx 正文（只读列表 + 文本视图端点 + 写拒绝） ----------


def _seed_docx(task_id: str, rel: str, title: str = "3.1 需求分析", paragraphs: str = "正文第一段。") -> None:
    from docx import Document

    p = artifact_store.work_dir(task_id) / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, 1)
    for line in paragraphs.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())
    doc.save(p)


def test_docx_listed_readonly_flag(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_docx(tid, "body/3.1 需求分析.docx")
    r = client.get("/api/workbench", params={"task_id": tid})
    files = {f["path"]: f for f in r.json()["files"]}
    assert "body/3.1 需求分析.docx" in files
    entry = files["body/3.1 需求分析.docx"]
    assert entry["editable"] is False and entry["revised"] is False
    assert entry["has_restore"] is False


def test_docx_view_endpoint(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_docx(tid, "body/技术部分/3.1 需求分析.docx", paragraphs="正文第一段。\n正文第二段。")
    r = client.get("/api/workbench/docx-view", params={"task_id": tid, "path": "body/技术部分/3.1 需求分析.docx"})
    assert r.status_code == 200
    d = r.json()
    assert "[P1]（Heading 1）3.1 需求分析" in d["lines"]
    assert "[P2]（Normal）正文第一段。" in d["lines"]
    assert d["abs_path"].endswith("3.1 需求分析.docx")
    # md 不走本端点
    _seed_out(tid, "analysis/evaluation.md", "内容")
    r2 = client.get("/api/workbench/docx-view", params={"task_id": tid, "path": "analysis/evaluation.md"})
    assert r2.status_code == 400


def test_docx_rejected_on_md_endpoints(client):
    task = create_task(client)["task"]
    tid = task["id"]
    _seed_docx(tid, "body/3.1 需求分析.docx")
    qs = {"task_id": tid, "path": "body/3.1 需求分析.docx"}
    assert client.get("/api/workbench/meta", params=qs).status_code == 400
    assert client.get("/api/workbench/content", params=qs).status_code == 400
    body = {"task_id": tid, "path": "body/3.1 需求分析.docx", "content": "x", "base_hash": "", "force": True}
    assert "Word" in client.put("/api/workbench/content", json=body).json()["detail"]
    assert client.post("/api/workbench/restore", json={"task_id": tid, "path": qs["path"]}).status_code == 400
