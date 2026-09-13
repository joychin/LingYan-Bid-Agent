"""知识库 API（v3 章节块模型）：上传归一/role 过滤/能力档位/说明/章节块端点/
排除开关/重建章节/确认/图片清单/全局视图。"""

import io
import json


def _upload(client, name: str, content: bytes):
    return client.post(
        "/api/kb/files",
        files={"file": (name, io.BytesIO(content), "application/octet-stream")},
    )


def _types(client):
    return {t["code"]: t for t in client.get("/api/kb/types").json()["types"]}


def _wait_ready(client, kid: str, timeout_s: float = 5.0):
    """等后台入库管线收敛（TestClient 里 fire-and-forget 会跨请求执行）。"""
    import time

    deadline = time.time() + timeout_s
    while time.time() < deadline:
        item = client.get(f"/api/kb/items/{kid}").json()
        if item.get("parse_status") in ("ready", "failed"):
            return item
        time.sleep(0.05)
    raise AssertionError("入库管线 5s 内未收敛")


def test_upload_no_bucket_param(client):
    """上传归一：不带任何分类参数；响应无 bucket。"""
    r = _upload(client, "a.txt", b"# t\n\ncontent")
    assert r.status_code == 201
    assert set(r.json()) == {"id", "file_name", "size"}


def test_types_payload_roles(client):
    types = _types(client)
    assert types["past_proposal"]["role"] == "writing"
    assert types["financial_report"]["role"] == "fact"
    # 社保缴纳证明（2026-09-08 补）：事实类、period 时间锚点，不再落 other
    assert types["social_security"]["role"] == "fact"
    assert types["social_security"]["time_fields"] == ["period"]
    assert "decompose" not in types["past_proposal"]  # 策略字段已删（role 驱动）
    assert "field_labels" in client.get("/api/kb/types").json()


def test_items_role_filter_and_capability(client):
    from app import db

    _upload(client, "证书.txt", b"# c\n\ncontent")
    _upload(client, "标书范文.txt", b"# s\n\ncontent")
    db.kb_update_item(
        next(it["id"] for it in db.kb_list_items() if it["file_name"] == "标书范文.txt"),
        doc_type="past_proposal",
    )
    # role 过滤（分组视图）
    writing = client.get("/api/kb/items", params={"role": "writing"}).json()["items"]
    assert [it["file_name"] for it in writing] == ["标书范文.txt"]
    assert writing[0]["role"] == "writing"
    fact = client.get("/api/kb/items", params={"role": "fact"}).json()["items"]
    assert all(it["role"] == "fact" for it in fact)
    # 能力档位（素材分离后无 enriched 档）；无素材字段
    assert all(it["capability"] in ("stored", "searchable", "typed") for it in fact + writing)
    assert all("material_count" not in it for it in fact + writing)
    # 素材端点已随素材分离删除
    assert client.get("/api/kb/materials").status_code in (404, 405)


def test_item_statement_and_freshness(client):
    from app import db

    r = _upload(client, "老证书.txt", b"# c\n\ncontent")
    kid = r.json()["id"]
    _wait_ready(client, kid)
    db.kb_update_item(kid, doc_type="qualification_certificate", parse_status="ready",
                      extract_status="done", suggested_metadata=json.dumps({
                          "doc_type": "qualification_certificate",
                          "statement": "ISO 认证，有效期至 2020-01-01（第1页）。",
                          "fields": {"valid_until": {"value": "2020-01-01", "source": "第1页"}},
                      }, ensure_ascii=False))
    item = next(it for it in client.get("/api/kb/items").json()["items"] if it["id"] == kid)
    assert item["capability"] == "typed"
    assert "ISO 认证" in item["suggested"]["statement"]
    assert {f["kind"] for f in item["freshness"]} == {"expired"}


def test_content_meta_image_count(client):
    from app.knowledge.ingest import run_ingest

    body = "# 概况\n\n" + "正文内容补充。" * 30
    r = _upload(client, "meta.txt", body.encode())
    run_ingest(r.json()["id"])
    meta = client.get(f"/api/kb/items/{r.json()['id']}/content").json()["meta"]
    assert meta["conversion"] == "txt-passthrough" and meta["chars"] > 0


def _docx_bytes(
    paras: int = 20,
    first: str = "一、方案",
    line_tpl: str = "运维服务方案内容文本第{i}行补充说明。",
) -> bytes:
    """最小素材 docx（一级标题 + N 段正文）——上传白名单收口为仅 .docx 后的 API 载体。"""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading(first, level=1)
    for i in range(1, paras + 1):
        doc.add_paragraph(line_tpl.format(i=i))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_kb_materials_endpoints_removed(client):
    """素材分离：知识库侧素材端点全部删除（素材域在 /api/materials）。"""
    r = _upload(client, "x.txt", b"# x\n\ncontent")
    kid = r.json()["id"]
    assert client.get(f"/api/kb/items/{kid}/materials").status_code in (404, 405)
    assert client.post(f"/api/kb/items/{kid}/decompose").status_code in (404, 405)
    assert client.put(f"/api/kb/items/{kid}/materials/m/excluded", json={"excluded": True}).status_code in (404, 405)


def test_materials_files_api_flow(client):
    """素材域 API 全流程：上传→解析→目录树→建块（多区间）→列表→改备注→删块→删文件。"""
    import hashlib
    import time

    body = _docx_bytes()
    r = client.post(
        "/api/materials/files",
        files={"file": ("库标书.docx", body, "application/octet-stream")},
    )
    assert r.status_code == 201
    fid = r.json()["id"]
    # 同 hash 重复上传 → 409
    dup = client.post(
        "/api/materials/files",
        files={"file": ("库标书.docx", body, "application/octet-stream")},
    )
    assert dup.status_code == 409

    # 解析收敛（TestClient 后台任务跨请求执行）
    deadline = time.time() + 5
    while time.time() < deadline:
        outline = client.get(f"/api/materials/files/{fid}/outline").json()
        if outline["parse_status"] in ("ready", "failed"):
            break
        time.sleep(0.05)
    assert outline["parse_status"] == "ready"
    assert outline["outline"] and outline["outline"][0]["标题"] == "一、方案"

    # 建块：多区间（勾两段不连续区间——用树区间与手造区间）
    block = client.post(
        f"/api/materials/files/{fid}/blocks",
        json={"title": "运维方案块", "note": "政务云运维", "ranges": [[1, 5], [10, 15]]},
    )
    assert block.status_code == 201
    bid = block.json()["id"]
    assert block.json()["ranges"] == [[1, 5], [10, 15]] and block.json()["chars"] > 0
    # 非法 ranges → 422
    assert client.post(
        f"/api/materials/files/{fid}/blocks",
        json={"title": "x", "note": "", "ranges": []},
    ).status_code == 422

    # 文件列表带块数；块列表带来源文件
    files = client.get("/api/materials/files").json()["files"]
    assert next(f for f in files if f["id"] == fid)["block_count"] == 1
    blocks = client.get("/api/materials/blocks").json()["blocks"]
    assert blocks[0]["file_name"] == "库标书.docx" and blocks[0]["note"] == "政务云运维"

    # 改备注 → 删块 → 删文件（连带）
    assert client.put(f"/api/materials/blocks/{bid}", json={"note": "改后备注"}).json()["note"] == "改后备注"
    assert client.delete(f"/api/materials/blocks/{bid}").status_code == 200
    assert client.get("/api/materials/blocks").json()["blocks"] == []
    assert client.delete(f"/api/materials/files/{fid}").status_code == 200
    assert client.get("/api/materials/files").json()["files"] == []
    assert hashlib.sha256(b"gone")  # 保持 import 使用


def test_materials_upload_docx_only(client):
    """上传白名单收口为仅 .docx：其余格式 400 + 人话出路（Word 另存为 docx）。"""
    r = client.post(
        "/api/materials/files",
        files={"file": ("范文.pdf", b"%PDF-1.4 fake", "application/octet-stream")},
    )
    assert r.status_code == 400
    detail = r.json()["detail"]
    assert ".docx" in detail and "另存为" in detail


def test_materials_file_content_endpoint(client):
    """挑章节预览端点：按行号切片 / 越界 clamp / 参数非法 422 / 404 / 未 ready 422。"""
    import time

    from app import db

    r = client.post(
        "/api/materials/files",
        files={"file": ("预览标书.docx", _docx_bytes(paras=40, first="预览标书"), "application/octet-stream")},
    )
    assert r.status_code == 201
    fid = r.json()["id"]
    deadline = time.time() + 5
    outline: dict = {}
    while time.time() < deadline:
        outline = client.get(f"/api/materials/files/{fid}/outline").json()
        if outline["parse_status"] in ("ready", "failed"):
            break
        time.sleep(0.05)
    assert outline["parse_status"] == "ready"
    assert outline["outline"], "docx 应解析出目录树"

    # 正常切片（闭区间；正文以标题行起）
    content = client.get(f"/api/materials/files/{fid}/content", params={"start": 2, "end": 4}).json()
    assert content["sections"][0]["start"] == 2 and content["sections"][0]["end"] == 4
    assert content["sections"][0]["text"]
    assert content["chars"] > 0

    # 越界 clamp：end 超总行数但跨度合法 → 夹到实际行数；跨度超上限 → 422
    total = outline["outline"][0]["end_line"]
    over = client.get(f"/api/materials/files/{fid}/content", params={"start": total - 2, "end": total + 50}).json()
    assert over["sections"] and over["sections"][0]["end"] == total
    assert client.get(f"/api/materials/files/{fid}/content", params={"start": 1, "end": 99999}).status_code == 422

    # 参数非法 → 422；文件不存在 → 404；未解析（pending）→ 422
    assert client.get(f"/api/materials/files/{fid}/content", params={"start": 0, "end": 5}).status_code == 422
    assert client.get(f"/api/materials/files/{fid}/content", params={"start": 5, "end": 2}).status_code == 422
    assert client.get("/api/materials/files/mt_nope/content", params={"start": 1, "end": 2}).status_code == 404
    fid2 = db.mt_insert_file(file_name="未解析预览.docx", file_hash="h-preview")
    assert client.get(f"/api/materials/files/{fid2['id']}/content", params={"start": 1, "end": 2}).status_code == 422


def test_materials_reparse_file(client):
    """reparse 端点：失败文件重试恢复 + 404 不存在 + 409 正在解析（单飞）。"""
    import time

    from app import db
    from app.knowledge import materials_lib as mlib

    r = client.post(
        "/api/materials/files",
        files={"file": ("重试标书.docx", _docx_bytes(paras=12), "application/octet-stream")},
    )
    fid = r.json()["id"]

    def _wait_terminal(left: str | None = None) -> dict:
        """等后台解析收敛；left=须先观察到离开的旧状态（防 reparse 任务未启动
        就读到重跑前的终态快照）。"""
        deadline = time.time() + 5
        while time.time() < deadline:
            f = db.mt_get_file(fid)
            if left is not None:
                if f["parse_status"] == left:
                    time.sleep(0.05)
                    continue
                left = None
            if f["parse_status"] in ("ready", "failed"):
                return f
            time.sleep(0.05)
        raise AssertionError("解析 5s 内未收敛")

    assert _wait_terminal()["parse_status"] == "ready"

    # 模拟历史失败（如解析程序 bug）→ reparse 恢复
    db.mt_update_file(fid, parse_status="failed", error="解析失败：name 'qn' is not defined")
    rr = client.post(f"/api/materials/files/{fid}/reparse")
    assert rr.status_code == 202 and rr.json() == {"ok": True}
    f = _wait_terminal(left="failed")
    assert f["parse_status"] == "ready" and f["error"] is None

    # 404 不存在；409 单飞——真实形态=parse_status=parsing + in-flight 标记双条件
    # （只剩标记而状态已终态=收尾调度窗口的陈旧标记，schedule_parse 自愈放行；
    #  2026-09-08 全量套件负载下实测该窗口曾致本测试偶发 409/202 竞态）
    assert client.post("/api/materials/files/mt_nope/reparse").status_code == 404
    db.mt_update_file(fid, parse_status="parsing", error=None)
    mlib._inflight.add(fid)
    try:
        assert client.post(f"/api/materials/files/{fid}/reparse").status_code == 409
    finally:
        mlib._inflight.discard(fid)
        db.mt_update_file(fid, parse_status="ready", error=None)


def test_materials_file_raw_endpoint(client):
    """原件字节端点（预览「原件」模式）：200=全量字节+下载名 / 404 不存在。"""
    body = _docx_bytes(paras=10, first="原件预览标书")
    r = client.post(
        "/api/materials/files",
        files={"file": ("原件预览.docx", body, "application/octet-stream")},
    )
    assert r.status_code == 201
    fid = r.json()["id"]
    raw = client.get(f"/api/materials/files/{fid}/raw")
    assert raw.status_code == 200
    assert raw.content == body
    # 中文文件名经 RFC 5987 编码（filename*=utf-8''…），断言形态而非原文
    assert "attachment" in raw.headers.get("content-disposition", "")
    assert raw.headers.get("content-disposition", "").endswith(".docx")
    assert client.get("/api/materials/files/mt_nope/raw").status_code == 404
    assert client.delete(f"/api/materials/files/{fid}").status_code == 200


def test_item_images_listing(client):
    """图片清单端点：内容页折叠区数据源（图片仅供查看，与素材无关）。"""
    from app.knowledge import store

    r = _upload(client, "无图.txt", b"# t\n\ncontent")
    kid = r.json()["id"]
    data = client.get(f"/api/kb/items/{kid}/images").json()
    assert data["images"] == []
    # 造两张盘上图片（绕过抽取，聚焦端点）
    img_dir = store.kb_images_dir("无图.txt")
    img_dir.mkdir(parents=True, exist_ok=True)
    (img_dir / "img_001.png").write_bytes(b"\x89PNG-fake")
    (img_dir / "img_002.png").write_bytes(b"\x89PNG-fake2")
    imgs = client.get(f"/api/kb/items/{kid}/images").json()["images"]
    assert [i["name"] for i in imgs] == ["img_001.png", "img_002.png"]
    assert imgs[0]["size"] == len(b"\x89PNG-fake")


def test_metadata_confirm_with_statement(client):
    from app import db

    r = _upload(client, "范文2.txt", b"# s\n\ncontent")
    kid = r.json()["id"]
    _wait_ready(client, kid)
    db.kb_update_item(kid, parse_status="ready", doc_type="other")

    ok = client.put(
        f"/api/kb/items/{kid}/metadata",
        json={"doc_type": "past_proposal",
              "statement": "共 3 章技术标。",
              "fields": {"project_name": "某园区项目"},
              "extra": {"来源": "2024 中标项目"}},
    )
    assert ok.status_code == 200
    data = ok.json()
    assert data["review_status"] == "confirmed" and data["role"] == "writing"
    assert data["business"]["statement"] == "共 3 章技术标。"


def test_metadata_confirm_with_questions(client):
    """PUT 检索问题：落 business + §questions 段可检中；坏形状/超长/超条数 422。"""
    from app import db
    from app.knowledge import fts

    r = _upload(client, "医院合同.txt", "# 合同\n\nXX医院智慧后勤平台项目，金额 380 万元。".encode("utf-8"))
    kid = r.json()["id"]
    _wait_ready(client, kid)
    db.kb_update_item(kid, parse_status="ready", doc_type="contract_case")

    ok = client.put(
        f"/api/kb/items/{kid}/metadata",
        json={"doc_type": "contract_case",
              "statement": "XX 医院智慧后勤平台合同（第1页）。",
              "questions": ["做过哪些医疗行业项目？", " 合同金额多大？ ", ""]},
    )
    assert ok.status_code == 200
    data = ok.json()
    assert data["business"]["questions"] == ["做过哪些医疗行业项目？", "合同金额多大？"]
    hits = db.kb_search_segments(fts.build_match_expr("医疗行业"), limit=5)
    assert any(h["item_id"] == kid and h.get("section_path") == "§questions" for h in hits)

    assert client.put(
        f"/api/kb/items/{kid}/metadata",
        json={"doc_type": "contract_case", "questions": "不是数组"},
    ).status_code == 422
    assert client.put(
        f"/api/kb/items/{kid}/metadata",
        json={"doc_type": "contract_case", "questions": ["超" * 41]},
    ).status_code == 422
    assert client.put(
        f"/api/kb/items/{kid}/metadata",
        json={"doc_type": "contract_case", "questions": [f"问题{i}" for i in range(11)]},
    ).status_code == 422


def test_delete_and_badge(client):
    r = _upload(client, "gone.txt", b"# g\n\ndel")
    kid = r.json()["id"]
    assert client.delete(f"/api/kb/items/{kid}").status_code == 200
    assert client.get(f"/api/kb/items/{kid}").status_code == 404
    assert client.get("/api/kb/badge").json() == {"pending": 0}


def test_materials_content_search_usage_flow(client):
    """完善批端点：块内容分节 / ?q= 正文检索 / outline 字数 / 引用打点透出。"""
    import time

    fid = client.post(
        "/api/materials/files",
        files={
            "file": (
                "完善标书.docx",
                _docx_bytes(paras=20, line_tpl="等保合规建设内容第{i}行补充。"),
                "application/octet-stream",
            )
        },
    ).json()["id"]
    deadline = time.time() + 5
    while time.time() < deadline:
        outline = client.get(f"/api/materials/files/{fid}/outline").json()
        if outline["parse_status"] in ("ready", "failed"):
            break
        time.sleep(0.05)
    assert outline["parse_status"] == "ready"
    assert outline["outline"][0]["chars"] > 0

    bid = client.post(
        f"/api/materials/files/{fid}/blocks",
        json={"title": "等保块", "note": "安全写法", "ranges": [[1, 6]]},
    ).json()["id"]

    # 块内容：分节区间 + 正文；不存在 404
    content = client.get(f"/api/materials/blocks/{bid}/content").json()
    assert content["sections"][0]["start"] == 1
    assert "等保合规建设内容" in content["sections"][0]["text"]
    assert client.get("/api/materials/blocks/blk_none/content").status_code == 404

    # ?q=：正文 FTS 命中（「建设内容」只在正文里）∪ 标题命中；无命中空表
    assert [b["id"] for b in client.get("/api/materials/blocks?q=建设内容").json()["blocks"]] == [bid]
    assert [b["id"] for b in client.get("/api/materials/blocks?q=等保块").json()["blocks"]] == [bid]
    assert client.get("/api/materials/blocks?q=绝不存在的词xyzzy").json()["blocks"] == []

    # 引用打点透出
    from app import db

    db.mt_touch_blocks([bid])
    row = next(b for b in client.get("/api/materials/blocks").json()["blocks"] if b["id"] == bid)
    assert row["use_count"] == 1 and row["last_used_at"]

    assert client.delete(f"/api/materials/files/{fid}").status_code == 200


def test_materials_block_content_preview(client):
    """块内容行内预览预算（2026-09-13 内存修复批）：sections 正文按累计预算截断、
    chars 保真、预算耗尽即停；preview<1 422；不传行为不变（写作指引表等消费方
    不再整块拉全文）。"""
    import time

    fid = client.post(
        "/api/materials/files",
        files={
            "file": (
                "预览预算.docx",
                _docx_bytes(paras=20, line_tpl="等保合规建设内容第{i}行补充。"),
                "application/octet-stream",
            )
        },
    ).json()["id"]
    deadline = time.time() + 5
    while time.time() < deadline:
        outline = client.get(f"/api/materials/files/{fid}/outline").json()
        if outline["parse_status"] in ("ready", "failed"):
            break
        time.sleep(0.05)
    assert outline["parse_status"] == "ready"

    bid = client.post(
        f"/api/materials/files/{fid}/blocks",
        json={"title": "双区块", "note": "", "ranges": [[2, 7], [9, 12]]},
    ).json()["id"]

    full = client.get(f"/api/materials/blocks/{bid}/content").json()
    assert len(full["sections"]) == 2 and full["chars"] > 60

    # 预算只够第一区的一部分：第一区截断、第二区不再出现；chars 仍是真实总字数
    pv = client.get(f"/api/materials/blocks/{bid}/content", params={"preview": 20}).json()
    assert len(pv["sections"]) == 1
    assert pv["sections"][0]["text"] == full["sections"][0]["text"][:20]
    assert pv["chars"] == full["chars"]

    # 预算跨区：两区都保留、合计恰等于预算
    budget = len(full["sections"][0]["text"]) + 5
    pv2 = client.get(f"/api/materials/blocks/{bid}/content", params={"preview": budget}).json()
    assert len(pv2["sections"]) == 2
    assert sum(len(s["text"]) for s in pv2["sections"]) == budget
    assert pv2["sections"][1]["text"] == full["sections"][1]["text"][:5]

    assert client.get(f"/api/materials/blocks/{bid}/content", params={"preview": 0}).status_code == 422
    assert client.delete(f"/api/materials/files/{fid}").status_code == 200


def test_check_result_passthrough_and_human_confirm_marker(client):
    """核对结果随条目下发（解析后的 JSON）；PUT 人工确认后 business 无 auto 标记。"""
    from app import db

    r = _upload(client, "核对证书.txt", b"# c\n\ncontent")
    kid = r.json()["id"]
    _wait_ready(client, kid)
    db.kb_update_item(
        kid, parse_status="ready", extract_status="done",
        doc_type="qualification_certificate",
        suggested_metadata=json.dumps({
            "doc_type": "qualification_certificate",
            "fields": {"valid_until": {"value": "2099-01-01"}},
        }, ensure_ascii=False),
        check_result=json.dumps({
            "status": "fail",
            "results": [{"field": "valid_until", "label": "有效期至",
                         "ok": False, "detail": "未在原文找到该日期——疑似抽取有误，请核对"}],
        }, ensure_ascii=False),
    )
    item = next(it for it in client.get("/api/kb/items").json()["items"] if it["id"] == kid)
    assert item["review_status"] == "pending_review"
    assert item["check_result"]["status"] == "fail"
    assert "未在原文找到该日期" in item["check_result"]["results"][0]["detail"]

    ok = client.put(f"/api/kb/items/{kid}/metadata",
                    json={"doc_type": "qualification_certificate",
                          "fields": {"valid_until": "2028-06-30"}})
    assert ok.status_code == 200
    data = ok.json()
    assert data["review_status"] == "confirmed"
    assert data["business"].get("confirmed_by") is None  # 人工语义：无 auto 标记
