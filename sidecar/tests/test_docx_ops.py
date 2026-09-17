"""docx 正文工具族：建节/读视图/素材元素注入/带修订标记的定向修订。

覆盖链路：素材 docx 上传解析（产 element_map）→ 勾选建块 → 节文件创建 →
元素注入（表格合并单元格/图片关系保真）→ 修订标记（跨 run 替换/插入/删除，
拒绝全部修订=原文的自校验）→ 错误路径（非 docx 素材/路径越界/find 不命中/
既有修订标记段落）。
"""

import json
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app import runctx
from app.knowledge import materials_lib as mlib
from app.tools.docx_ops import (
    _accept_revisions_inplace,
    _accepted_text,
    _flatten_rejected,
    _flow_to_mermaid,
    docx_assemble_volume,
    docx_comment_add,
    docx_diagram_insert,
    docx_image_insert,
    docx_material_inject,
    docx_section_create,
    docx_section_read,
    docx_section_revise,
    docx_source_inject,
    section_text_lines,
)
from tests.util import init_env

COMPANY = "北京华信科技有限公司"


def _tiny_png(path: Path) -> None:
    """1×1 红点 PNG（图片关系迁移的最小载体）。"""

    def chunk(t: bytes, d: bytes) -> bytes:
        c = t + d
        return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b"")
    )


def _make_material_docx(path: Path, png: Path) -> None:
    """素材原件：标题 + 跨 run 公司名段 + 合并单元格表格 + 图片段。"""
    doc = Document()
    doc.add_heading("运维服务方案", 1)
    p = doc.add_paragraph()
    p.add_run("本项目由")
    p.add_run("北京华信")
    p.add_run("科技有限公司承建，服务期三年。")
    t = doc.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "人员角色"
    t.cell(0, 1).merge(t.cell(0, 2)).text = "配置要求"
    t.cell(1, 0).merge(t.cell(2, 0)).text = "项目经理"
    t.cell(1, 1).text = "1 人"
    t.cell(1, 2).text = "PMP"
    t.cell(2, 1).text = "工程师"
    t.cell(2, 2).text = "5 人"
    doc.add_picture(str(png), width=914400 // 2)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _make_tender_docx(path: Path) -> None:
    """招标原件：两章结构 + 投标函空栏段 + 含空格子的报价表（格式件拷贝/填空载体）。"""
    doc = Document()
    doc.add_heading("第一章 招标公告", 1)
    doc.add_paragraph("本项目公开招标，欢迎合格的投标人参加投标。")
    doc.add_heading("第四章 投标文件格式", 1)
    doc.add_heading("投标函", 2)
    doc.add_paragraph("致：______（招标人名称）")
    t = doc.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "项目名称"
    t.rows[0].cells[1].text = ""
    t.rows[0].cells[2].text = "报价（元）"
    t.rows[1].cells[0].text = ""
    t.rows[1].cells[1].text = "报价大写"
    t.rows[1].cells[2].text = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _tender_source(env) -> Path:
    """把招标原件放进任务 sources/，返回绝对路径。"""
    from app.artifact_store import sources_dir

    p = sources_dir(env["task"]["id"]) / "招标文件.docx"
    _make_tender_docx(p)
    return p


@pytest.fixture
def env(tmp_path, monkeypatch):
    task, conv = init_env(tmp_path, monkeypatch)
    runctx.set_run(conv["id"], "r_test", task["id"])
    mlib.ensure_dirs()
    png = tmp_path / "dot.png"
    _tiny_png(png)
    src = mlib.mt_files_dir() / "历史运维方案.docx"
    _make_material_docx(src, png)
    from app import db

    f = db.mt_insert_file("历史运维方案.docx", "hash_docx_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(
        f["id"], "运维整章", "历史项目运维章节",
        ranges=[[1, 10 ** 6]],  # 全文（_squash_ranges 会夹紧到实际行数）
    )
    yield {"task": task, "conv": conv, "file": f, "block": block, "png": png}
    runctx.clear_run()


def _make_section(name: str = "技术部分/3.1 运维方案.docx") -> str:
    rel = name if name.startswith("body/") else f"body/{name}"
    r = docx_section_create.invoke({"path": rel, "title": "3.1 运维方案"})
    assert r.startswith("[已创建]"), r
    return rel


# ---------- 建节与读视图 ----------

def test_create_and_read(env):
    from app.artifact_store import work_dir

    r = docx_section_create.invoke({"path": "body/技术部分/3.1 需求分析", "title": "3.1 需求分析", "paragraphs": "第一段。\n第二段。"})
    assert r.startswith("[已创建]") and "正文 2 段" in r
    # 路径自动补 .docx
    assert (work_dir(env["task"]["id"]) / "body" / "技术部分" / "3.1 需求分析.docx").is_file()
    view = docx_section_read.invoke({"path": "body/技术部分/3.1 需求分析.docx"})
    assert "[P1]（Heading 1）3.1 需求分析" in view
    assert "[P2]（Tender Body）第一段。" in view
    assert "共 3 个段落" in view


def test_create_refuses_overwrite(env):
    _make_section()
    r = docx_section_create.invoke({"path": "body/技术部分/3.1 运维方案.docx", "title": "x"})
    assert r.startswith("[创建失败]") and "已存在" in r


def test_path_containment(env):
    assert "只允许落在任务 work/body/" in docx_section_create.invoke({"path": "out/x.docx", "title": "x"})
    assert "路径越界" in docx_section_create.invoke({"path": "body/../../x.docx", "title": "x"})
    assert "只允许落在任务 work/body/" in docx_section_read.invoke({"path": "drafts/x.docx"})


def test_read_missing(env):
    assert "文件不存在" in docx_section_read.invoke({"path": "body/不存在.docx"})


# ---------- 素材注入 ----------

def test_material_inject_full_chain(env):
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[已注入]"), r
    assert "表格 1 张" in r and "图片 1 张" in r
    from app import db
    from app.artifact_store import work_dir

    # 注入落盘成功 → AI 引用打点 +1
    assert db.mt_get_block(env["block"]["id"])["use_count"] == 1

    dst = work_dir(env["task"]["id"]) / section
    chk = Document(str(dst))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "运维服务方案" in texts and COMPANY in texts
    assert len(chk.tables) == 1
    xml = chk.tables[0]._tbl.xml
    assert "vMerge" in xml and "gridSpan" in xml  # 合并单元格随元素拷贝保留
    for blip in chk.element.body.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        assert rid in chk.part.related_parts  # 图片关系完好


def test_material_inject_rebuilds_element_map(env):
    """旧解析产物无 element_map：注入自动幂等重跑补产。"""
    mlib.mt_parse_dir("历史运维方案.docx").joinpath("element_map.json").unlink()
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[已注入]"), r
    assert mlib.read_element_map("历史运维方案.docx")


def test_material_inject_rejects_stale_element_map_v1(env):
    """坐标系修复（2026-09-10 表格逐行）前的 v1 旧 map 被版本护栏拒收：
    读侧返回 None → 注入走既有「幂等重跑 + md 漂移比对」路径换得 v2 正确坐标
    （重跑 md 逐字节不变，比对放行）。手搓错位坐标证明是护栏在挡而非碰巧一致。
    """
    mp = mlib.mt_parse_dir("历史运维方案.docx") / "element_map.json"
    stale = json.loads(mp.read_text(encoding="utf-8"))
    assert stale["version"] >= 2  # 新解析产物已是现行版本
    stale["version"] = 1
    stale["element_lines"] = [[el, s + 100, e + 100] for el, s, e in stale["element_lines"]]
    mp.write_text(json.dumps(stale, ensure_ascii=False), encoding="utf-8")
    assert mlib.read_element_map("历史运维方案.docx") is None
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[已注入]"), r
    fresh = json.loads(mp.read_text(encoding="utf-8"))
    assert fresh["version"] >= 2 and mlib.read_element_map("历史运维方案.docx") is not None


def test_material_inject_requires_docx_source(env, tmp_path):
    """pdf/txt 等原件无可注入元素：明确拒绝并指路文本方式。"""
    from app import db

    txt = mlib.mt_files_dir() / "服务清单.txt"
    txt.write_text("# 服务清单\n\n提供驻场运维服务。\n", encoding="utf-8")
    f = db.mt_insert_file("服务清单.txt", "hash_txt_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(f["id"], "服务清单", "", ranges=[[1, 100]])
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": block["id"], "dest": section})
    assert r.startswith("[注入失败]") and "不是 docx" in r


def test_material_inject_missing_block(env):
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": "blk_000000000000", "dest": section})
    assert r.startswith("[注入失败]") and "不存在" in r


# ---------- 招标原件拷贝（docx_source_inject） ----------


def test_material_inject_body_style_adoption(env):
    """素材适配模板（2026-09-08 用户拍板：素材库拷贝归顺模板、招标格式件
    保真）：无样式引用的素材正文段挂 Tender Body（标书缩进/行距，与 AI 正文
    同观感——此前吃中性 Normal 不缩进是观感割裂主因）；表格整表与表内文字
    不受影响；素材标题段引用内置 Heading 自动吃宿主（模板）定义；招标格式件
    拷贝不挂样式（格式由招标文件定死）。"""
    section = _make_section()
    assert docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": section}
    ).startswith("[已注入]")
    doc = Document(str(_abs(env, section)))
    body = [p for p in doc.paragraphs if p.text.startswith("本项目由")]
    assert body and all(p.style.name == "Tender Body" for p in body)
    assert any(
        p.text == "运维服务方案" and p.style.name == "Heading 1" for p in doc.paragraphs
    )
    # 表格整表保真：单元格段落不被挂样式（无显式引用，渲染吃 Normal）
    cell_p = doc.tables[0].rows[0].cells[0].paragraphs[0]._p
    cell_ppr = cell_p.find(qn("w:pPr"))
    assert cell_ppr is None or cell_ppr.find(qn("w:pStyle")) is None

    # 招标格式件：保真不挂（同款无样式正文段保持原样）
    _tender_source(env)
    section2 = _make_section("技术部分/投标函.docx")
    assert docx_source_inject.invoke(
        {"source": "招标文件.docx", "dest": section2}
    ).startswith("[已注入]")
    chk = Document(str(_abs(env, section2)))
    zh = [p for p in chk.paragraphs if p.text == "致：______（招标人名称）"]
    assert zh
    zh_ppr = zh[0]._p.find(qn("w:pPr"))
    assert zh_ppr is None or zh_ppr.find(qn("w:pStyle")) is None


def test_source_inject_whole_file(env):
    _tender_source(env)
    section = _make_section()
    r = docx_source_inject.invoke({"source": "招标文件.docx", "dest": section})
    assert r.startswith("[已注入]") and "整份文件" in r, r
    assert "表格 1 张" in r
    # 返回语不得教写内联占位（2026-09-16 prompt 一致性批）：占位文字
    # validate_body 判不过，缺值正规出口=Word 批注——旧句「未定值【待补：…】」
    # 曾让照做的写手自查必挂一次
    assert "【待补" not in r, r
    assert "docx_comment_add" in r, r
    chk = Document(str(_abs(env, section)))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "第一章 招标公告" in texts and "致：______（招标人名称）" in texts
    assert len(chk.tables) == 1 and chk.tables[0].cell(0, 0).text == "项目名称"

    # 同区间重复注入拦截（2026-09-10 review，与素材侧同口径）：append 语义下
    # 重拷=格式件翻倍可直达交付合册。第二次同源注入被拦、节文件未被追加
    r2 = docx_source_inject.invoke({"source": "招标文件.docx", "dest": section})
    assert r2.startswith("[注入失败]") and "已注入过" in r2 and "replace=true" in r2
    chk2 = Document(str(_abs(env, section)))
    assert len(chk2.paragraphs) == len(chk.paragraphs) and len(chk2.tables) == 1

    # 不同节不受影响（拦截按目标节当前内容比对，非全局记账）
    other = _make_section("技术部分/投标函副本.docx")
    assert docx_source_inject.invoke({"source": "招标文件.docx", "dest": other}).startswith("[已注入]")


def test_source_inject_line_range(env):
    """行号区间拷（L 前缀容忍）：只带第四章格式章，不带第一章公告。"""
    from app.parse import convert as parse_convert

    src = _tender_source(env)
    md_lines = parse_convert(src).md.splitlines()
    start = next(i for i, ln in enumerate(md_lines, 1) if ln.startswith("# 第四章"))
    section = _make_section()
    r = docx_source_inject.invoke(
        {"source": "招标文件.docx", "dest": section, "lines": f"L{start}-{len(md_lines)}"}
    )
    assert r.startswith("[已注入]") and f"行号区间 {start}-{len(md_lines)}" in r, r
    chk = Document(str(_abs(env, section)))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "第四章 投标文件格式" in texts and "投标函" in texts
    assert "招标公告" not in texts and "欢迎合格的投标人" not in texts
    assert len(chk.tables) == 1  # 报价表随格式章带入


def test_element_lines_coords_match_md_after_tables(tmp_path):
    """坐标系统一回归（2026-09-10）：表格逐行展开后 element_lines 行号=最终 md
    行号。修复前表格整块 append（内嵌换行不占下标），映射坐标在表格后累计
    漂移——表格后段落按 md 行号找不到元素（实测 13 表拉开 140 行）。"""
    from app.parse.docx import convert

    doc = Document()
    doc.add_heading("第一章 公告", 1)
    doc.add_paragraph("公告正文第一段。")
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for r in range(3):
        t.cell(r, 0).text = f"行{r}"
        t.cell(r, 1).text = f"值{r}"
    doc.add_heading("附件14：承诺书", 2)
    doc.add_paragraph("我单位承诺：拟派项目经理准时到位。")
    p = tmp_path / "原件.docx"
    doc.save(p)

    res = convert(p)
    md_lines = res.md.splitlines()
    line_no = next(i for i, ln in enumerate(md_lines, 1) if "我单位承诺" in ln)
    els = res.info["element_lines"]
    assert max(e for _el, s, e in els) == len(md_lines), "映射未覆盖到 md 末行（坐标系漂移）"
    hit = [(s, e) for _el, s, e in els if s <= line_no <= e]
    assert hit, "表格后段落未映射（坐标系漂移回归）"
    assert any("我单位承诺" in "\n".join(md_lines[s - 1 : e]) for s, e in hit)


def test_source_inject_outline_range_after_table(env):
    """端到端：按 outline.json 的行号区间注入「表格之后」的附件——写手探针
    风暴的真实形态（承诺书在多张表后，修复前按 outline 区间注入必「未映射」，
    写手被迫逐段试探）。"""
    from app.artifact_store import sources_dir
    from app.parse import convert as parse_convert
    from app.parse import outline_with_lines

    doc = Document()
    doc.add_heading("第一章 招标公告", 1)
    doc.add_paragraph("公告正文。")
    for ti in range(2):
        t = doc.add_table(rows=3, cols=2)
        t.style = "Table Grid"
        for r in range(3):
            t.cell(r, 0).text = f"表{ti}行{r}"
            t.cell(r, 1).text = "值"
    doc.add_heading("附件14：承诺书", 1)
    doc.add_paragraph("我单位承诺：拟派项目经理将准时到位。")
    sroot = sources_dir(env["task"]["id"])
    sroot.mkdir(parents=True, exist_ok=True)
    doc.save(sroot / "谈判文件.docx")

    def find(tree, kw):
        for n in tree:
            if kw in n["标题"]:
                return n
            hit = find(n.get("children") or [], kw)
            if hit:
                return hit
        return None

    node = find(outline_with_lines(parse_convert(sroot / "谈判文件.docx").md), "承诺书")
    assert node is not None
    section = _make_section()
    r = docx_source_inject.invoke(
        {"source": "谈判文件.docx", "dest": section, "lines": f"{node['start_line']}-{node['end_line']}"}
    )
    assert r.startswith("[已注入]"), r
    chk = Document(str(_abs(env, section)))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "附件14：承诺书" in texts and "我单位承诺" in texts
    assert "招标公告" not in texts and len(chk.tables) == 0  # 表前的表不被误带


def test_material_inject_block_after_table(env, tmp_path):
    """素材块勾选区间落在表格之后：映射须命中正确元素（修复前两套坐标漂移，
    表格后区块注入错元素或「未映射」——素材库「勾选区间需复核」的部分根因）。"""
    from app import db

    doc = Document()
    doc.add_heading("运维方案", 1)
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for r in range(3):
        t.cell(r, 0).text = f"行{r}"
        t.cell(r, 1).text = "值"
    doc.add_paragraph("独家承诺文本：现场服务响应时间两小时。")
    src = mlib.mt_files_dir() / "补充素材.docx"
    doc.save(src)
    f = db.mt_insert_file("补充素材.docx", "hash_docx_after_table")
    mlib.run_parse(f["id"])
    md_lines = mlib.mt_parse_paths("补充素材.docx")[0].read_text(encoding="utf-8").splitlines()
    ln = next(i for i, line in enumerate(md_lines, 1) if "独家承诺文本" in line)
    block = mlib.create_block(f["id"], "表格后承诺段", "勾选表格后段落", ranges=[[ln, ln]])

    section = _make_section()
    r = docx_material_inject.invoke({"block_id": block["id"], "dest": section})
    assert r.startswith("[已注入]"), r
    chk = Document(str(_abs(env, section)))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "独家承诺文本：现场服务响应时间两小时。" in texts
    assert len(chk.tables) == 0  # 只注入勾选段，表格不带


# ---------- 素材注入：lines 区间自选（2026-09-15 契约修正批） ----------


def test_material_inject_lines_subset(env):
    """lines=块内行号区间：只注入该区间元素（块=拷贝授权范围非注入原子——
    大块只取本节相关章节）；区间外标题/表格/图片不带，成功摘要带区间。"""
    md_lines = mlib.mt_parse_paths("历史运维方案.docx")[0].read_text(encoding="utf-8").splitlines()
    para_ln = next(i for i, ln in enumerate(md_lines, 1) if "北京华信" in ln)
    section = _make_section()
    r = docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": section, "lines": f"L{para_ln}-L{para_ln}"}
    )
    assert r.startswith("[已注入]"), r
    assert f"（L{para_ln}-L{para_ln}）" in r  # 摘要带区间 scope
    chk = Document(str(_abs(env, section)))
    texts = "\n".join(p.text for p in chk.paragraphs)
    assert "北京华信" in texts                  # 区间内段落注入
    assert "运维服务方案" not in texts          # 区间外标题不带
    assert len(chk.tables) == 0                 # 区间外表格不带
    assert not chk.element.body.findall(".//" + qn("a:blip"))  # 区间外图片不带


def test_material_inject_lines_fence(env):
    """授权围栏：lines 请求区间须完全落在块勾选范围内——块是用户授权的拷贝
    范围，块外内容机械拒绝（报错带块的合法区间）；格式错给人话。"""
    md_lines = mlib.mt_parse_paths("历史运维方案.docx")[0].read_text(encoding="utf-8").splitlines()
    n = len(md_lines)
    para_ln = next(i for i, ln in enumerate(md_lines, 1) if "北京华信" in ln)
    narrow = mlib.create_block(env["file"]["id"], "窄块", "", ranges=[[para_ln, para_ln + 1]])
    section = _make_section()
    # 完全在勾选范围外
    r = docx_material_inject.invoke(
        {"block_id": narrow["id"], "dest": section, "lines": f"L{n}-L{n}"}
    )
    assert r.startswith("[注入失败]") and "超出素材块《窄块》" in r
    assert f"L{para_ln}-L{para_ln + 1}" in r  # 合法区间随报错给出
    # 跨出勾选边界（部分越界同样拒）
    r = docx_material_inject.invoke(
        {"block_id": narrow["id"], "dest": section, "lines": f"L{para_ln - 1}-L{para_ln}"}
    )
    assert r.startswith("[注入失败]") and "超出素材块" in r
    # 格式错误
    r = docx_material_inject.invoke({"block_id": narrow["id"], "dest": section, "lines": "第三段"})
    assert r.startswith("[注入失败]") and "lines 须为" in r


def test_material_inject_lines_multi_range_and_dup(env):
    """多区间一次注入（段落+表格）；同区间二次注入被拦（防重基数=本次选中
    元素、摘要带区间），另一区间不受牵连——同块不同区间进不同内容的正道。"""
    md_lines = mlib.mt_parse_paths("历史运维方案.docx")[0].read_text(encoding="utf-8").splitlines()
    para_ln = next(i for i, ln in enumerate(md_lines, 1) if "北京华信" in ln)
    tbl_lns = [i for i, ln in enumerate(md_lines, 1) if "|" in ln]
    section = _make_section()
    # 多区间（中文逗号分隔也认）：正文段 + 表格
    r = docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": section,
         "lines": f"L{para_ln}-L{para_ln}，L{tbl_lns[0]}-L{tbl_lns[-1]}"}
    )
    assert r.startswith("[已注入]") and "表格 1 张" in r, r
    chk = Document(str(_abs(env, section)))
    assert "北京华信" in "\n".join(p.text for p in chk.paragraphs)
    assert len(chk.tables) == 1
    # 同区间二次注入 → 拦（追加语义，重复=翻倍）
    r2 = docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": section, "lines": f"L{para_ln}-L{para_ln}"}
    )
    assert r2.startswith("[注入失败]") and "已注入过" in r2
    assert f"（L{para_ln}-L{para_ln}）" in r2


def test_source_inject_rejects(env):
    from app.artifact_store import sources_dir

    sroot = sources_dir(env["task"]["id"])
    sroot.mkdir(parents=True, exist_ok=True)
    (sroot / "格式附件.txt").write_text("投标函格式", encoding="utf-8")
    section = _make_section()
    r = docx_source_inject.invoke({"source": "格式附件.txt", "dest": section})
    assert r.startswith("[注入失败]") and "不是 docx" in r
    assert "没有文件" in docx_source_inject.invoke({"source": "不存在的附件.docx", "dest": section})
    _tender_source(env)
    assert "起-止" in docx_source_inject.invoke(
        {"source": "招标文件.docx", "dest": section, "lines": "全部"}
    )
    assert "文件不存在" in docx_source_inject.invoke(
        {"source": "招标文件.docx", "dest": "body/未建节.docx"}
    )


# ---------- 修订标记 ----------

def _injected_section(env) -> tuple[str, Document]:
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[已注入]"), r
    return section, Document(str(_abs(env, section)))


def _abs(env, rel: str) -> Path:
    from app.artifact_store import work_dir

    return work_dir(env["task"]["id"]) / rel


def test_revise_replace_across_runs(env):
    """公司名拆在多个 run：替换落成 w:del/w:ins，拒绝修订=原文。"""
    section, _ = _injected_section(env)
    view = docx_section_read.invoke({"path": section})
    para_no = next(
        int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln
    )
    edits = [
        {"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"},
        # 同段第二处：穿过本轮已落的修订标记定位（换公司名 + 换工期是常态场景）
        {"para": para_no, "action": "replace", "find": "服务期三年", "text": "服务期两年"},
    ]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and "2 处" in r, r

    chk = Document(str(_abs(env, section)))
    para = chk.paragraphs[para_no - 1]
    del_text = "".join(
        t.text or "" for t in para._p.iter(qn("w:delText"))
    )
    assert COMPANY in del_text and "服务期三年" in del_text
    ins_text = "".join(
        t.text or "" for e in para._p.iter(qn("w:ins")) for t in e.iter(qn("w:t"))
    )
    assert "上海中信科技有限公司" in ins_text and "服务期两年" in ins_text
    # 段内其余文字未被标记动过（接受视角：前缀+两处新值+后缀齐整）
    from app.tools.docx_ops import _accepted_text

    assert _accepted_text(para._p) == (
        "本项目由上海中信科技有限公司承建，服务期两年。"
    )


def test_revise_insert_and_delete(env):
    section, _ = _injected_section(env)
    n_before = len(Document(str(_abs(env, section))).paragraphs)
    view = docx_section_read.invoke({"path": section})
    head_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if "运维服务方案" in ln)
    tail_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if "服务期三年" in ln)
    edits = [
        {"para": tail_no, "action": "insert_after", "text": "本章按本次招标文件要求编制。"},
        {"para": head_no, "action": "delete"},
    ]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and "2 处" in r, r

    chk = Document(str(_abs(env, section)))
    assert len(chk.paragraphs) == n_before + 1  # 删的是标记、插的是真段落
    # 插入段：段落标记与内容均为 ins（拒绝修订=整段消失）
    ins_para = chk.paragraphs[tail_no]  # 插在 tail_no 段后
    ppr = ins_para._p.find(qn("w:pPr"))
    assert ppr is not None and ppr.find(qn("w:rPr")).find(qn("w:ins")) is not None
    # 视图按「接受修订后」文本展示——新段内容对模型可见（python-docx .text 不含 ins）
    view = docx_section_read.invoke({"path": section})
    assert "本章按本次招标文件要求编制" in view
    assert "〔删除修订" in view  # 被删段在视图中明确标出
    # 删除段：run 文本全部转 delText
    gone = chk.paragraphs[head_no - 1]
    assert not gone.text or all(
        t.tag == qn("w:delText") for t in gone._p.iter() if t.tag in (qn("w:t"), qn("w:delText"))
    )


def test_revise_return_lists_edits_and_appends_view(env):
    """返回逐条结果；插段批附最新视图、新段最终序号与落盘事实一致（回读收敛批
    2026-09-13：此前返回只报计数，写手改完只能整读视图重锚定/确认，3.5 次/节）。"""
    section, _ = _injected_section(env)
    view = docx_section_read.invoke({"path": section})
    para_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln)
    tail_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if "服务期三年" in ln)
    edits = [
        {"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"},
        {"para": tail_no, "action": "delete"},
        {"para": tail_no, "action": "insert_after", "text": "本章按本次招标文件要求编制。"},
    ]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and "3 处" in r, r
    assert f"P{para_no} 替换「" in r and "上海中信科技" in r
    assert f"P{tail_no} 删除（原位保留删除标记" in r
    assert "新段（插在原 P" in r
    assert "最新读视图如下" in r  # 含插段 → 附视图（序号已位移）
    # 返回里的新段序号与落盘事实一致：该序号段落确为新插内容
    import re

    m = re.search(r"P(\d+) 新段", r)
    chk = Document(str(_abs(env, section)))
    assert "本章按本次招标文件要求编制" in _accepted_text(
        chk.paragraphs[int(m.group(1)) - 1]._p
    )
    # 附带视图里的序号即新序号：新插内容出现在视图块中
    assert "本章按本次招标文件要求编制" in r.split("最新读视图如下", 1)[1]


def test_revise_return_pure_replace_carries_no_view(env):
    """纯替换批不附视图（删段/替换序号不变），返回带免回读注记。"""
    section, _ = _injected_section(env)
    view = docx_section_read.invoke({"path": section})
    para_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln)
    edits = [{"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"}]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and f"P{para_no} 替换「" in r
    assert "最新读视图如下" not in r
    assert "不改变段落序号" in r


def test_revise_find_mismatch_keeps_file(env):
    section, _ = _injected_section(env)
    before = _abs(env, section).read_bytes()
    edits = [{"para": 2, "action": "replace", "find": "根本不存在的文本", "text": "x"}]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[修订失败]") and "未找到期望文本" in r
    assert _abs(env, section).read_bytes() == before  # 失败不落盘


def test_revise_bad_payload(env):
    section = _make_section()
    assert "不是合法 JSON" in docx_section_revise.invoke({"path": section, "edits": "{oops"})
    assert "非空 JSON 数组" in docx_section_revise.invoke({"path": section, "edits": "[]"})
    assert "replace 需 find 与 text" in docx_section_revise.invoke(
        {"path": section, "edits": json.dumps([{"para": 1, "action": "replace", "text": "x"}])}
    )
    assert "超出范围" in docx_section_revise.invoke(
        {"path": section, "edits": json.dumps([{"para": 99, "action": "delete"}])}
    )
    assert "不可同时传" in docx_section_revise.invoke(
        {"path": section, "edits": json.dumps([
            {"para": 1, "table": 1, "row": 1, "col": 1, "action": "fill", "text": "x"},
        ])}
    )
    assert "三件套" in docx_section_revise.invoke(
        {"path": section, "edits": json.dumps([{"action": "fill", "text": "x"}])}
    )


# ---------- 表格单元格修订 ----------


def _format_section(env) -> str:
    """拷好招标格式件的节（格式件填空的测试载体）。"""
    _tender_source(env)
    section = _make_section()
    assert docx_source_inject.invoke(
        {"source": "招标文件.docx", "dest": section}
    ).startswith("[已注入]")
    return section


def test_view_table_cell_coordinates(env):
    """表格视图逐行展开格坐标：空格显式（空）、合并格标注。"""
    section = _format_section(env)
    view = docx_section_read.invoke({"path": section})
    assert "[T1] 表格 2行×3列" in view
    assert "[T1] R1：C1=项目名称 C2=（空） C3=报价（元）" in view
    assert "[T1] R2：C1=（空） C2=报价大写 C3=（空）" in view


def test_view_merge_annotations(env):
    """合并格视图标注：横向合并标「同左」、纵向合并标「同上」。"""
    section, _ = _injected_section(env)
    view = docx_section_read.invoke({"path": section})
    assert "[T1] R1：C1=人员角色 C2=配置要求 C3=（同左合并格）" in view
    assert "[T1] R2：C1=项目经理" in view
    assert "[T1] R3：C1=（同上合并格）" in view


def test_revise_cell_fill_and_replace(env):
    """格式件填空：空格 fill、旧值 replace，段落与格混在同一次提交互不影响。"""
    from app.tools.docx_ops import _accepted_text

    section = _format_section(env)
    edits = [
        {"table": 1, "row": 1, "col": 2, "action": "fill", "text": "智慧园区运维项目"},
        {"table": 1, "row": 2, "col": 3, "action": "fill", "text": "【待补：报价】"},
        {"table": 1, "row": 2, "col": 2, "action": "replace", "find": "报价大写", "text": "报价大写金额"},
        {"para": 2, "action": "insert_after", "text": "函件正文补充行。"},
    ]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and "4 处" in r, r

    chk = Document(str(_abs(env, section)))
    c12 = chk.tables[0].cell(0, 1)
    assert _accepted_text(c12.paragraphs[0]._p) == "智慧园区运维项目"
    ins_text = "".join(
        t.text or "" for e in c12._tc.iter(qn("w:ins")) for t in e.iter(qn("w:t"))
    )
    assert "智慧园区运维项目" in ins_text  # fill 落成插入修订（Word 审阅可见）
    c22 = chk.tables[0].cell(1, 1)
    assert _accepted_text(c22.paragraphs[0]._p) == "报价大写金额"
    view = docx_section_read.invoke({"path": section})  # 视图（接受视角）可见填入值
    assert "C2=智慧园区运维项目" in view
    from app.tools.docx_ops import section_text_lines

    assert any("智慧园区运维项目" in ln for ln in section_text_lines(chk))


def test_revise_merged_cell_via_continuation_coordinate(env):
    """合并格的续坐标（R3C1）与原点坐标指向同一格：任一坐标都能改到内容。"""
    from app.tools.docx_ops import _accepted_text

    section, _ = _injected_section(env)
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"table": 1, "row": 3, "col": 1, "action": "replace", "find": "项目经理", "text": "项目总监"},
    ])})
    assert r.startswith("[已修订]"), r
    chk = Document(str(_abs(env, section)))
    assert _accepted_text(chk.tables[0].cell(1, 0).paragraphs[0]._p) == "项目总监"


def test_revise_cell_fill_nonempty_rejected(env):
    section = _format_section(env)
    before = _abs(env, section).read_bytes()
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"table": 1, "row": 1, "col": 1, "action": "fill", "text": "x"},
    ])})
    assert r.startswith("[修订失败]") and "非空格子" in r
    assert _abs(env, section).read_bytes() == before  # 失败不落盘


def test_revise_cell_out_of_range_and_find_miss(env):
    section = _format_section(env)
    assert "超出范围" in docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"table": 1, "row": 9, "col": 1, "action": "fill", "text": "x"},
    ])})
    assert "未找到期望文本" in docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"table": 1, "row": 1, "col": 1, "action": "replace", "find": "根本不存在", "text": "x"},
    ])})


def test_flatten_rejected_covers_table_cells(env):
    """自校验视角覆盖格内段落——表格修订写坏（拒绝视角文本变化）可被拦截。"""
    from app.tools.docx_ops import _flatten_rejected

    section = _format_section(env)
    flattened = _flatten_rejected(Document(str(_abs(env, section))))
    assert "项目名称" in flattened and "报价大写" in flattened


# ---------- replace 重写与恢复点 ----------


def test_create_replace_rotates_restore_points(env):
    section = _make_section()
    r = docx_section_create.invoke(
        {"path": section, "title": "3.1 运维方案", "paragraphs": "初稿", "replace": True}
    )
    assert r.startswith("[已重建]"), r
    d = _abs(env, section).with_name(_abs(env, section).name + ".restorepoints")
    assert len([p for p in d.iterdir() if p.suffix == ".bak"]) == 1
    # 连续重建 → 恢复点只留 3 个（栈深与工作台编辑/产物一致）
    for i in range(4):
        docx_section_create.invoke(
            {"path": section, "title": "x", "paragraphs": f"v{i}", "replace": True}
        )
    assert len([p for p in d.iterdir() if p.suffix == ".bak"]) == 3


# ---------- 样式与编号定义迁移 ----------


def test_material_inject_migrates_style_and_numbering(env):
    """注入元素引用自定义样式（basedOn 链）与多级编号：定义随迁，Word 打开不丢观感。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    src = mlib.mt_files_dir() / "编号样式素材.docx"
    doc = Document()
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="BideCustom">'
        f'<w:name w:val="BideCustom"/><w:basedOn w:val="BideCustomBase"/></w:style>'
    ))
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="BideCustomBase">'
        f'<w:name w:val="BideCustomBase"/></w:style>'
    ))
    np = doc.part.numbering_part.element
    np.append(parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77">'
        f'<w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        f'<w:lvlText w:val="%1"/></w:lvl></w:abstractNum>'
    ))
    np.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="77"><w:abstractNumId w:val="77"/></w:num>'))
    p = doc.add_paragraph("自定义样式编号段落")
    p.style = doc.styles["BideCustom"]
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="77"/></w:numPr>'
    ))
    doc.save(src)

    from app import db

    f = db.mt_insert_file("编号样式素材.docx", "hash_style_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(f["id"], "自定义样式章", "", ranges=[[1, 1000]])
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": block["id"], "dest": section})
    assert r.startswith("[已注入]") and "随迁样式定义" in r, r

    chk = Document(str(_abs(env, section)))
    ids = {s.get(qn("w:styleId")) for s in chk.styles.element.findall(qn("w:style"))}
    assert {"BideCustom", "BideCustomBase"} <= ids  # basedOn 链递归补齐
    nums = {n.get(qn("w:numId")) for n in chk.part.numbering_part.element.findall(qn("w:num"))}
    assert "77" in nums  # numId 连其 abstractNum 一并补拷


# ---------- 编号定义冲突重映射（numbering id 只是文档内部门牌号） ----------


def _add_cn_numbering(doc, num_id: str, abs_id: str) -> None:
    """素材侧自定义中文编号：先替换模板自带的同 id 定义（素材内 id 唯一），
    与节文件模板的同 id 形成跨文档冲突——模拟历史标书的独立编号空间。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    np = doc.part.numbering_part.element
    for el in list(np.findall(qn("w:num"))):
        if el.get(qn("w:numId")) == num_id:
            np.remove(el)
    for el in list(np.findall(qn("w:abstractNum"))):
        if el.get(qn("w:abstractNumId")) == abs_id:
            np.remove(el)
    np.append(parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">'
        f'<w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="chineseCountingThousand"/>'
        f'<w:lvlText w:val="一、"/><w:lvlJc w:val="left"/></w:lvl></w:abstractNum>'
    ))
    np.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abs_id}"/></w:num>'
    ))


def _numbered_para(doc, text: str, num_id: str) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    p = doc.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))


def _numbering_root(chk):
    return chk.part.numbering_part.element


def _numbered_para_ids(chk) -> list[str]:
    out = []
    for p in chk.paragraphs:
        ppr = p._p.find(qn("w:pPr"))
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        if numpr is not None and numpr.find(qn("w:numId")) is not None:
            out.append(numpr.find(qn("w:numId")).get(qn("w:val")))
    return out


def _abs_ref_of(chk, num_id: str) -> str:
    num = next(n for n in _numbering_root(chk).findall(qn("w:num"))
               if n.get(qn("w:numId")) == num_id)
    return num.find(qn("w:abstractNumId")).get(qn("w:val"))


def _lvl_text_of_abs(chk, abs_id: str) -> str | None:
    abs_el = next(a for a in _numbering_root(chk).findall(qn("w:abstractNum"))
                  if a.get(qn("w:abstractNumId")) == abs_id)
    lvl0 = abs_el.find(qn("w:lvl"))
    return lvl0.find(qn("w:lvlText")).get(qn("w:val")) if lvl0 is not None else None


def test_material_inject_numbering_conflict_remapped(env):
    """numId 撞建节模板（模板自带 numId 1-9）：定义不同则重映射新 id、改写引用，
    素材编号定义保真、模板原定义不动——修复「同 id 即沿用目标定义」的静默错配。"""
    src = mlib.mt_files_dir() / "编号冲突素材.docx"
    doc = Document()
    doc.add_heading("人员配置", 1)
    _add_cn_numbering(doc, "2", "2")  # numId 与 abstractNumId 双撞模板
    _numbered_para(doc, "项目经理一名", "2")
    doc.save(src)

    from app import db

    f = db.mt_insert_file("编号冲突素材.docx", "hash_numc_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(f["id"], "编号冲突章", "", ranges=[[1, 1000]])
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": block["id"], "dest": section})
    assert r.startswith("[已注入]") and "编号定义 1 组" in r, r

    chk = Document(str(_abs(env, section)))  # 重新打开不炸 = 迁入 XML 可解析
    used = _numbered_para_ids(chk)
    assert used == [used[0]] and used[0] != "2"  # 引用已重映射，不再指向模板 numId 2
    assert _lvl_text_of_abs(chk, _abs_ref_of(chk, used[0])) == "一、"  # 素材定义保真
    assert _lvl_text_of_abs(chk, _abs_ref_of(chk, "2")) != "一、"  # 模板原定义未被改动


def test_material_inject_numbering_sequence_semantics(env):
    """同一块内的多段共享同一重映射（列表编号连续）；另一块再注入=独立序列
    （跨调用不复用映射——不同素材的同定义列表不得被错误并接成连续编号）。"""
    src = mlib.mt_files_dir() / "编号序列素材.docx"
    doc = Document()
    doc.add_heading("服务承诺", 1)
    _add_cn_numbering(doc, "2", "2")
    _numbered_para(doc, "七乘二十四小时响应", "2")
    _numbered_para(doc, "两小时内到场", "2")
    doc.add_heading("培训计划", 1)
    _numbered_para(doc, "管理员培训一期", "2")
    doc.save(src)

    from app import db

    f = db.mt_insert_file("编号序列素材.docx", "hash_nums_1")
    mlib.run_parse(f["id"])
    block1 = mlib.create_block(f["id"], "承诺章", "", ranges=[[1, 3]])
    block2 = mlib.create_block(f["id"], "培训章", "", ranges=[[4, 1000]])
    section = _make_section()
    assert docx_material_inject.invoke({"block_id": block1["id"], "dest": section}).startswith("[已注入]")
    chk = Document(str(_abs(env, section)))
    first_two = _numbered_para_ids(chk)
    assert len(first_two) == 2 and first_two[0] == first_two[1]  # 块内两段同一编号方案

    assert docx_material_inject.invoke({"block_id": block2["id"], "dest": section}).startswith("[已注入]")
    chk = Document(str(_abs(env, section)))
    ids = _numbered_para_ids(chk)
    assert len(ids) == 3
    assert ids[2] != ids[0]  # 跨块独立序列（不被并接）
    assert _lvl_text_of_abs(chk, _abs_ref_of(chk, ids[2])) == "一、"  # 定义仍保真


def test_material_inject_numbering_abs_conflict_only(env):
    """numId 不撞但 abstractNumId 撞模板：定义不同则 abstractNum 换新 id 迁入，
    num 保留原 id 改指向——此前这条轴同样静默错配。"""
    src = mlib.mt_files_dir() / "编号半撞素材.docx"
    doc = Document()
    doc.add_heading("售后条款", 1)
    _add_cn_numbering(doc, "77", "3")  # numId 不撞；abstractNumId=3 撞模板且定义不同
    _numbered_para(doc, "质保期三年", "77")
    doc.save(src)

    from app import db

    f = db.mt_insert_file("编号半撞素材.docx", "hash_numh_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(f["id"], "编号半撞章", "", ranges=[[1, 1000]])
    section = _make_section()
    assert docx_material_inject.invoke({"block_id": block["id"], "dest": section}).startswith("[已注入]")

    chk = Document(str(_abs(env, section)))
    used = _numbered_para_ids(chk)
    assert used == ["77"]  # numId 不撞 → 原 id 保留
    new_abs = _abs_ref_of(chk, "77")
    assert new_abs != "3"  # abstractNum 已换新 id（不指向模板的 3）
    assert _lvl_text_of_abs(chk, new_abs) == "一、"  # 素材定义保真
    assert _lvl_text_of_abs(chk, "3") != "一、"  # 模板 abstractNum 3 未被改动


# ---------- 2026-09-08 机制修复：分节符/字体/丢图/锚定漂移/重复注入 ----------


def test_material_inject_strips_inner_sectpr(env):
    """素材段内分节符（横向页等版式设置）不得随元素拷贝进入节文件——否则
    分节属性中途生效、版式突变（合册同引擎同防线）。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    src = mlib.mt_files_dir() / "分节素材.docx"
    doc = Document()
    doc.add_heading("带分节的章", 1)
    p = doc.add_paragraph("章末段落")
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:sectPr {nsdecls("w")}><w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/></w:sectPr>'
    ))
    doc.save(src)

    from app import db

    f = db.mt_insert_file("分节素材.docx", "hash_sect_1")
    mlib.run_parse(f["id"])
    block = mlib.create_block(f["id"], "分节章", "", ranges=[[1, 1000]])
    section = _make_section()
    assert docx_material_inject.invoke({"block_id": block["id"], "dest": section}).startswith("[已注入]")
    chk = Document(str(_abs(env, section)))
    assert len(chk.element.body.findall(".//" + qn("w:sectPr"))) == 1  # 仅节文件文末分节


def _east(doc, style: str):
    rpr = doc.styles[style].element.find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    return fonts.get(qn("w:eastAsia")) if fonts is not None else None


def test_base_template_layout_on_create_and_assemble(env):
    """建节/合册从标书基准模板起建（格式与内容分离，模板由
    scripts/make_base_template.py 维护）：中文 eastAsia 定死、正文段挂
    Tender Body（1.5 倍行距+首行缩进 2 字符）、标题黑体加粗黑色（默认英文
    模板的蓝色英文脸是「生成的 Word 难看」的根源）、封面两档居中、目录样式带
    点线前导、A4+右下角页码域。版式取值对照用户版式参考件实测档（2026-09-13）。"""
    from docx.enum.text import (
        WD_ALIGN_PARAGRAPH,
        WD_TAB_ALIGNMENT,
        WD_TAB_LEADER,
    )
    from docx.shared import Pt, RGBColor

    from app import publish
    from app.artifact_store import work_dir

    section = _make_section("项目理解与需求分析.docx")
    docx_section_create.invoke({  # 已存在 → 覆盖重写，带初始正文段
        "path": section, "title": "项目理解与需求分析",
        "paragraphs": "本项目团队对采购需求的理解如下。", "replace": True,
    })
    doc = Document(str(work_dir(env["task"]["id"]) / section))
    assert _east(doc, "Normal") == "宋体"
    assert _east(doc, "Heading 1") == "黑体"
    h1 = doc.styles["Heading 1"]
    assert h1.font.color.rgb == RGBColor(0, 0, 0) and h1.font.size == Pt(16)  # 三号
    assert doc.styles["Heading 2"].font.size == Pt(15)  # 小三
    assert doc.styles["Heading 3"].font.size == Pt(14)  # 四号
    assert doc.styles["Heading 4"].font.size == Pt(12)  # 小四
    # 标题中西文同族（参考件标题 ascii 也是黑体，不吃 Times 混排）
    assert _east(doc, "Heading 1") == "黑体"
    h1_fonts = h1.element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert h1_fonts.get(qn("w:ascii")) == "黑体"
    assert h1.paragraph_format.keep_together is True
    # 主题引用防线（ＭＳ 明朝回归守卫，2026-09-13）：*Theme 属性会压住显式字体名，
    # 而 python-docx 默认模板的 Heading/Title 天生带引用、theme 东亚字形为空串，
    # 落空即回退应用默认东亚字体（Mac Office=ＭＳ 明朝）——显式名必须无引用残留
    import re
    import zipfile

    from app.tools import docx_ops as _docx_ops

    for sid_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
        rf = doc.styles[sid_name].element.find(qn("w:rPr")).find(qn("w:rFonts"))
        for theme_attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            assert rf.get(qn(f"w:{theme_attr}")) is None, f"{sid_name} 残留 {theme_attr}"
    tpl_zip = zipfile.ZipFile(_docx_ops._BASE_TEMPLATE)
    theme_xml = tpl_zip.read("word/theme/theme1.xml").decode("utf-8")
    eas = re.findall(r'<a:(?:majorFont|minorFont)>.*?<a:ea typeface="([^"]+)"', theme_xml, re.S)
    assert len(eas) == 2 and all(eas), f"theme 东亚字形应非空（兜素材合并样式），实际 {eas}"
    body = doc.styles["Tender Body"]
    assert body.paragraph_format.line_spacing == 1.5
    ind = body.element.get_or_add_pPr().find(qn("w:ind"))
    assert ind.get(qn("w:firstLineChars")) == "200"
    cover = doc.styles["Tender Cover"]
    assert _east(doc, "Tender Cover") == "黑体"
    assert cover.font.size == Pt(22) and cover.font.bold is True  # 二号加粗
    assert cover.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.styles["Tender Cover Sub"]
    assert sub.font.size == Pt(15) and sub.font.bold is True  # 小三加粗
    assert sub.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    # 目录：「目录」二字三号黑体居中，条目小四宋体 + 右对齐点线前导
    assert doc.styles["TOC Heading"].font.size == Pt(16)
    assert doc.styles["TOC Heading"].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    tab = doc.styles["toc 1"].paragraph_format.tab_stops[0]
    assert tab.alignment == WD_TAB_ALIGNMENT.RIGHT and tab.leader == WD_TAB_LEADER.DOTS
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.paragraphs[1].style.name == "Tender Body"
    # 模板带的样式示例段在建节产物中整段剥离（body 只留版面）
    assert len(doc.paragraphs) == 2
    assert "样式示例" not in "\n".join(p.text for p in doc.paragraphs)
    pgsz = doc.element.body.find(qn("w:sectPr")).find(qn("w:pgSz"))
    assert pgsz.get(qn("w:w")) == "11906"  # A4
    footer = doc.sections[0].footer._element
    assert any("PAGE" in (t.text or "") for t in footer.iter(qn("w:instrText")))

    publish.publish_artifact(_KEY, _DIR_SINGLE, task_id=env["task"]["id"], conversation_id=env["conv"]["id"])
    assert docx_assemble_volume.invoke({}).startswith("[已合册]")
    vol = Document(str(work_dir(env["task"]["id"]) / "body" / "整本-技术部分.docx"))
    assert _east(vol, "Normal") == "宋体"
    # 模板页脚自带页码域，合册不再手拼（双重页码回归守卫）
    n_fields = sum(
        1 for t in vol.sections[0].footer._element.iter(qn("w:instrText"))
        if "PAGE" in (t.text or "")
    )
    assert n_fields == 1


def test_material_inject_image_only_para_survives(env):
    """纯图片段落在「连续第 3 空行」位不再被空行折叠吃出元素映射——占位行
    保命，图片可注入（2026-09-08 丢图实证的修复）。"""
    src = mlib.mt_files_dir() / "空窗图片素材.docx"
    png = src.parent / "_tiny.png"
    _tiny_png(png)
    doc = Document()
    doc.add_heading("图集", 1)
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_picture(str(png))  # 连续第 3 个空行位（原实现会折叠丢图）
    doc.add_heading("收尾", 1)
    doc.save(src)

    from app import db

    f = db.mt_insert_file("空窗图片素材.docx", "hash_imgwin_1")
    mlib.run_parse(f["id"])
    md = (mlib.mt_parse_dir("空窗图片素材.docx") / "空窗图片素材.docx.md").read_text(encoding="utf-8")
    assert "![](图片)" in md  # 占位行可见（勾选界面/预览同见）
    block = mlib.create_block(f["id"], "图集章", "", ranges=[[1, 1000]])
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": block["id"], "dest": section})
    assert r.startswith("[已注入]") and "图片 1 张" in r  # 图注入成功
    chk = Document(str(_abs(env, section)))
    assert chk.element.body.findall(".//" + qn("a:blip"))


def test_material_inject_reparse_drift_rejected(env):
    """补跑映射时 md 与盘上不一致（模拟解析代码升级重排）：拒绝注入点名重勾，
    不静默按错位区间拷错元素。"""
    md_path = mlib.mt_parse_dir("历史运维方案.docx") / "历史运维方案.docx.md"
    md_path.write_text("被篡改的旧 md，行号与重解析产物必然不同。\n" * 3, encoding="utf-8")
    (md_path.parent / "element_map.json").unlink()
    section = _make_section()
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[注入失败]") and "勾选区间会错位" in r
    assert "重新确认该文件的勾选" in r


def test_material_inject_duplicate_blocked(env):
    """同块二次注入同一节文件被拦：注入是追加语义，重复=内容翻倍，无正当场景。"""
    section = _make_section()
    assert docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section}).startswith("[已注入]")
    r = docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section})
    assert r.startswith("[注入失败]") and "已注入过" in r


# ---------- 整本合册 ----------

_KEY = "tender.directory/tender-response-docs@1"

_DIR_SINGLE = {
    "response_documents": [
        {
            "name": "技术部分",
            "scope": "",
            "directory": [
                {"目录名称": "技术方案", "level": 1, "children": [
                    {"目录名称": "项目理解与需求分析", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": ["REQ-01"]},
                    {"目录名称": "总体设计方案", "level": 2, "children": [],
                     "交付形态": "混合", "来源位置": ["SCORE-02"]},
                ]},
                {"目录名称": "附件：资质证书复印件", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": ["MAND-02"]},
            ],
        }
    ]
}


def _seed_dir_artifact(env, content):
    from app import publish

    publish.publish_artifact(
        _KEY, content, task_id=env["task"]["id"], conversation_id=env["conv"]["id"]
    )


def test_assemble_tree_order_headings_and_missing(env):
    """树序合册：容器标题层级随树深、标题按树序自动编号（缺省第X章+1.1）、
    节文件自带标题跳过（对账用裸名，不受编号影响）、缺失节点名、
    模板填充叶子未产出按附件对待（点名不占位）；重跑覆盖且整本不算孤儿。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析",
         "paragraphs": "项目理解正文第一段。"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    assert "合并 1 节" in r
    assert "缺失 1 节未并入：总体设计方案" in r
    assert "模板填充类未产出 1 节（按附件对待，不占整本位）：附件：资质证书复印件" in r

    # 合册成功即自动发布 tender.volume 产物（2026-09-13）：索引行 + 包内 docx 字节一致
    from app import artifact_store as _astore
    from app import db as _adb

    assert "已发布为整本标书成果" in r
    assert "（第 1 版）" in r  # 版本号钉结局：真发新版本 vs 未重发（2026-09-16）
    vrows = [r_ for r_ in _adb.list_artifact_index() if r_["kind"] == "tender.volume"]
    assert len(vrows) == 1
    assert vrows[0]["display_name"] == "技术部分"
    assert vrows[0]["content_seq"] == 1
    pkg_docx = _astore.package_files(vrows[0]["artifact_id"], vrows[0])
    assert len(pkg_docx) == 1 and pkg_docx[0].name == "整本-技术部分.docx"
    assert pkg_docx[0].read_bytes() == _abs(env, "body/整本-技术部分.docx").read_bytes()

    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    paras = chk.paragraphs
    assert paras[0].style.name == "Title" and paras[0].text == "技术部分"
    assert paras[1].style.name == "Heading 1" and paras[1].text == "第一章\u3000技术方案"
    assert paras[2].style.name == "Heading 2" and paras[2].text == "1.1 项目理解与需求分析"
    texts = [p.text for p in paras]
    assert texts.count("1.1 项目理解与需求分析") == 1  # 节文件自带标题段已跳过（编号后标题仍只一份）
    assert "项目理解正文第一段。" in texts
    assert not any("资质证书复印件" in t for t in texts)  # 未产出格式件不进整本
    # 页脚页码域（可打印闭环）
    assert "PAGE" in chk.sections[0].footer.paragraphs[0]._p.xml
    # 重复合册：整本自身不算孤儿（派生产物；缺失节点名的「未并入」仍在）；
    # docx 内容未变 → 不重复发布（seq 不动，无新卡）
    r2 = docx_assemble_volume.invoke({})
    assert "个节文件未并入" not in r2
    assert "未重复发布" in r2
    assert "（第 1 版）" in r2  # 未重发分支同样带当前版本号
    assert _adb.get_artifact_index(vrows[0]["artifact_id"])["content_seq"] == 1


_DIR_COVER = {
    "response_documents": [
        {
            "name": "技术部分",
            "scope": "",
            "directory": [
                {"目录名称": "封面", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "技术方案", "level": 1, "children": [
                    {"目录名称": "项目理解与需求分析", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": ["REQ-01"]},
                ]},
            ],
        }
    ]
}


def test_revise_insert_with_style(env):
    """insert_after 可带 style 字段挂指定版式（封面行）；样式不存在回落正文
    样式且返回注记——用户自定义版式缺封面样式名不失败。"""
    section = _make_section()
    edits = json.dumps([
        {"para": 1, "action": "insert_after", "text": "XX 项目投标文件",
         "style": "Tender Cover"},
        {"para": 1, "action": "insert_after", "text": "投标人：XX 有限公司",
         "style": "Tender Cover Sub"},
        {"para": 1, "action": "insert_after", "text": "未知样式行",
         "style": "不存在的样式"},
    ], ensure_ascii=False)
    r = docx_section_revise.invoke({"path": section, "edits": edits})
    assert r.startswith("[已修订]")
    assert "不存在的样式" in r and "正文样式" in r
    chk = Document(str(_abs(env, section)))
    # 插入段在 w:ins 修订包裹里，p.text 读不到——按接受视角取文本
    by_text = {_accepted_text(p._p): p.style.name for p in chk.paragraphs}
    assert by_text["XX 项目投标文件"] == "Tender Cover"
    assert by_text["投标人：XX 有限公司"] == "Tender Cover Sub"
    assert by_text["未知样式行"] == "Tender Body"  # 回落正文样式


def test_assemble_cover_node(env):
    """封面=树首一级叶子（名「封面」）：整本首页即封面页——无册名 Title、
    无「封面」标题行（节文件标题段照旧剥）、封面行挂 Tender Cover、开
    「首页不同」（封面页无页眉页脚）、第一章分页让封面独占一页。"""
    _seed_dir_artifact(env, _DIR_COVER)
    docx_section_create.invoke({"path": "body/封面", "title": "封面"})
    docx_section_revise.invoke({"path": "body/封面", "edits": json.dumps([
        {"para": 1, "action": "insert_after", "text": "XX 项目投标文件（技术部分）",
         "style": "Tender Cover"},
        {"para": 1, "action": "insert_after", "text": "投标人：XX 有限公司",
         "style": "Tender Cover Sub"},
    ], ensure_ascii=False)})
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析",
         "paragraphs": "项目理解正文第一段。"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]") and "缺失" not in r and "未产出" not in r

    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    paras = chk.paragraphs
    # 封面行是修订插入（run 在 w:ins 里），文本按接受视角取
    texts = [_accepted_text(p._p) for p in paras]
    assert paras[0].style.name == "Tender Cover" and texts[0] == "XX 项目投标文件（技术部分）"
    assert not any(p.style.name == "Title" for p in paras)  # 册名大标题被跳过
    assert "封面" not in texts  # 封面节点自身标题不发、节文件标题段被剥
    # 封面不占编号序：第一个内容章仍是「第一章」
    chapter = next(p for t, p in zip(texts, paras) if t == "第一章\u3000技术方案")
    assert chapter.paragraph_format.page_break_before is True  # 封面独占首页
    assert chk.sections[0].different_first_page_header_footer is True
    assert "PAGE" in chk.sections[0].footer.paragraphs[0]._p.xml  # 默认页脚仍在


def test_assemble_cover_missing_reported(env):
    """封面缺节文件的两种形态照旧点名（正文编写→缺失正文节；模板或附件填充→
    按附件对待），册名 Title 同样跳过（封面位由树决定，不因缺文件回退双标题）。"""
    _seed_dir_artifact(env, _DIR_COVER)
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析",
         "paragraphs": "项目理解正文第一段。"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "缺失 1 节未并入：封面" in r
    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert chk.paragraphs[0].style.name == "Heading 1"  # 无 Title，直接进第一章

    import copy

    non_prose = copy.deepcopy(_DIR_COVER)
    non_prose["response_documents"][0]["directory"][0]["交付形态"] = "模板或附件填充"
    _seed_dir_artifact(env, non_prose)
    r2 = docx_assemble_volume.invoke({})
    assert r2.startswith("[已合册]")
    assert "模板填充类未产出 1 节（按附件对待，不占整本位）：封面" in r2
    chk2 = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert chk2.paragraphs[0].style.name == "Heading 1"


def test_assemble_includes_produced_format_node(env):
    """模板填充叶子产出节文件即按树序并进整本（格式件是标书组成部分），
    与正文统一编号（拍板 2026-09-10：格式件章也占章序）。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析",
         "paragraphs": "项目理解正文第一段。"}
    )
    docx_section_create.invoke(
        {"path": "body/附件：资质证书复印件", "title": "附件：资质证书复印件",
         "paragraphs": "（复印件附后）"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]") and "合并 2 节" in r, r
    assert "模板填充类未产出" not in r
    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    texts = [p.text for p in chk.paragraphs]
    assert "（复印件附后）" in texts
    assert texts.count("第二章\u3000附件：资质证书复印件") == 1  # 树序标题一份（格式件章统一编号；节文件自带标题已跳过）
    assert "个节文件未并入" not in r  # 格式件文件已消费，不算孤儿


def test_assemble_accepts_revisions_and_keeps_images(env):
    """整本=交付态：节内修订标记并入时按「接受全部修订」压平——整本零
    w:ins/w:del（未在 Word 里接受修订前不再新旧并存，目录乱象主诉）、替换后
    文本在场/被删文本消失、图片关系完好；节文件层修订原样保留（审阅入口
    在节级，改内容回节级改再重合册）。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "运维方案", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": ["REQ-01"]},
            ]}
        ]
    })
    section = "body/运维方案.docx"
    assert docx_section_create.invoke({"path": section, "title": "运维方案"}).startswith("[已创建]")
    assert docx_material_inject.invoke({"block_id": env["block"]["id"], "dest": section}).startswith("[已注入]")
    view = docx_section_read.invoke({"path": section})
    para_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln)
    assert docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"},
    ])}).startswith("[已修订]")
    # 整段删除修订（_tracked_delete 形态：run 全删+段落标记删）——压平后整段消失
    assert docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": 1, "action": "insert_after", "text": "删除演示段"},
    ])}).startswith("[已修订]")
    view = docx_section_read.invoke({"path": section})
    del_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if "删除演示段" in ln)
    assert docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": del_no, "action": "delete"},
    ])}).startswith("[已修订]")

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]") and "含图片 1 张" in r, r
    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert chk.element.body.find(".//" + qn("w:ins")) is None  # 交付态：修订已压平
    assert chk.element.body.find(".//" + qn("w:del")) is None
    texts = [p.text for p in chk.paragraphs]
    assert any("上海中信科技有限公司" in t for t in texts)  # 替换后的新文本在场（压平后 p.text 直读）
    assert not any(COMPANY in t for t in texts)  # 被删的旧公司名不在交付稿
    assert "删除演示段" not in texts  # 整段删除修订的段不复存在
    for blip in chk.element.body.findall(".//" + qn("a:blip")):
        assert blip.get(qn("r:embed")) in chk.part.related_parts  # 图片关系完好
    # 节文件层修订保留（审阅入口不动——面板/Word 审阅在节级）
    sec = Document(str(_abs(env, section)))
    assert sec.element.body.find(".//" + qn("w:ins")) is not None
    assert sec.element.body.find(".//" + qn("w:del")) is not None


def test_accept_revisions_keeps_table_cell_block_child():
    """格内唯一段带「段落标记删除+内容删净」压平后留空段：OOXML 硬要求
    w:tc 至少一个块级子元素，空格会让整本被 Word 判损坏（病态素材防御，
    正常 Word 编辑产生不了此形态——标记删除语义=与后续段合并）。"""
    tbl = parse_xml(
        f'<w:tbl {nsdecls("w")}><w:tr><w:tc><w:tcPr/>'
        "<w:p><w:pPr><w:rPr><w:del/></w:rPr></w:pPr>"
        '<w:del><w:r><w:t>格内唯一段</w:t></w:r></w:del>'
        "</w:p></w:tc></w:tr></w:tbl>"
    )
    assert _accept_revisions_inplace(tbl)
    tc = tbl.find(".//" + qn("w:tc"))
    assert tc.find(qn("w:p")) is not None  # 空段保命，格不再只剩 tcPr
    # 多段格：删净段移除后仍有兄弟段，不额外补空段
    tbl2 = parse_xml(
        f'<w:tbl {nsdecls("w")}><w:tr><w:tc><w:tcPr/>'
        "<w:p><w:pPr><w:rPr><w:del/></w:rPr></w:pPr>"
        '<w:del><w:r><w:t>第一段</w:t></w:r></w:del></w:p>'
        "<w:p><w:r><w:t>存活段</w:t></w:r></w:p>"
        "</w:tc></w:tr></w:tbl>"
    )
    assert _accept_revisions_inplace(tbl2)
    paras = tbl2.find(".//" + qn("w:tc")).findall(qn("w:p"))
    assert len(paras) == 1 and _accepted_text(paras[0]) == "存活段"


def test_assemble_demotes_off_tree_headings(env):
    """树外标题摘出大纲：节内小标题/素材自带章标题并入时显式 outlineLvl=9
    （正文级）——整本导航窗格/自动目录只剩合册器按树发的章节骨架；Heading
    样式与文本不动（视觉零变化）；合册器自发的标题无 outlineLvl 覆盖
    （样式自带层级，骨架完整）。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "运维实施方案", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    assert docx_section_create.invoke({"path": "body/运维实施方案", "title": "运维实施方案"}).startswith("[已创建]")
    # 节内小标题（写作模型自加，无编号，Heading 样式进大纲——乱象主诉之二）
    assert docx_section_revise.invoke({"path": "body/运维实施方案", "edits": json.dumps([
        {"para": 1, "action": "insert_after", "text": "沟通与报告机制", "style": "Heading 2"},
    ], ensure_ascii=False)}).startswith("[已修订]")
    # 素材自带章标题（与树标题不匹配的 mismatched 形态，同样进拷入面）
    assert docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": "body/运维实施方案.docx"}
    ).startswith("[已注入]")

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    headings = []
    for p in chk.paragraphs:
        if not p.style.name.startswith("Heading"):
            continue
        ppr = p._p.find(qn("w:pPr"))
        lvl = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        headings.append((p.style.name, lvl.get(qn("w:val")) if lvl is not None else None, p.text))
    # 骨架=无 outlineLvl 覆盖的标题（合册器按树发，样式自带层级）
    assert [h for h in headings if h[1] is None] == [
        ("Heading 1", None, "第一章\u3000运维实施方案")
    ]
    # 树外标题全部显式降 9（导航/自动目录不再收录）；样式按树深度降级
    # （2026-09-16 拍板：本叶深度 1，源 H2 小标题目标级 2 不变；素材章标题
    # 源 H1 → 目标级 2——视觉层级有意下沉，替代 09-12「样式不动」旧口径）
    demoted = {h[2]: h for h in headings if h[1] == "9"}
    assert demoted["沟通与报告机制"][0] == "Heading 2"
    assert demoted["运维服务方案"][0] == "Heading 2"  # 素材章标题摘大纲且降级


def test_assemble_toc_reconciliation(env):
    """目录页对账（探测+提示，不是门禁）：目录页条目 vs 实收章节——列了整本
    没有的（按附件对待未产出）、漏了实有的，逐项点名；编号前缀/点线页码写法
    差异不误报；超长说明行不当条目；「目录」自身两侧豁免。"""
    note = "说明" + "长" * 65  # 67 字：目录页说明行，超 60 不当条目
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "目录", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "总体设计", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "资质证明", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": []},
            ]}
        ]
    })
    # 目录页：项目理解（编号+点线页码写法）、资质证明（实为按附件对待未产出）、
    # 说明行——漏列总体设计
    assert docx_section_create.invoke({"path": "body/目录", "title": "目录", "paragraphs":
        f"1.1 项目理解…………3\n资质证明\n{note}"}).startswith("[已创建]")
    assert docx_section_create.invoke(
        {"path": "body/项目理解", "title": "项目理解", "paragraphs": "正文"}).startswith("[已创建]")
    assert docx_section_create.invoke(
        {"path": "body/总体设计", "title": "总体设计", "paragraphs": "正文"}).startswith("[已创建]")

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    warn = next(ln for ln in r.splitlines() if "目录页与正文对账不符" in ln)
    assert "资质证明" in warn  # 列了但整本没有（按附件对待未产出，目录页不该照列）
    assert "总体设计" in warn  # 整本有但目录页未列
    assert "项目理解" not in warn  # 编号/点线页码写法差异不误报
    assert "说明" not in warn  # 超长说明行不当条目
    assert "模板填充类未产出 1 节" in r  # 既有行不受影响


# ---------- 机械目录页 / 前置区编号 / 空容器抑制 / 附件壳节（2026-09-14 结构缺口批） ----------

_DIR_TOC = {
    "response_documents": [
        {
            "name": "技术部分",
            "scope": "",
            "directory": [
                {"目录名称": "封面", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "编制索引", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "目录", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": []},
                {"目录名称": "技术方案", "level": 1, "children": [
                    {"目录名称": "项目理解与需求分析", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": ["REQ-01"]},
                ]},
                {"目录名称": "附件：资质证书复印件", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": ["MAND-02"]},
            ],
        }
    ]
}


def test_assemble_mechanical_toc_and_front_matter(env):
    """目录节点无节文件→机械目录页（Word 目录域+缓存清单，插在目录节点位置）；
    前置区（目录前的非封面节点=编制索引）不占章号，正文从目录后第一章起编；
    未产出附件节在缓存清单标「另附」；返回行报告条目数。"""
    _seed_dir_artifact(env, _DIR_TOC)
    docx_section_create.invoke({"path": "body/封面", "title": "封面"})
    docx_section_create.invoke({"path": "body/编制索引", "title": "编制索引",
                                "paragraphs": "索引表正文。"})
    docx_section_create.invoke({"path": "body/项目理解与需求分析",
                                "title": "项目理解与需求分析",
                                "paragraphs": "项目理解正文第一段。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    assert "目录页已生成（4 条" in r  # 编制索引 + 第一章 技术方案 + 1.1 + 附件另附
    assert "未产出附件节标「另附」1 项" in r
    assert "模板填充类未产出 1 节" in r

    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    texts = [_accepted_text(p._p) for p in chk.paragraphs]
    # 前置区不占章号：编制索引无「第一章」前缀；正文章从目录后起编
    assert not any(t.startswith("第一章\u3000编制索引") for t in texts)
    assert "第一章\u3000技术方案" in texts
    # 机械目录页：标题非 Heading 样式（防目录域/导航自引用），位于编制索引之后、第一章之前
    toc_i = next(i for i, p in enumerate(chk.paragraphs) if p.text == "目录")
    assert chk.paragraphs[toc_i].style.name not in ("Heading 1", "Heading 2", "Title")
    idx_i = next(i for i, t in enumerate(texts) if t == "编制索引")
    ch1_i = next(i for i, t in enumerate(texts) if t.startswith("第一章\u3000技术方案"))
    assert idx_i < toc_i < ch1_i
    # Word 目录域三件套 + 指令文本（域包裹缓存清单——更新域整体替换得带页码目录）
    xml = chk.element.body.xml
    assert xml.count('w:fldCharType="begin"') >= 1
    assert ' TOC \\o "1-3"' in xml
    assert 'w:fldCharType="separate"' in xml and 'w:fldCharType="end"' in xml
    # 缓存清单条目：前置页无编号、正文带编号、缺文件正文节点照列、附件标另附
    assert any(t.strip() == "编制索引" for t in texts)
    assert any(t.strip() == "第一章\u3000技术方案" for t in texts)
    assert any(t.strip() == "1.1 项目理解与需求分析" for t in texts)
    assert any("附件：资质证书复印件" in t and "（另附）" in t for t in texts)


def test_assemble_toc_node_with_file_unnumbered(env):
    """手写目录节文件优先：并入+对账照旧，但「目录」标题不占章号（此前会是
    「第一章 目录」），后续正文从第一章起编。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "目录", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    assert docx_section_create.invoke({"path": "body/目录", "title": "目录",
                                       "paragraphs": "项目理解"}).startswith("[已创建]")
    assert docx_section_create.invoke({"path": "body/项目理解", "title": "项目理解",
                                       "paragraphs": "正文。"}).startswith("[已创建]")
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "目录页与正文对账不符" not in r
    assert "目录页已生成" not in r  # 手写路径不走机械生成
    texts = [_accepted_text(p._p) for p in Document(
        str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert not any(t.startswith("第一章\u3000目录") for t in texts)
    assert "第一章\u3000项目理解" in texts


def test_assemble_attachment_shell_section(env):
    """附件壳节（2026-09-14 拍板：知识库全无命中也建节）=模板填充叶子带节文件
    →按树序正常并入整本（贴入位行随正文走），不再进「按附件对待」未产出点名。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    docx_section_create.invoke({"path": "body/项目理解与需求分析",
                                "title": "项目理解与需求分析", "paragraphs": "正文。"})
    docx_section_create.invoke({"path": "body/附件：资质证书复印件",
                                "title": "附件：资质证书复印件",
                                "paragraphs": "（此处贴入：资质证书复印件，加盖公章）"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "模板填充类未产出" not in r
    texts = [_accepted_text(p._p) for p in Document(
        str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert any("此处贴入：资质证书复印件" in t for t in texts)
    assert any(t.startswith("第二章\u3000附件：资质证书复印件") for t in texts)


def test_assemble_empty_container_suppressed(env):
    """空容器抑制（兜底）：容器子树无任何将产出叶子→不发章标题（防「第X章
    其他资料」光杆空壳）；有产出叶子的容器照常发标题。子节点照旧走未产出
    点名，信息不丢。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "技术方案", "level": 1, "children": [
                    {"目录名称": "项目理解与需求分析", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": ["REQ-01"]},
                ]},
                {"目录名称": "其他资料", "level": 1, "children": [
                    {"目录名称": "社保缴纳票据", "level": 2, "children": [],
                     "交付形态": "模板或附件填充", "来源位置": []},
                    {"目录名称": "信用查询资料", "level": 2, "children": [],
                     "交付形态": "模板或附件填充", "来源位置": []},
                ]},
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/项目理解与需求分析",
                                "title": "项目理解与需求分析", "paragraphs": "正文。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "模板填充类未产出 2 节" in r
    texts = [_accepted_text(p._p) for p in Document(
        str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert not any("其他资料" in t for t in texts)  # 空容器标题被抑制
    assert any(t.startswith("第一章\u3000技术方案") for t in texts)  # 部分产出容器照常


def test_assemble_suppressed_container_reported(env):
    """空容器抑制的容器标题点名：此前容器名从整本/目录页/报告三处同时消失，
    其「另附」叶子悬挂在无父标题下——补一行报告交代。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "附件章", "level": 1, "children": [
                    {"目录名称": "附件：证书", "level": 2, "children": [],
                     "交付形态": "模板或附件填充", "来源位置": []},
                ]},
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/项目理解", "title": "项目理解",
                                "paragraphs": "正文。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "容器章节整体未产出" in r
    suppressed_line = next(ln for ln in r.splitlines() if "容器章节整体未产出" in ln)
    assert "附件章" in suppressed_line


def test_assemble_no_toc_node_advisory(env):
    """无「目录」节点：全部一级节点按章编号（旧行为兼容）但现在可见——
    一行事实提示，供用户判断是否要前置区。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    r = docx_assemble_volume.invoke({})
    assert "目录树无「目录」节点" in r and "按章编号" in r


def test_assemble_cover_variant_advisory(env):
    """「封面页」等变体不识别为封面（清洗后≠「封面」）：占号+发标题照旧，
    补事实提示指路改名，不做模糊匹配。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "封面页", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/封面页", "title": "封面页",
                                "paragraphs": "封面内容。"})
    docx_section_create.invoke({"path": "body/项目理解", "title": "项目理解",
                                "paragraphs": "正文。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "含「封面」但非约定名「封面」" in r


def test_assemble_duplicate_leaf_titles_merge_once(env):
    """同册同名叶子（清洗后同文件）：内容只并入一次——重复并入=整本里同一份
    内容出现两遍；⚠️ 点名请改目录标题其一。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "项目理解与需求分析", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解与需求分析 ", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},  # 尾随空格清洗后同名
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/项目理解与需求分析",
                                "title": "项目理解与需求分析",
                                "paragraphs": "唯一一份正文。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "清洗后同名的节点" in r and "内容只并入一次" in r
    assert "合并 1 节" in r
    texts = [_accepted_text(p._p) for p in Document(
        str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert texts.count("唯一一份正文。") == 1


def test_assemble_duplicate_volume_names_skips_second(env):
    """册名清洗后重名：第二册整册跳过（避免节目录与整本文件互覆），⚠️ 点名。"""
    content = {
        "response_documents": [
            {"name": "商务部分", "scope": "", "directory": [
                {"目录名称": "售后承诺", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []}]},
            {"name": "商务部分 ", "scope": "", "directory": [  # 尾随空格清洗后同名
                {"目录名称": "售后承诺", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []}]},
        ]
    }
    _seed_dir_artifact(env, content)
    docx_section_create.invoke({"path": "body/商务部分/售后承诺", "title": "售后承诺",
                                "paragraphs": "甲册内容。"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "清洗后同名" in r and "本册未合册" in r
    assert "合并 1 节" in r  # 只有第一册合册


def test_assemble_self_volume_named_leaf_guarded(env):
    """叶子标题清洗后与整本输出同名：上一轮整本产物不当节内容并入（自噬翻倍
    断路），按缺失处理并单独点名。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "整本-技术部分", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/项目理解", "title": "项目理解",
                                "paragraphs": "正文。"})
    r1 = docx_assemble_volume.invoke({})
    assert "缺失 1 节未并入" in r1  # 首轮：叶子「整本-技术部分」尚无文件
    r2 = docx_assemble_volume.invoke({})
    assert r2.startswith("[已合册]")
    assert "个节点标题与整本产物文件同名" in r2
    texts = [_accepted_text(p._p) for p in Document(
        str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert texts.count("正文。") == 1  # 上一轮整本没有整份被并进来


def test_assemble_volume_prefix_accounting(env):
    """「整本-」前缀收窄（2026-09-15 审计）：标题恰以「整本-」开头的合法节文件
    正常并入；子目录里错放的 整本-*.docx 按孤儿点名；根层整本产物不报孤儿。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "整本-特别篇", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
                {"目录名称": "项目理解与需求分析", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    docx_section_create.invoke({"path": "body/整本-特别篇", "title": "整本-特别篇",
                                "paragraphs": "特别篇正文。"})
    docx_section_create.invoke({"path": "body/项目理解与需求分析",
                                "title": "项目理解与需求分析", "paragraphs": "普通节正文。"})
    stray = _abs(env, "body/子目录/整本-错放.docx")
    stray.parent.mkdir(parents=True, exist_ok=True)
    Document().save(stray)
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "合并 2 节" in r  # 「整本-特别篇」合法节文件照常并入
    orphan_seg = r.split("个节文件未并入", 1)[1]
    assert "body/子目录/整本-错放.docx" in orphan_seg
    assert "整本-特别篇" not in orphan_seg  # 已并入的节文件不算孤儿
    r2 = docx_assemble_volume.invoke({})
    orphan_seg2 = r2.split("个节文件未并入", 1)[1]
    assert "body/整本-技术部分.docx" not in orphan_seg2  # 根层整本产物不报孤儿


def test_assemble_merged_zero_unfilled_reported(env):
    """全册只有未产出的模板填充叶子：merged==0 报告列出 unfilled 名单，
    不再误报「目录树无叶子节点」；不产出整本文件。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "附件：资质证书复印件", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": []},
            ]}
        ]
    })
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "无已写节文件" in r
    assert "模板填充类未产出 1 节" in r and "附件：资质证书复印件" in r
    assert "目录树无叶子节点" not in r
    assert not _abs(env, "body/整本-技术部分.docx").exists()


def test_assemble_migrates_comments_and_warns_inline(env):
    """节内待办批注迁入整本（id 重映射、标记不悬空）；内联占位兜底扫描点名。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析",
         "paragraphs": "项目理解正文第一段。"}
    )
    assert docx_comment_add.invoke({
        "path": "body/项目理解与需求分析.docx", "after": "2",
        "text": "缺项目批复文件编号，待用户确认",
    }).startswith("[已加批注]")
    docx_section_create.invoke(
        {"path": "body/总体设计方案", "title": "总体设计方案",
         "paragraphs": "方案正文。\n【待补：分项报价表金额】"}
    )
    assert docx_comment_add.invoke({
        "path": "body/总体设计方案.docx",
        "text": "分项报价表金额缺口径，待澄清",
    }).startswith("[已加批注]")

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]") and "含批注 2 条待处理" in r, r
    assert "正文含 1 处内联占位" in r
    assert "总体设计方案(P3)：【待补：分项报价表金额】" in r

    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    cmts = list(chk.comments)
    assert {c.text for c in cmts} == {"缺项目批复文件编号，待用户确认", "分项报价表金额缺口径，待澄清"}
    # 正文标记 id 已重映射、各自指向存在的批注（不悬空、不串号）
    starts = chk.element.body.findall(".//" + qn("w:commentRangeStart"))
    assert {s.get(qn("w:id")) for s in starts} == {str(c.comment_id) for c in cmts}


def test_assemble_multi_volume_and_orphans(env):
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "项目理解", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]},
            {"name": "商务部分", "scope": "", "directory": [
                {"目录名称": "售后服务承诺", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]},
        ]
    })
    docx_section_create.invoke({"path": "body/技术部分/项目理解", "title": "项目理解", "paragraphs": "正文"})
    docx_section_create.invoke({"path": "body/旧版遗留节", "title": "旧版遗留节", "paragraphs": "旧稿"})
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert _abs(env, "body/整本-技术部分.docx").is_file()
    assert "[合册失败]" not in r
    assert "商务部分：无已写节文件" in r  # 缺失整册也只报事实，不炸
    assert "未并入" in r and "body/旧版遗留节.docx" in r  # 孤儿节文件点名


def test_assemble_without_directory_errors(env):
    assert "[合册失败] 无投标目录产物" in docx_assemble_volume.invoke({})


_DIR_NUMBERING = {
    "response_documents": [
        {
            "name": "技术部分", "scope": "",
            "directory": [
                {"目录名称": "总体部署", "level": 1, "children": [
                    {"目录名称": "部署原则", "level": 2, "children": [
                        {"目录名称": "安全原则", "level": 3, "children": [],
                         "交付形态": "正文编写", "来源位置": []},
                    ]},
                    {"目录名称": "分步实施", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": []},
                ]},
                {"目录名称": "售后服务", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ],
        },
        {
            "name": "商务部分", "scope": "",
            "directory": [
                {"目录名称": "报价说明", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ],
        },
    ]
}


def test_assemble_numbering_schemes(env):
    """章节编号=树位置的纯函数：缺省 chapter（第X章全角空格+1.1）、decimal、
    gov、none 四格式按树序生成；多册各自从首章重起；编号后 _is_title_para
    对账剥标题仍按裸名（自带标题段只剥一份）。"""
    import copy

    for name in ("安全原则", "分步实施", "售后服务"):
        docx_section_create.invoke(
            {"path": f"body/技术部分/{name}", "title": name, "paragraphs": "正文"}
        )
    docx_section_create.invoke(
        {"path": "body/商务部分/报价说明", "title": "报价说明", "paragraphs": "正文"}
    )

    def headings() -> dict[str, list[str]]:
        out: dict[str, list[str]] = {}
        for vol in ("技术部分", "商务部分"):
            doc = Document(str(_abs(env, f"body/整本-{vol}.docx")))
            out[vol] = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        return out

    # 缺省（产物无 numbering 字段）=chapter
    _seed_dir_artifact(env, _DIR_NUMBERING)
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]") and "⚠️" not in r, r
    assert headings()["技术部分"] == [
        "第一章\u3000总体部署", "1.1 部署原则", "1.1.1 安全原则",
        "1.2 分步实施", "第二章\u3000售后服务",
    ]
    assert headings()["商务部分"] == ["第一章\u3000报价说明"]  # 各册从首章重起

    for scheme, expected in (
        ("decimal", ["1 总体部署", "1.1 部署原则", "1.1.1 安全原则", "1.2 分步实施", "2 售后服务"]),
        ("gov", ["一、总体部署", "（一）部署原则", "1.安全原则", "（二）分步实施", "二、售后服务"]),
        ("none", ["总体部署", "部署原则", "安全原则", "分步实施", "售后服务"]),
    ):
        content = copy.deepcopy(_DIR_NUMBERING)
        content["numbering"] = scheme
        _seed_dir_artifact(env, content)
        assert docx_assemble_volume.invoke({}).startswith("[已合册]")
        assert headings()["技术部分"] == expected, scheme


def test_assemble_warns_self_numbered_node_names(env):
    """目录节点名自带编号（树格式红线违例）：合册照常编号（双重编号如实可见）
    + ⚠️ 点名请用户修目录产物——探测+提示裁决，不是门禁。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "第一章 技术方案", "level": 1, "children": [
                    {"目录名称": "1.1 项目理解", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": []},
                ]},
            ]}
        ]
    })
    docx_section_create.invoke(
        {"path": "body/1.1 项目理解", "title": "1.1 项目理解", "paragraphs": "正文"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]")
    assert "2 个目录节点名自带编号" in r
    assert "第一章 技术方案" in r and "1.1 项目理解" in r
    assert "双重编号" in r
    texts = [p.text for p in Document(str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert "第一章\u3000第一章 技术方案" in texts  # 双重编号如实可见
    assert "1.1 1.1 项目理解" in texts


# ---------- 终稿视角文本抽取 ----------


def test_section_text_lines_accepted_view(env):
    """校验对象=接受全部修订后的终稿：插入修订计入、删除修订不计。"""
    from app.tools.docx_ops import section_text_lines

    section = _make_section()
    assert docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": 1, "action": "insert_after", "text": "插入的新段落"},
    ])}).startswith("[已修订]")
    assert "插入的新段落" in section_text_lines(Document(str(_abs(env, section))))

    section2 = "body/删除演示.docx"
    assert docx_section_create.invoke(
        {"path": section2, "title": "X", "paragraphs": "保留段\n删除段"}
    ).startswith("[已创建]")
    assert docx_section_revise.invoke({"path": section2, "edits": json.dumps([
        {"para": 3, "action": "delete"},  # P3=删除段（P1=标题、P2=保留段）
    ])}).startswith("[已修订]")
    # 被删段以空行占位：P 段号与读视图对齐，终稿视角内容消失
    assert section_text_lines(Document(str(_abs(env, section2)))) == ["X", "保留段", ""]


# ---------- review 修复批（2026-09-07）----------


def test_revise_replace_across_revision_boundary(env):
    """二轮修订替换横跨「上一轮插入内容 + 原文」边界：每个匹配 run 原位包
    del（ins 内变 ins>del 嵌套、原文留 p>del），拒绝视角原文完整——共用单一
    del 挂首 run 位置会把原文 run 挪进 ins 子树，自校验必拒且模型无法自救。"""
    from app.tools.docx_ops import _accepted_text, _flatten_rejected

    section, _ = _injected_section(env)
    # 第一轮：换公司名（产生 w:ins）
    view = docx_section_read.invoke({"path": section})
    para_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln)
    r1 = docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"},
    ])})
    assert r1.startswith("[已修订]"), r1
    rejected_before = _flatten_rejected(Document(str(_abs(env, section))))

    # 第二轮：find 横跨「新公司名（ins 内容）+ 原文『承建』」边界
    r2 = docx_section_revise.invoke({"path": section, "edits": json.dumps([
        {"para": para_no, "action": "replace", "find": "上海中信科技有限公司承建", "text": "我司承建"},
    ])})
    assert r2.startswith("[已修订]"), r2  # 自校验通过=落盘成功

    chk = Document(str(_abs(env, section)))
    para = chk.paragraphs[para_no - 1]
    assert _accepted_text(para._p) == "本项目由我司承建，服务期三年。"
    assert _flatten_rejected(chk) == rejected_before  # 拒绝全部修订=原文逐字一致


def test_tools_survive_corrupt_docx(env):
    """损坏 docx（用户上传的招标/素材原件是外部输入，半截文件是现实情况）：
    工具返回失败文案，不打崩 run。"""
    section = _make_section()
    _abs(env, section).write_bytes(b"not a zip file")
    assert docx_section_read.invoke({"path": section}).startswith("[读取失败]")
    assert docx_section_revise.invoke(
        {"path": section, "edits": json.dumps([{"para": 1, "action": "delete"}])}
    ).startswith("[修订失败]")
    # 素材原件损坏：element_map 是解析落盘的独立 json 仍可用，Document() 打开才炸
    src = mlib.mt_files_dir() / env["file"]["file_name"]
    src.write_bytes(b"\x50\x4b broken zip")
    section2 = _make_section("技术部分/3.2 另一节.docx")
    assert docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": section2}
    ).startswith("[注入失败]")


def test_dest_path_lexical_escape(env):
    """body/../ 单级穿越：词法前缀检查放行、resolve 后落 work/ 根——resolve 后的
    body 目录祖先链校验必须拦截。"""
    from app.artifact_store import work_dir

    r = docx_section_create.invoke({"path": "body/../escaped", "title": "x"})
    assert r.startswith("[创建失败]") and "work/body/" in r, r
    assert not (work_dir(env["task"]["id"]) / "escaped.docx").exists()


def test_revise_insert_after_keeps_submission_order(env):
    """同段连续 insert_after：文档顺序与提交顺序一致（游标锚=上次插入的新段，
    否则 addnext 紧跟定位段导致整批倒序）。"""
    section = _make_section()
    edits = [
        {"para": 1, "action": "insert_after", "text": "第一条"},
        {"para": 1, "action": "insert_after", "text": "第二条"},
        {"para": 1, "action": "insert_after", "text": "第三条"},
    ]
    r = docx_section_revise.invoke({"path": section, "edits": json.dumps(edits)})
    assert r.startswith("[已修订]") and "3 处" in r, r
    from app.tools.docx_ops import _accepted_text

    reopened = Document(str(_abs(env, section)))
    texts = [_accepted_text(p._p) for p in reopened.paragraphs]
    i = texts.index("第一条")
    assert texts[i + 1] == "第二条" and texts[i + 2] == "第三条"
    # 插入段挂模板正文样式（Tender Body：标书行距/缩进），不吃中性 Normal
    assert reopened.paragraphs[i].style.name == "Tender Body"


def test_assemble_keeps_mismatched_own_heading(env):
    """节文件自带标题与树标题写法不一致：不吞（判据=样式+文本匹配节标题；
    只看 Heading1 会把素材自带的章标题静默剥掉——宁可重复发可见可修）。"""
    _seed_dir_artifact(env, {
        "response_documents": [
            {"name": "技术部分", "scope": "", "directory": [
                {"目录名称": "运维实施方案", "level": 1, "children": [],
                 "交付形态": "正文编写", "来源位置": []},
            ]}
        ]
    })
    # 建档标题写法与树节点不一致（素材整章注入后，素材自己的章标题同款形态）
    assert docx_section_create.invoke({"path": "body/运维实施方案", "title": "运维方案"}).startswith("[已创建]")
    assert docx_material_inject.invoke(
        {"block_id": env["block"]["id"], "dest": "body/运维实施方案.docx"}
    ).startswith("[已注入]")
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    texts = [p.text for p in Document(str(_abs(env, "body/整本-技术部分.docx"))).paragraphs]
    assert "第一章\u3000运维实施方案" in texts  # 树标题照发（带程序编号）
    assert "运维方案" in texts  # 节自带标题不被吞
    assert "运维服务方案" in texts  # 素材章标题内容保留


# ---------- 单图插入（docx_image_insert：独立图片进正文的放图通道） ----------


def _kb_image(file_name: str = "ISO27001证书.pdf") -> str:
    """知识库抽取图：造进 knowledge/parse/<文件名>/images/，返回工作区相对路径。"""
    from app.knowledge import store as kb_store

    img_dir = kb_store.kb_images_dir(file_name)
    img_dir.mkdir(parents=True, exist_ok=True)
    _tiny_png(img_dir / "img_001.png")
    return f"knowledge/parse/{file_name}/images/img_001.png"


def test_image_insert_after_paragraph(env):
    """知识库图插在指定段之后：位置/双修订标记/居中/全宽/拒绝视角不漂移/视图可见。"""
    rel = "body/技术部分/3.2 公司资质.docx"
    r = docx_section_create.invoke({
        "path": rel, "title": "3.2 公司资质",
        "paragraphs": "公司持有 ISO27001 信息安全管理体系认证证书。\n证书复印件见下图。",
    })
    assert r.startswith("[已创建]"), r
    r = docx_image_insert.invoke({"dest": rel, "image": _kb_image(), "after": "1"})
    assert r.startswith("[已插图]") and "P1 之后" in r, r

    doc = Document(str(_abs(env, rel)))
    paras = doc.paragraphs  # P1 标题、P2 图片、P3/P4 正文
    img_p = paras[1]._p
    assert len(img_p.findall(".//" + qn("a:blip"))) == 1
    # 段落标记修订（拒绝修订=整段含图消失）+ 内容修订（w:ins 包含图片 run）
    ppr = img_p.find(qn("w:pPr"))
    assert ppr.find(qn("w:rPr")).find(qn("w:ins")) is not None
    ins = img_p.find(qn("w:ins"))
    assert ins is not None and ins.find(".//" + qn("a:blip")) is not None
    # 居中、不挂正文样式（Tender Body 首行缩进会推偏图片）
    assert ppr.find(qn("w:jc")).get(qn("w:val")) == "center"
    assert ppr.find(qn("w:pStyle")) is None
    # 全宽适配：图片宽 = 页宽 − 左右边距（直取 wp:extent——python-docx 的
    # inline_shapes 视图只认 w:p 直接挂 w:r 的形态，修订包裹后匹配不到）
    sec = doc.sections[-1]
    cx = int(img_p.find(".//" + qn("wp:extent")).get("cx"))
    assert cx == int(sec.page_width - sec.left_margin - sec.right_margin)
    # 拒绝视角无空行漂移（工具内自校验的同款断言）
    assert _flatten_rejected(doc) == "3.2 公司资质\n公司持有 ISO27001 信息安全管理体系认证证书。\n证书复印件见下图。"
    # 读视图图片标签可见（模型后续按标签定位）
    assert "〔图×1〕" in docx_section_read.invoke({"path": rel})


def test_image_insert_append_at_end(env):
    """after 留空=追加到节末（sectPr 前）。"""
    rel = _make_section("技术部分/附图.docx")
    r = docx_image_insert.invoke({"dest": rel, "image": _kb_image()})
    assert r.startswith("[已插图]") and "节末" in r, r
    doc = Document(str(_abs(env, rel)))
    assert doc.paragraphs[-1]._p.findall(".//" + qn("a:blip"))


def test_image_insert_pdf_page_render(env):
    """PDF 原件按页现场渲染（抽取缺图的兜底）；任务前缀形态归一；页号越界人话报错。"""
    from app.artifact_store import sources_dir
    from tests import pdfgen

    pdf = sources_dir(env["task"]["id"]) / "证书扫描件.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdfgen.Pdf(pdf).text(72, 72, "ISO27001 CERTIFICATE").save()

    rel = _make_section("技术部分/3.2 公司资质.docx")
    r = docx_image_insert.invoke({"dest": rel, "image": "sources/证书扫描件.pdf", "page": 1})
    assert r.startswith("[已插图]") and "第 1 页" in r, r
    assert any(p._p.findall(".//" + qn("a:blip")) for p in Document(str(_abs(env, rel))).paragraphs)

    # 任务前缀形态（模型引用任务文件常带前缀）同样可插
    r = docx_image_insert.invoke({
        "dest": rel, "image": f"{env['task']['id']}/sources/证书扫描件.pdf", "after": "P1",
    })
    assert r.startswith("[已插图]"), r
    doc = Document(str(_abs(env, rel)))
    assert sum(len(p._p.findall(".//" + qn("a:blip"))) for p in doc.paragraphs) == 2

    # 页号超范围
    r = docx_image_insert.invoke({"dest": rel, "image": "sources/证书扫描件.pdf", "page": 5})
    assert r.startswith("[插图失败]") and "共 1 页" in r


def test_image_insert_webp_transcoded(env):
    """webp 图源自动转 png 插入。

    知识库上传白名单收 webp（IMAGE_EXTS）、python-docx 原生只认
    png/jpg/bmp/gif/tiff——此前「收得进、插不进」，报错让用户换格式
    （2026-09-14 复核批收口：Pillow 解码转 PNG 后插入，透明通道保留）。
    """
    from PIL import Image

    from app.knowledge import store as kb_store

    img_dir = kb_store.kb_images_dir("产品手册.pdf")
    img_dir.mkdir(parents=True, exist_ok=True)
    Image.new("RGBA", (2, 2), (255, 0, 0, 255)).save(img_dir / "img_001.webp", format="WEBP")

    rel = _make_section("技术部分/系统界面.docx")
    r = docx_image_insert.invoke({"dest": rel, "image": "knowledge/parse/产品手册.pdf/images/img_001.webp"})
    assert r.startswith("[已插图]") and "webp 已转 png" in r, r

    doc = Document(str(_abs(env, rel)))
    p_el = next(p._p for p in doc.paragraphs if p._p.findall(".//" + qn("a:blip")))
    blip = p_el.find(".//" + qn("a:blip"))
    part = doc.part.related_parts[blip.get(qn("r:embed"))]
    assert part.content_type == "image/png"  # 落进包里的部件是 PNG 而非 webp


def test_image_insert_rejects(env):
    """容错：越界/不存在/坏格式图源/段落号超范围/目标节缺失，全部人话报错。"""
    rel = _make_section("技术部分/附图.docx")
    assert docx_image_insert.invoke({"dest": rel, "image": "/etc/passwd"}).startswith("[插图失败]")
    assert "不存在" in docx_image_insert.invoke(
        {"dest": rel, "image": "knowledge/parse/无此文件.pdf/images/img_001.png"})
    # 坏 webp：解码失败人话报错（完好 webp 走转码分支，见 test_image_insert_webp_transcoded）
    from app.knowledge import store as kb_store

    webp_dir = kb_store.kb_images_dir("证书.webp")
    webp_dir.mkdir(parents=True, exist_ok=True)
    (webp_dir / "img_001.webp").write_bytes(b"fake-webp")
    out = docx_image_insert.invoke({"dest": rel, "image": "knowledge/parse/证书.webp/images/img_001.webp"})
    assert out.startswith("[插图失败]") and "webp" in out
    # docx 原件不作图源
    _tender_source(env)
    out = docx_image_insert.invoke({"dest": rel, "image": "sources/招标文件.docx"})
    assert out.startswith("[插图失败]") and "docx" in out
    # 段落号超范围
    assert "超范围" in docx_image_insert.invoke(
        {"dest": rel, "image": _kb_image(), "after": "99"})
    # 目标节未创建
    assert "不存在" in docx_image_insert.invoke(
        {"dest": "body/没有建过.docx", "image": _kb_image()})


def test_image_insert_into_revised_section(env):
    """已有修订标记的节里插图：rev_id 续号、自校验（拒绝视角不变）依旧通过。"""
    rel = _make_section("技术部分/3.3 业绩.docx")
    r = docx_section_revise.invoke({"path": rel, "edits": json.dumps([
        {"para": 1, "action": "insert_after", "text": "公司近三年业绩一览。"},
    ])})
    assert r.startswith("[已修订]"), r
    r = docx_image_insert.invoke({"dest": rel, "image": _kb_image(), "after": "2"})
    assert r.startswith("[已插图]"), r
    # 修订过的节拒绝视角：文本段保留（w:ins 插段整段消失）、图片段消失，无空行
    assert _flatten_rejected(Document(str(_abs(env, rel)))) == "3.1 运维方案"


# ---------- 待办批注（docx_comment_add：缺料/待澄清的正规落点） ----------


def test_comment_add_anchor_view_and_errors(env):
    """批注锚定指定段/默认节末；读视图行尾可见、终稿投影干净；错误路径人话。"""
    rel = "body/技术部分/3.3 项目团队.docx"
    docx_section_create.invoke({
        "path": rel, "title": "3.3 项目团队",
        "paragraphs": "团队拟投入 5 人。\n项目经理持一级建造师证书。",
    })
    r = docx_comment_add.invoke({
        "path": rel, "after": "2", "text": "项目经理证书编号缺失，需向用户确认后回填",
    })
    assert r.startswith("[已加批注]") and "P2" in r and "共 1 条" in r, r
    r2 = docx_comment_add.invoke({"path": rel, "text": "驻场人数是否含后台待澄清"})
    assert r2.startswith("[已加批注]") and "P3" in r2 and "共 2 条" in r2, r2

    view = docx_section_read.invoke({"path": rel})
    assert "待办批注 2 条" in view
    assert "〔批注：项目经理证书编号缺失，需向用户确认后回填〕" in view
    # 终稿视角（接受修订后评委看到的文本）不含批注内容——批注不是正文
    doc = Document(str(_abs(env, rel)))
    assert not any("证书编号缺失" in t or "待澄清" in t for t in section_text_lines(doc))
    # 加批注后再修订：自校验（拒绝视角）不漂移、批注保留
    assert docx_section_revise.invoke({"path": rel, "edits": json.dumps([
        {"para": 2, "action": "insert_after", "text": "证书编号见附件。"},
    ])}).startswith("[已修订]")
    assert "待办批注 2 条" in docx_section_read.invoke({"path": rel})

    # 错误路径
    assert docx_comment_add.invoke({"path": rel, "text": ""}).startswith("[批注失败]")
    assert "段落号超范围" in docx_comment_add.invoke({"path": rel, "after": "99", "text": "x"})
    assert "after 须为段落号" in docx_comment_add.invoke({"path": rel, "after": "P", "text": "x"})
    assert "文件不存在" in docx_comment_add.invoke({"path": "body/无此节.docx", "text": "x"})


def test_revision_and_comment_authors_are_branded(env):
    """修订/批注署名钉在 _AUTHOR/_COMMENT_AUTHOR 单点上（Word 审阅侧栏用户可见）。"""
    from app.tools.docx_ops import _AUTHOR, _COMMENT_AUTHOR

    assert _AUTHOR == _COMMENT_AUTHOR == "灵燕智能", (_AUTHOR, _COMMENT_AUTHOR)

    section, _ = _injected_section(env)
    view = docx_section_read.invoke({"path": section})
    para_no = next(int(ln.split("]")[0][2:]) for ln in view.splitlines() if COMPANY in ln)
    assert docx_section_revise.invoke({
        "path": section,
        "edits": json.dumps([
            {"para": para_no, "action": "replace", "find": COMPANY, "text": "上海中信科技有限公司"},
        ]),
    }).startswith("[已修订]")
    assert docx_comment_add.invoke({
        "path": section, "after": str(para_no), "text": "署名核对",
    }).startswith("[已加批注]")

    doc = Document(str(_abs(env, section)))
    # 只看目标段（素材自带修订不参与，署名断言不因夹具形态漂移）
    assert {
        el.get(qn("w:author"))
        for el in doc.paragraphs[para_no - 1]._p.iter()
        if el.tag in (qn("w:ins"), qn("w:del"))
    } == {_AUTHOR}
    assert {c.author for c in doc.comments if (c.text or "").strip()} == {_COMMENT_AUTHOR}


# ---------- 样式/编号迁移去重 + 拷贝卫生（2026-09-13 目录乱号批） ----------
# 实证背景：素材自带「styleId=3 name="heading 2" + 挂多级编号」样式随注入/合册
# 迁入，与目标内建 heading 2 同名并存——WPS/LibreOffice 按名解析把整本全部二级
# 标题套上无意义连续序号（真实任务 98 行正文污染）；撞 id 定义不同则被静默跳过、
# 拷贝段绑到无关样式；直挂 outlineLvl 段（格式件标题/被改写的正文段）漏摘进导航。


def _add_para_style(doc, sid: str, name: str, ppr: str = "", default: bool = False) -> None:
    d = ' w:default="1"' if default else ""
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="{sid}"{d}>'
        f'<w:name w:val="{name}"/>{ppr}</w:style>'
    ))


def _ref_para(doc, sid: str, text: str = "段"):
    p = doc.add_paragraph(text)
    p._p.get_or_add_pPr().insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{sid}"/>'))
    return p._p


def _style_name_map(doc) -> dict[str, str]:
    out = {}
    for st in doc.styles.element.findall(qn("w:style")):
        ne = st.find(qn("w:name"))
        out[st.get(qn("w:styleId"))] = (ne.get(qn("w:val")) if ne is not None else "") or ""
    return out


def test_merge_styles_same_name_renames_and_disarms_numbering():
    """同名样式迁入改名（治 D1）：目标仍只有一个 "heading 2"（内建、无编号），
    素材样式改名 "heading 2 2"——按名合并的污染路径（整本标题被套连续序号）
    被断掉。2026-09-16 拍板收窄：迁入**标题样式**的编号引用剥除（素材编号
    标题不再原样渲染，节内小标题一律裸标题）。"""
    from app.tools.docx_ops import _merge_missing_styles

    src = Document()
    _add_para_style(
        src, "3", "heading 2",
        '<w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="11"/></w:numPr></w:pPr>',
    )
    el = _ref_para(src, "3", "素材小节标题")
    dst = Document()  # 默认模板自带 heading 1-9

    moved, remap = _merge_missing_styles(src, dst, [el])
    names = _style_name_map(dst)
    h2 = [sid for sid, nm in names.items() if nm.lower() == "heading 2"]
    assert len(h2) == 1 and h2[0] == "Heading2"  # 内建唯一、未被素材顶掉
    assert moved == 1 and remap == {}  # id 不撞：原 id 迁入
    assert names.get("3") == "heading 2 2"  # 同名 → 改名迁入
    st3 = next(s for s in dst.styles.element.findall(qn("w:style"))
               if s.get(qn("w:styleId")) == "3")
    st3_ppr = st3.find(qn("w:pPr"))
    assert st3_ppr is None or st3_ppr.find(qn("w:numPr")) is None  # 标题样式编号已剥
    assert el.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "3"


def test_merge_styles_body_list_style_keeps_numbering():
    """非标题样式（正文列表）迁入编号保真：剥编号只命中标题类样式，
    圆点/(1) 型列表样式及其编号定义照常迁入。"""
    from app.tools.docx_ops import _merge_missing_styles

    src = Document()
    _add_para_style(
        src, "21", "List Paragraph",
        '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="12"/></w:numPr></w:pPr>',
    )
    el = _ref_para(src, "21", "列表项一条")
    dst = Document()

    moved, _remap = _merge_missing_styles(src, dst, [el])
    assert moved == 1
    st21 = next(s for s in dst.styles.element.findall(qn("w:style"))
                if s.get(qn("w:styleId")) == "21")
    assert st21.find(qn("w:pPr")).find(qn("w:numPr")) is not None  # 正文列表保真


def test_merge_styles_id_collision_different_def_remaps():
    """撞 id 且定义不同 → 新 id 迁入 + 引用改写（治 H1）：此前静默跳过会把
    拷贝段绑到目标里同 id 的无关样式（id 先到先得随合并序漂移，观感错位）。"""
    from app.tools.docx_ops import _merge_missing_styles

    src = Document()
    _add_para_style(src, "9", "my heading", "<w:pPr><w:b/></w:pPr>")
    el = _ref_para(src, "9")
    dst = Document()
    _add_para_style(dst, "9", "my body", "<w:pPr><w:i/></w:pPr>")

    moved, remap = _merge_missing_styles(src, dst, [el])
    assert moved == 1 and set(remap) == {"9"}
    new_sid = remap["9"]
    names = _style_name_map(dst)
    assert names[new_sid] == "my heading"  # 无同名冲突：名字不改
    assert names["9"] == "my body"  # 目标原样式不动
    assert el.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == new_sid


def test_merge_styles_equivalent_definition_reuses():
    """撞 id 但定义等价 → 沿用目标（原行为回归），不重复迁入。"""
    from app.tools.docx_ops import _merge_missing_styles

    src = Document()
    _add_para_style(src, "77", "shared style", "<w:pPr><w:b/></w:pPr>")
    el = _ref_para(src, "77")
    dst = Document()
    _add_para_style(dst, "77", "shared style", "<w:pPr><w:b/></w:pPr>")
    n_before = len(dst.styles.element.findall(qn("w:style")))

    moved, remap = _merge_missing_styles(src, dst, [el])
    assert moved == 0 and remap == {}
    assert len(dst.styles.element.findall(qn("w:style"))) == n_before
    assert el.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "77"


def test_merge_styles_strips_default_flag():
    """源文档的默认样式标记不迁入（w:default 顶掉目标默认会让全文档换底样式）。"""
    from app.tools.docx_ops import _merge_missing_styles

    src = Document()
    _add_para_style(src, "88", "src normal", default=True)
    el = _ref_para(src, "88")
    dst = Document()

    moved, _remap = _merge_missing_styles(src, dst, [el])
    st = next(s for s in dst.styles.element.findall(qn("w:style"))
              if s.get(qn("w:styleId")) == "88")
    assert moved == 1 and st.get(qn("w:default")) is None


def test_numbering_style_links_guarded_and_remapped():
    """迁入编号的 pStyle 绑定（治 H2）：解析到目标内建标题/Normal 的剥除——
    章节编号按树序写进标题文本，样式绑定会把拷入标题一起计数打乱章序；
    样式换 id 后按重映射表改写，不绑到目标里同 id 的无关样式。"""
    from app.tools.docx_ops import _merge_missing_numbering

    def _armed(src, abs_id: str, num_id: str, pstyle: str):
        np_ = src.part.numbering_part.element
        ab = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            f'<w:lvlText w:val="%1"/><w:pStyle w:val="{pstyle}"/></w:lvl></w:abstractNum>'
        )
        first = np_.find(qn("w:num"))
        first.addprevious(ab) if first is not None else np_.append(ab)
        np_.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abs_id}"/></w:num>'
        ))
        p = src.add_paragraph("带列表段")
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
        ))
        return p._p

    # ① 绑定指向目标内建 Heading2 → 迁入时剥除
    src1, dst1 = Document(), Document()
    el1 = _armed(src1, "90", "30", "Heading2")
    _merge_missing_numbering(src1, dst1, [el1])
    ab1 = next(a for a in dst1.part.numbering_part.element.findall(qn("w:abstractNum"))
               if a.get(qn("w:abstractNumId")) == "90")
    assert ab1.find(qn("w:lvl")).find(qn("w:pStyle")) is None

    # ② 样式 id 重映射 → 绑定改写到新 id
    src2, dst2 = Document(), Document()
    el2 = _armed(src2, "91", "31", "4")
    _merge_missing_numbering(src2, dst2, [el2], style_id_remap={"4": "77"})
    ab2 = next(a for a in dst2.part.numbering_part.element.findall(qn("w:abstractNum"))
               if a.get(qn("w:abstractNumId")) == "91")
    assert ab2.find(qn("w:lvl")).find(qn("w:pStyle")).get(qn("w:val")) == "77"


def test_demote_extra_headings_direct_outline():
    """摘大纲扩判据（治 D2）：无标题样式、直挂 outlineLvl<9 的段落也摘——
    招标格式件标题/被写手改写过的正文段的漏网形态；普通段不动。"""
    from app.tools.docx_ops import _demote_extra_headings

    doc = Document()
    p1 = doc.add_paragraph("直挂大纲段")
    p1._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
    p2 = doc.add_paragraph("普通段")

    n = _demote_extra_headings(doc, [p1._p, p2._p])
    ol = p1._p.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert n == 1 and ol is not None and ol.get(qn("w:val")) == "9"
    p2pr = p2._p.find(qn("w:pPr"))
    assert p2pr is None or p2pr.find(qn("w:outlineLvl")) is None


def test_source_inject_strips_outline_and_toc_bookmarks(env):
    """拷入卫生：招标件段落的直挂大纲级别与 _Toc 死书签被剥（保真=版式，
    不含导航元数据/源文档内部书签——同段拷进多节会造成书签 id 重复）；
    普通书签保留。"""
    from app.artifact_store import sources_dir

    p = sources_dir(env["task"]["id"]) / "格式附件.docx"
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("资格一览表", 1)
    para = doc.add_paragraph("附件11人员资格一览表")
    para._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
    para._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="146" w:name="_Toc268191677"/>'))
    para._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="146"/>'))
    para._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="147" w:name="keepme"/>'))
    para._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="147"/>'))
    doc.save(p)

    section = _make_section()
    r = docx_source_inject.invoke({"source": "格式附件.docx", "dest": section})
    assert r.startswith("[已注入]"), r
    body = Document(str(_abs(env, section))).element.body
    assert body.findall(f".//{qn('w:outlineLvl')}") == []
    starts = list(body.iter(qn("w:bookmarkStart")))
    assert not [b for b in starts if (b.get(qn("w:name")) or "").startswith("_")]
    assert any(b.get(qn("w:name")) == "keepme" for b in starts)


def _make_armed_material_docx(path: Path) -> None:
    """复刻腾励/云枢形态的「武装」素材：styleId=3 name="heading 2" 挂多级编号
    （abstractNum 带 pStyle 链接 2/3），标题段**同时**直挂 numPr（真实素材
    双通道形态）——拷贝即引入同名编号标题样式；另带一条正文列表段（编号保真面）。"""
    doc = Document()
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="3">'
        '<w:name w:val="heading 2"/><w:basedOn w:val="1"/><w:next w:val="1"/>'
        '<w:uiPriority w:val="9"/><w:qFormat/>'
        '<w:pPr><w:keepNext/><w:numPr><w:ilvl w:val="1"/><w:numId w:val="11"/></w:numPr>'
        '<w:outlineLvl w:val="1"/></w:pPr></w:style>'
    ))
    doc.styles.element.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="21">'
        '<w:name w:val="List Paragraph"/>'
        '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="12"/></w:numPr></w:pPr></w:style>'
    ))
    np_ = doc.part.numbering_part.element
    ab = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="95">'
        '<w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1"/><w:pStyle w:val="2"/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1.%2"/><w:pStyle w:val="3"/></w:lvl></w:abstractNum>'
    )
    ab_bullet = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="96">'
        '<w:multiLevelType w:val="hybridMultilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        '<w:lvlText w:val="·"/></w:lvl></w:abstractNum>'
    )
    first = np_.find(qn("w:num"))
    if first is not None:
        first.addprevious(ab)
        first.addprevious(ab_bullet)
    else:
        np_.append(ab)
        np_.append(ab_bullet)
    np_.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="11"><w:abstractNumId w:val="95"/></w:num>'))
    np_.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="12"><w:abstractNumId w:val="96"/></w:num>'))
    p = doc.add_paragraph("素材小节标题甲")
    ppr = p._p.get_or_add_pPr()
    ppr.insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="3"/>'))
    # 直挂 numPr（素材真实形态：样式级与段落级编号并存，段落级优先渲染）
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="11"/></w:numPr>'
    ))
    bl = doc.add_paragraph("素材列表项一条")
    bl._p.get_or_add_pPr().insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="21"/>'))
    doc.add_paragraph("素材正文一句，内容足够构成检索块。")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def test_assemble_disarms_armed_styles_and_strips_outline(env):
    """端到端（注入+合册两层同治）：武装素材注入节文件、合册成整本后——
    目标唯一 "heading 2"（骨架样式无编号）；素材标题样式迁入但**编号剥除**
    （2026-09-16 拍板收窄：样式级+段落级双通道都摘，节内小标题裸标题）、
    标题段直挂 numPr 剥除；正文列表样式与编号定义保真迁入；无直挂大纲
    级别段、无 _Toc 死书签。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    section = "body/项目理解与需求分析.docx"
    r = docx_section_create.invoke(
        {"path": section, "title": "项目理解与需求分析", "paragraphs": "项目理解正文第一段。"}
    )
    assert r.startswith("[已创建]"), r

    src = mlib.mt_files_dir() / "同名编号标题素材.docx"
    _make_armed_material_docx(src)
    from app import db

    f = db.mt_insert_file("同名编号标题素材.docx", "hash_armed_1")
    mlib.run_parse(f["id"])
    blk = mlib.create_block(f["id"], "公司介绍", "奥哲介绍素材", ranges=[[1, 10 ** 6]])
    ri = docx_material_inject.invoke({"block_id": blk["id"], "dest": section})
    assert ri.startswith("[已注入]"), ri
    # 节文件层：同名即改名（注入层与合册层共用同一迁移函数）；标题段已裸化
    sec_doc = Document(str(_abs(env, section)))
    sec_names = _style_name_map(sec_doc)
    assert [s for s, nm in sec_names.items() if nm.lower() == "heading 2"] == ["Heading2"]
    sec_body = sec_doc.element.body
    sec_heading_p = next(
        p for p in sec_body.findall(f".//{qn('w:p')}")
        if "素材小节标题甲" in "".join(p.itertext())
    )
    sec_hppr = sec_heading_p.find(qn("w:pPr"))
    assert sec_hppr.find(qn("w:numPr")) is None  # 段落直挂编号已剥
    sec_st3 = next(s for s in sec_doc.styles.element.findall(qn("w:style"))
                   if s.get(qn("w:styleId")) == "3")
    sec_st3_ppr = sec_st3.find(qn("w:pPr"))
    assert sec_st3_ppr.find(qn("w:numPr")) is None  # 样式级编号已剥
    sec_bl = next(
        p for p in sec_body.findall(f".//{qn('w:p')}")
        if "素材列表项一条" in "".join(p.itertext())
    )
    # 正文列表的编号保真走样式通道（节内直挂本就无）
    assert sec_bl.find(qn("w:pPr")).find(qn("w:pStyle")) is not None

    # 存量残留模拟：历史节文件里的直挂大纲段 + _Toc 书签（注入剥除对存量无效，
    # 合册侧 _strip_copy_residue + _demote_extra_headings 兜底）
    sec_doc = Document(str(_abs(env, section)))
    p = sec_doc.add_paragraph("历史残留的直挂大纲段")
    p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
    p._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="146" w:name="_TocX"/>'))
    p._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="146"/>'))
    sec_doc.save(str(_abs(env, section)))

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    assert "节内小标题降级 1 段" in r  # 武装素材标题段按树深度降级（深度 2 + 源级 2 → H3）
    final = Document(str(_abs(env, "body/整本-技术部分.docx")))
    names = _style_name_map(final)
    assert [s for s, nm in names.items() if nm.lower() == "heading 2"] == ["Heading2"]
    # 2026-09-16 降级批：拷入标题段改挂内建 Heading3（与骨架同族、层级下沉），
    # 素材标题样式不再被引用 → 整本完全不迁入（styles 无 styleId=3）
    body = final.element.body
    final_heading_p = next(
        p for p in body.findall(f".//{qn('w:p')}")
        if "素材小节标题甲" in "".join(p.itertext())
    )
    fh_ps = final_heading_p.find(qn("w:pPr")).find(qn("w:pStyle"))
    assert fh_ps.get(qn("w:val")) == "Heading3"
    assert "3" not in names
    # 降级后 demote 照旧钉大纲 9（导航/目录不收录不变）
    fh_ol = final_heading_p.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert fh_ol is not None and fh_ol.get(qn("w:val")) == "9"
    # 正文列表样式在整本里编号保真（剥除只命中标题类）
    st21 = next(s for s in final.styles.element.findall(qn("w:style"))
                if s.get(qn("w:styleId")) == "21")
    assert st21.find(qn("w:pPr")).find(qn("w:numPr")) is not None
    # 骨架样式无编号（编号按树序写进标题文本，不走样式绑定）
    h2 = next(s for s in final.styles.element.findall(qn("w:style"))
              if s.get(qn("w:styleId")) == "Heading2")
    h2ppr = h2.find(qn("w:pPr"))
    assert h2ppr is None or h2ppr.find(qn("w:numPr")) is None
    # 迁入编号不绑内建标题/Normal
    guard = {f"heading {i}" for i in range(1, 10)} | {"normal"}
    for a in final.part.numbering_part.element.findall(qn("w:abstractNum")):
        for lv in a.findall(qn("w:lvl")):
            ps = lv.find(qn("w:pStyle"))
            if ps is not None:
                assert names.get(ps.get(qn("w:val")), "").lower() not in guard
    # 直挂大纲段被剥（余下的 outlineLvl 全是摘除标记 9）、无 _Toc 书签
    ols = body.findall(f".//{qn('w:outlineLvl')}")
    assert ols and all(o.get(qn("w:val")) == "9" for o in ols)
    assert not [b for b in body.iter(qn("w:bookmarkStart"))
                if (b.get(qn("w:name")) or "").startswith("_")]


def _numpr_para(doc, text: str, *, pstyle: str | None, num_id: int, ilvl: int = 0,
                outline: int | None = None):
    p = doc.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    if pstyle:
        ppr.insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{pstyle}"/>'))
    if outline is not None:
        ppr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{outline}"/>'))
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))
    return p._p


def test_strip_copy_residue_strips_heading_numbering():
    """拷入卫生剥标题编号（2026-09-16 批）：标题样式段/直挂大纲段的直挂
    numPr 整棵剥；正文列表段（非标题）保真；heading_ids 不传=旧行为（不碰
    numPr，兼容）。"""
    from app.tools.docx_ops import _strip_copy_residue

    doc = Document()
    head = _numpr_para(doc, "素材标题", pstyle="Heading2", num_id=11, ilvl=4)
    outl = _numpr_para(doc, "格式件标题", pstyle=None, num_id=12, outline=2)
    plain = _numpr_para(doc, "列表项", pstyle=None, num_id=13)

    ids = {"Heading2"}
    for el in (head, outl, plain):
        _strip_copy_residue(el, ids)

    assert head.find(qn("w:pPr")).find(qn("w:numPr")) is None  # 标题样式段剥
    assert head.find(qn("w:pPr")).find(qn("w:pStyle")) is not None
    outl_ppr = outl.find(qn("w:pPr"))
    assert outl_ppr.find(qn("w:numPr")) is None  # 直挂大纲段剥
    assert outl_ppr.find(qn("w:outlineLvl")) is None  # 大纲级别照旧剥
    assert plain.find(qn("w:pPr")).find(qn("w:numPr")) is not None  # 正文列表保真

    # 不传 heading_ids：旧行为，numPr 不动（存量调用兼容）
    doc2 = Document()
    keep = _numpr_para(doc2, "素材标题", pstyle="Heading2", num_id=11)
    _strip_copy_residue(keep)
    assert keep.find(qn("w:pPr")).find(qn("w:numPr")) is not None


# ---------- 节内小标题按树深度降级（2026-09-16 整本平铺观感批） ----------
# 实证背景：真实任务整本 218 个裸 H2 小标题与 54 个真节标题（5.1 式）同级平铺，
# 评委视角分不清层级；节文件层级本来就对（节标题 H1 > 小标题 H2/H3），撞级只发生
# 在整本（节变 H2）——修合册这一层：目标级 = clamp(树深度 + 源级 − 1, 2, 9)。


def _subhead_section(env, path: str, title: str, subheads: list[tuple[str, int]],
                     table_head: str | None = None) -> None:
    """建节并往里塞写手小标题（(文本, 级) 列表），可选表格内标题段。"""
    r = docx_section_create.invoke({"path": path, "title": title, "paragraphs": f"{title}正文。"})
    assert r.startswith("[已创建]"), r
    p = _abs(env, path if path.endswith(".docx") else path + ".docx")
    doc = Document(str(p))
    for text, lvl in subheads:
        para = doc.add_paragraph(text)
        para.style = doc.styles[f"Heading {lvl}"]
    if table_head:
        tbl = doc.add_table(rows=1, cols=1)
        cell_p = tbl.rows[0].cells[0].paragraphs[0]
        cell_p.text = table_head
        cell_p.style = doc.styles["Heading 2"]
    doc.save(str(p))


def _vol_pstyle(vol_doc, text: str) -> tuple[str, str | None]:
    """按文本找整本段落，回 (pStyleId, outlineLvl|None)。"""
    p = next(
        p_ for p_ in vol_doc.element.body.findall(f".//{qn('w:p')}")
        if text in "".join(p_.itertext())
    )
    ppr = p.find(qn("w:pPr"))
    ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
    ol = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    return (
        ps.get(qn("w:val")) if ps is not None else None,
        ol.get(qn("w:val")) if ol is not None else None,
    )


def _vol_ptext(vol_doc, text: str) -> str:
    """按子串找整本段落，回接受视角全文（编号断言用）。"""
    p = next(
        p_ for p_ in vol_doc.element.body.findall(f".//{qn('w:p')}")
        if text in "".join(p_.itertext())
    )
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def test_assemble_restyles_subheads_by_tree_depth(env):
    """深度 2 节：H2→Heading3、H3→Heading4（层级下沉一档）、表格内标题段
    同降（规则统一）；demote 照旧钉 olvl=9（导航/目录不收录不变）；
    节文件零改动（工作台层级本来就对）。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    _subhead_section(
        env, "body/项目理解与需求分析", "项目理解与需求分析",
        [("二级小标题", 2), ("三级小标题", 3)], table_head="表内标题",
    )

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    assert "节内小标题降级 3 段" in r

    vol = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert _vol_pstyle(vol, "二级小标题") == ("Heading3", "9")
    assert _vol_pstyle(vol, "三级小标题") == ("Heading4", "9")
    assert _vol_pstyle(vol, "表内标题") == ("Heading3", "9")  # 表格内同降
    # 程序拼编号（2026-09-16 二批）：H3/H4 按节内出现顺序接续节号 1.1
    assert _vol_ptext(vol, "二级小标题") == "1.1.1 二级小标题"
    assert _vol_ptext(vol, "三级小标题") == "1.1.1.1 三级小标题"
    assert _vol_ptext(vol, "表内标题") == "1.1.2 表内标题"
    # 骨架节标题不受影响（仍是 H2）
    assert _vol_pstyle(vol, "1.1 项目理解与需求分析")[0] == "Heading2"
    # 节文件原样未动（重开断言仍是写手层级）
    sec = Document(str(_abs(env, "body/项目理解与需求分析.docx")))
    sec_ps = lambda t: next(  # noqa: E731
        p_.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val"))
        for p_ in sec.element.body.findall(f".//{qn('w:p')}")
        if t in "".join(p_.itertext())
    )
    assert sec_ps("二级小标题") == "Heading2"
    assert sec_ps("三级小标题") == "Heading3"


def test_assemble_subhead_depth1_keeps_level(env):
    """深度 1 章叶（培训方案式，章下直接挂内容）：源 H2 → 目标级 2 不变
    （公式钉住：章叶小标题本就是节级观感，不越级下沉）。"""
    tree = {
        "response_documents": [
            {
                "name": "技术部分", "scope": "",
                "directory": [
                    {"目录名称": "培训方案", "level": 1, "children": [],
                     "交付形态": "正文编写", "来源位置": []},
                ],
            }
        ]
    }
    _seed_dir_artifact(env, tree)
    _subhead_section(env, "body/培训方案", "培训方案", [("培训目标与总体安排", 2)])

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    vol = Document(str(_abs(env, "body/整本-技术部分.docx")))
    sid, olvl = _vol_pstyle(vol, "培训目标与总体安排")
    assert sid == "Heading2"  # 深度 1 + 源级 2 → 目标 2（观感不降）
    assert olvl == "9"  # 导航/目录照旧不收录
    assert _vol_ptext(vol, "培训目标与总体安排") == "培训目标与总体安排"  # 章叶不编号


_DIR_SUBHEAD_GUARD = {
    "response_documents": [
        {
            "name": "技术部分", "scope": "",
            "directory": [
                {"目录名称": "编制说明章", "level": 1, "children": [
                    {"目录名称": "索引明细", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": []},
                ]},
                {"目录名称": "目录", "level": 1, "children": [],
                 "交付形态": "模板或附件填充", "来源位置": []},
                {"目录名称": "格式章", "level": 1, "children": [
                    {"目录名称": "投标函格式", "level": 2, "children": [],
                     "交付形态": "模板或附件填充", "来源位置": []},
                ]},
                {"目录名称": "正文章", "level": 1, "children": [
                    {"目录名称": "正文节", "level": 2, "children": [],
                     "交付形态": "正文编写", "来源位置": []},
                ]},
            ],
        }
    ]
}


def test_assemble_subhead_guard_frontmatter_and_nonprose(env):
    """降级护栏（守卫放在深度 2 上才能观察——深度 1 时目标级=源级，降与不降
    同形）：前置区（目录节点前）与格式件章（NON_PROSE 招标件零改动保真）的
    标题段保持原 pStyle 不降；同一次合册里正文节照常降级（对照组）。"""
    _seed_dir_artifact(env, _DIR_SUBHEAD_GUARD)
    _subhead_section(env, "body/索引明细", "索引明细", [("前置区小标题", 2)])
    _subhead_section(env, "body/投标函格式", "投标函格式", [("格式件内部标题", 2)])
    _subhead_section(env, "body/正文节", "正文节", [("正文小标题", 2)])

    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    vol = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert _vol_pstyle(vol, "前置区小标题")[0] == "Heading2"  # 前置区不降
    assert _vol_pstyle(vol, "格式件内部标题")[0] == "Heading2"  # 格式件保真不降
    assert _vol_pstyle(vol, "正文小标题")[0] == "Heading3"  # 对照组：正文照降
    # 护栏同时免编号；对照组正文节接续节号（前置区不占章号：格式章=第一章、
    # 正文章=第二章 → 正文节 2.1 → 小标题 2.1.1）
    assert _vol_ptext(vol, "前置区小标题") == "前置区小标题"
    assert _vol_ptext(vol, "格式件内部标题") == "格式件内部标题"
    assert _vol_ptext(vol, "正文小标题") == "2.1.1 正文小标题"


def test_assemble_subhead_numbering_gov_and_none(env):
    """小标题编号格式矩阵：gov 接续中文层级（节（一）下 H3→1.、H4→（1））；
    numbering=none 全册零编号（骨架与小标题都不加）。"""
    tree = {
        "numbering": "gov",
        "response_documents": [
            {
                "name": "技术部分", "scope": "",
                "directory": [
                    {"目录名称": "建设方案", "level": 1, "children": [
                        {"目录名称": "总体设计", "level": 2, "children": [],
                         "交付形态": "正文编写", "来源位置": []},
                    ]},
                ],
            }
        ],
    }
    _seed_dir_artifact(env, tree)
    _subhead_section(env, "body/总体设计", "总体设计",
                     [("政务一级小标题", 2), ("政务二级小标题", 3)])
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    vol = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert _vol_ptext(vol, "政务一级小标题") == "1.政务一级小标题"
    assert _vol_ptext(vol, "政务二级小标题") == "（1）政务二级小标题"

    tree["numbering"] = "none"
    _seed_dir_artifact(env, tree)
    r = docx_section_create.invoke({"path": "body/总体设计", "title": "总体设计",
                                    "paragraphs": "总体设计正文。", "replace": True})
    assert r.startswith(("[已创建]", "[已重建]")), r
    p = _abs(env, "body/总体设计.docx")
    doc = Document(str(p))
    para = doc.add_paragraph("无编号小标题")
    para.style = doc.styles["Heading 2"]
    doc.save(str(p))
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    vol = Document(str(_abs(env, "body/整本-技术部分.docx")))
    assert _vol_ptext(vol, "无编号小标题") == "无编号小标题"
    texts = [_vol_ptext(vol, "总体设计")]
    assert all(not t.startswith(("一、", "（一）", "1.", "1 ")) for t in texts)


# ---------- 同文件并发写防线（2026-09-13 事故批） ----------


def test_parallel_comment_add_same_file(env):
    """并发写坏节文件事故的回归哨兵（机制断言）。

    事故形态：模型一 turn 连发 8 条 docx_comment_add 打同一节文件，ToolNode
    真并行执行，原地 save 互相覆盖把 zip 写坏（BadZipFile）→ 写手转而 read_file
    读证书图 → base64 内联撑爆上下文。修法=按路径互斥 + 原子存盘；本测试
    Barrier 对齐 8 线程同时开火，断言三条：全部成功、8 条批注一条不丢（锁
    防丢更新）、无 .tmp 残件且文件完好可开（原子替换）。
    """
    import contextvars
    import threading

    from app.artifact_store import work_dir
    from app.tools.docx_ops import _comment_texts

    rel = _make_section("body/并发批注.docx")
    n = 8
    barrier = threading.Barrier(n)
    results: list[str] = []
    rlock = threading.Lock()

    def worker(i: int):
        try:
            barrier.wait(timeout=10)
            r = docx_comment_add.invoke({"path": rel, "after": "P1", "text": f"待办 {i}：需用户确认"})
        except Exception as e:  # noqa: BLE001
            r = f"[异常] {type(e).__name__}: {e}"
        with rlock:
            results.append(r)

    # 生产形态：ToolNode 的 ContextThreadPoolExecutor 逐任务拷贝 context
    # （runctx 的 task_id 走 contextvar）——裸 Thread 不带上下文，须显式 ctx.run
    threads = [
        threading.Thread(target=contextvars.copy_context().run, args=(lambda i=i: worker(i),))
        for i in range(n)
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=30)

    assert len(results) == n
    bad = [r for r in results if not r.startswith("[已加批注]")]
    assert not bad, bad  # 零 BadZipFile、零异常
    dst = work_dir(env["task"]["id"]) / "body" / "并发批注.docx"
    doc = Document(str(dst))  # 文件完好可开（未写坏）
    assert len(_comment_texts(doc)) == n  # 8 条批注全在——并行不丢更新
    assert not list(dst.parent.glob("*.tmp"))  # 原子存盘无残件


# ---------- 表格通道批（2026-09-14）：建节混排 / 图示生成 / 合册图注编号 ----------


def test_create_body_blocks_mixed(env):
    """body 块序列建节：段落/表题/表格按序落、表题挂 Tender Caption 在表格上方、
    表头加粗+浅灰底纹；paragraphs 参数向后兼容。"""
    body = json.dumps(
        [
            {"type": "p", "text": "我方组建专职实施团队，按招标岗位口径配置，如下表："},
            {
                "type": "table",
                "caption": "项目团队岗位配置",
                "header": ["岗位", "人数", "职责"],
                "rows": [["项目经理", "1", "对本项目交付负总责"], ["架构师", "1", "主持架构设计"]],
            },
            {"type": "p", "text": "同一成员可兼任多个角色，兼任不降低履行质量。"},
        ],
        ensure_ascii=False,
    )
    r = docx_section_create.invoke({"path": "body/团队配置", "title": "团队配置", "body": body})
    assert r.startswith("[已创建]"), r
    assert "正文 2 段+表格 1 张" in r
    doc = Document(str(_abs(env, "body/团队配置.docx")))
    kinds = [e.tag.split("}")[-1] for e in doc.element.body]
    assert kinds == ["p", "p", "p", "tbl", "p", "sectPr"]  # 标题/引言/表题/表格/收尾
    assert doc.paragraphs[2].style.name == "Tender Caption"  # 表题在表格上方
    tbl = doc.tables[0]
    assert [c.text for c in tbl.rows[0].cells] == ["岗位", "人数", "职责"]
    shd = tbl.cell(0, 0)._tc.find(qn("w:tcPr")).find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "F2F2F2"
    assert tbl.cell(0, 0).paragraphs[0].runs[0].bold is True
    view = docx_section_read.invoke({"path": "body/团队配置"})
    assert "[T1] R1：C1=岗位 C2=人数 C3=职责" in view


def test_create_body_blocks_validation(env):
    """body 参数校验：非 JSON / 列超限 / 单格超长 / rows 缺失，报错给修复写法。"""
    r = docx_section_create.invoke({"path": "body/bad1", "title": "x", "body": "不是json"})
    assert r.startswith("[创建失败]") and "body 不是合法 JSON" in r
    wide = json.dumps([{"type": "table", "header": [f"列{i}" for i in range(9)], "rows": [["x"] * 9]}])
    r = docx_section_create.invoke({"path": "body/bad2", "title": "x", "body": wide})
    assert "列数 9 超上限" in r
    long_cell = json.dumps([{"type": "table", "header": ["a"], "rows": [["x" * 201]]}])
    r = docx_section_create.invoke({"path": "body/bad3", "title": "x", "body": long_cell})
    assert "201 字超上限" in r
    r = docx_section_create.invoke({"path": "body/bad4", "title": "x",
                                    "body": json.dumps([{"type": "table", "header": ["a"]}])})
    assert "缺 rows" in r
    # 行宽超 header 报错不截断（2026-09-14 review 修复）：静默截断会丢格
    fat_row = json.dumps([{"type": "table", "header": ["a", "b", "c"],
                           "rows": [["1", "2", "3", "4", "5"]]}])
    r = docx_section_create.invoke({"path": "body/bad5", "title": "x", "body": fat_row})
    assert r.startswith("[创建失败]") and "第 1 行 5 列超过 header 的 3 列" in r
    # 少列的宽容保留：补空串建表成功
    thin = json.dumps([{"type": "table", "header": ["a", "b", "c"], "rows": [["only"]]}])
    r = docx_section_create.invoke({"path": "body/thin", "title": "x", "body": thin})
    assert r.startswith("[已创建]") and "表格 1 张" in r
    assert len(Document(str(_abs(env, "body/thin.docx"))).tables[0].columns) == 3


def test_diagram_layered_gantt_radial(env):
    """三种网格图示：layered 层间箭头行/域头深灰、gantt 周数推定+填色格+里程碑行、
    radial 中列纵向合并；表格不占 P 序号、锚定插入与图注跟随。"""
    docx_section_create.invoke({"path": "body/进度", "title": "进度", "paragraphs": "第一段。\n第二段。"})
    r = docx_diagram_insert.invoke(
        {
            "dest": "body/进度",
            "kind": "layered",
            "spec": json.dumps(
                {"layers": [{"title": "决策层", "items": ["领导小组"]},
                            {"title": "执行层", "items": ["开发组", "测试组"]}]},
                ensure_ascii=False,
            ),
            "after": "2",
            "caption": "项目组织架构",
        }
    )
    assert r.startswith("[已插图示]"), r
    assert "T1" in r and "2 层" in r and "P2 之后" in r
    doc = Document(str(_abs(env, "body/进度.docx")))
    t1 = doc.tables[0]
    assert len(t1.rows) == 3 and len(t1.columns) == 5  # 2 层 + 1 箭头行
    assert t1.cell(1, 2).text == "↓"
    shd = t1.cell(0, 0)._tc.find(qn("w:tcPr")).find(qn("w:shd"))
    assert shd.get(qn("w:fill")) == "D9D9D9"  # 域头深灰
    # 表格不占 P 序号：插在 P2（第一段）后，图注段紧随表格、第二段顺延为 P4
    view = docx_section_read.invoke({"path": "body/进度"})
    assert "[P2]（Tender Body）第一段。" in view
    assert "[P3]（Tender Caption）项目组织架构" in view
    assert "[P4]（Tender Body）第二段。" in view

    r2 = docx_diagram_insert.invoke(
        {
            "dest": "body/进度",
            "kind": "gantt",
            "spec": json.dumps(
                {"tasks": [{"name": "需求调研", "start": 1, "end": 3}],
                 "milestones": [{"week": 3, "label": "需求评审"}]}
            ),
        }
    )
    assert "1 项任务×3 周" in r2  # weeks 省略按最晚结束周推定
    doc = Document(str(_abs(env, "body/进度.docx")))
    t2 = doc.tables[1]
    assert len(t2.rows) == 3 and len(t2.columns) == 4  # 表头+任务+里程碑
    assert t2.cell(0, 3).text == "W3"
    assert t2.cell(1, 3)._tc.find(qn("w:tcPr")).find(qn("w:shd")) is not None  # 区间格填色
    assert t2.cell(2, 3).text == "◆" and t2.cell(2, 0).text == "需求评审"

    r3 = docx_diagram_insert.invoke(
        {
            "dest": "body/进度",
            "kind": "radial",
            "spec": json.dumps(
                {"center": "人员主记录", "left": ["公司与组织", "职位与任职"],
                 "right": ["任免业务"]},
                ensure_ascii=False,
            ),
        }
    )
    assert r3.startswith("[已插图示]")
    doc = Document(str(_abs(env, "body/进度.docx")))
    t3 = doc.tables[2]
    assert t3.cell(0, 2)._tc is t3.cell(1, 2)._tc  # 中列纵向合并


def test_diagram_validation(env):
    """图示参数校验：kind 词表 / 甘特周超限给聚合出路 / layered 缺字段。"""
    docx_section_create.invoke({"path": "body/dv", "title": "dv", "paragraphs": "x。"})
    r = docx_diagram_insert.invoke({"dest": "body/dv", "kind": "wireframe", "spec": "{}"})
    assert r.startswith("[图示失败]") and "kind 须为" in r and "flow" in r
    r = docx_diagram_insert.invoke(
        {"dest": "body/dv", "kind": "gantt",
         "spec": json.dumps({"tasks": [{"name": "t", "start": 1, "end": 99}]})}
    )
    assert "超上限（≤36）" in r and "按月" in r
    r = docx_diagram_insert.invoke(
        {"dest": "body/dv", "kind": "layered", "spec": json.dumps({"layers": [{"title": "A层"}]})}
    )
    assert "缺 title 或 items" in r


def test_assemble_caption_numbering(env):
    """合册图注/表题全局重编号：表题（表格上方）编「表 X-Y」、图注（图示下方）编
    「图 X-Y」，跨节累计、章号随树序；与章节编号同一原则（节文件不带编号）。"""
    _seed_dir_artifact(env, _DIR_SINGLE)
    body = json.dumps(
        [
            {"type": "p", "text": "团队配置如下："},
            {"type": "table", "caption": "岗位配置", "header": ["岗位"], "rows": [["项目经理"]]},
        ],
        ensure_ascii=False,
    )
    docx_section_create.invoke(
        {"path": "body/项目理解与需求分析", "title": "项目理解与需求分析", "body": body}
    )
    docx_section_create.invoke({"path": "body/总体设计方案", "title": "总体设计方案", "paragraphs": "方案段。"})
    docx_diagram_insert.invoke(
        {"dest": "body/总体设计方案", "kind": "layered",
         "spec": json.dumps({"layers": [{"title": "应用层", "items": ["门户"]}]}, ensure_ascii=False),
         "caption": "系统架构"}
    )
    r = docx_assemble_volume.invoke({})
    assert r.startswith("[已合册]"), r
    assert "图注/表题编号 2 处" in r
    chk = Document(str(_abs(env, "body/整本-技术部分.docx")))
    texts = [p.text for p in chk.paragraphs]
    assert "表 1-1 岗位配置" in texts  # 表题在表格上方 → 表系列
    assert "图 1-1 系统架构" in texts  # 图示图注在表格下方 → 图系列


# ---------- 流程图（kind=flow，mermaid 消费方 2026-09-14 批三） ----------


def test_flow_spec_validation(env):
    """flow spec 校验：缺 nodes/edges、超节点上限、悬空边、非法 shape、id 重复。"""
    docx_section_create.invoke({"path": "body/fv", "title": "fv", "paragraphs": "x。"})
    r = docx_diagram_insert.invoke({"dest": "body/fv", "kind": "flow", "spec": "{}"})
    assert r.startswith("[图示失败]") and "flow 须带 nodes" in r
    many = {"nodes": [{"id": f"n{i}", "label": f"节点{i}"} for i in range(21)], "edges": [["n0", "n1"]]}
    r = docx_diagram_insert.invoke({"dest": "body/fv", "kind": "flow", "spec": json.dumps(many)})
    assert "节点数 21 超上限" in r and "拆成两张图" in r
    r = docx_diagram_insert.invoke({"dest": "body/fv", "kind": "flow",
        "spec": json.dumps({"nodes": [{"id": "a", "label": "甲"}], "edges": [["a", "zz"]]})})
    assert "不存在的节点" in r
    r = docx_diagram_insert.invoke({"dest": "body/fv", "kind": "flow",
        "spec": json.dumps({"nodes": [{"id": "a", "label": "甲", "shape": "circle"}],
                            "edges": [["a", "a"]]})})
    assert "shape 须为" in r
    r = docx_diagram_insert.invoke({"dest": "body/fv", "kind": "flow",
        "spec": json.dumps({"nodes": [{"id": "a", "label": "甲"}, {"id": "a", "label": "乙"}],
                            "edges": [["a", "a"]]})})
    assert "重复" in r


def test_flow_to_mermaid_translation():
    """JSON 拓扑 → mermaid 文本：形状语法/边标签/转义（引号→#quot;、竖线→全角）。"""
    parsed = {
        "nodes": [
            {"id": "a", "label": "提交故障", "shape": "round"},
            {"id": "b", "label": "判定级别", "shape": "diamond"},
            {"id": "c", "label": '含"引号"的节点', "shape": "rect"},
        ],
        "edges": [["a", "b"], ["b", "c", "重|大"]],
    }
    m = _flow_to_mermaid(parsed)
    assert m.startswith("flowchart TD")
    assert 'a("提交故障")' in m          # round → 圆括号
    assert 'b{"判定级别"}' in m          # diamond → 花括号
    assert 'c["含#quot;引号#quot;的节点"]' in m  # 引号转义
    assert "a --> b" in m
    assert "b -->|重｜大| c" in m        # 边标签 + 竖线转全角


def test_flow_via_render_queue(env):
    """flow 经渲染队列：登记 mermaid 载荷 → 回执 PNG → 节文件含图；无回执降级。"""
    import contextvars
    import threading

    from fastapi.testclient import TestClient

    from app.main import app

    docx_section_create.invoke({"path": "body/fv", "title": "fv", "paragraphs": "x。"})
    spec = json.dumps({
        "nodes": [{"id": "a", "label": "提交故障"},
                  {"id": "b", "label": "判定级别", "shape": "diamond"},
                  {"id": "c", "label": "应急处置"}],
        "edges": [["a", "b"], ["b", "c", "重大"], ["c", "b", "重判"]],
    }, ensure_ascii=False)
    client = TestClient(app)
    out: dict = {}
    ctx = contextvars.copy_context()

    def run():
        out["r"] = docx_diagram_insert.invoke(
            {"dest": "body/fv", "kind": "flow", "spec": spec, "caption": "故障处置流程"})

    t = threading.Thread(target=lambda: ctx.run(run))
    t.start()
    # 等 pending 出现,断言载荷是 flow+mermaid 文本(非模型手写——程序翻译产物)
    import time as _t

    pend = []
    for _ in range(100):
        pend = client.get("/api/render/pending").json()["requests"]
        if pend:
            break
        _t.sleep(0.05)
    assert len(pend) == 1
    assert pend[0]["kind"] == "flow"
    assert pend[0]["mermaid"].startswith("flowchart TD")
    assert 'b{"判定级别"}' in pend[0]["mermaid"]
    rr = client.post("/api/render/figure", data={"request_id": pend[0]["request_id"]},
                     files={"file": ("f.png", _tiny_png_bytes(), "image/png")})
    assert rr.status_code == 200
    t.join(timeout=5)
    assert out["r"].startswith("[已插图示]") and "流程图 1 张" in out["r"]
    assert "节点 3/连线 3" in out["r"]
    doc = Document(str(_abs(env, "body/fv.docx")))
    assert len(doc.element.body.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")) == 1


def _tiny_png_bytes() -> bytes:
    import struct as _s
    import zlib as _z

    def chunk(t: bytes, d: bytes) -> bytes:
        c = t + d
        return _s.pack(">I", len(d)) + c + _s.pack(">I", _z.crc32(c) & 0xFFFFFFFF)

    ihdr = _s.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
            + chunk(b"IDAT", _z.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b""))


def test_image_height_cap(env):
    """插图高度封顶（2026-09-14 实测修复）：mermaid 长链流程图显示高 34.7~62.5cm
    占满整页还溢出——按版心宽等比缩放后超 18cm 的按高度反缩（等比、变窄居中）。"""
    import io

    from PIL import Image as PILImage

    from app.tools.docx_ops import _tracked_image_paragraph

    buf = io.BytesIO()
    PILImage.new("RGB", (100, 1000), "white").save(buf, format="PNG")  # 细高图 1:10
    docx_section_create.invoke({"path": "body/hc", "title": "hc", "paragraphs": "x。"})
    from docx import Document as _D

    doc = _D(str(_abs(env, "body/hc.docx")))
    p_el = _tracked_image_paragraph(doc, buf.getvalue())
    ext = p_el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    if ext is None:  # python-docx 命名空间形态兜底
        ext = next(e for e in p_el.iter() if e.tag.endswith("}extent"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    assert cy <= 18 * 360000 + 1, f"显示高 {cy/360000:.1f}cm 超封顶"
    assert abs(cx / cy - 100 / 1000) < 0.02, "反缩须等比（宽高比不变）"


def test_revise_returns_now_text_without_reread(env):
    """纯改动批（无插段）返回「改动后现文」——replace/格编辑后模型直接看到结果，
    回读动机消除（2026-09-14 回读再收敛）；删除段标已删除；插段批维持全视图。"""
    import json as _json

    docx_section_create.invoke(
        {"path": "body/现文节", "title": "现文节",
         "paragraphs": "第一段原文甲。\n第二段原文乙。\n第三段原文丙。"}
    )
    edits = _json.dumps([
        {"para": 2, "action": "replace", "find": "原文甲", "text": "改后之甲"},
        {"para": 3, "action": "delete"},
    ], ensure_ascii=False)
    r = docx_section_revise.invoke({"path": "body/现文节", "edits": edits})
    assert "[已修订]" in r and "改动后现文（接受修订视角，无需回读确认）" in r
    assert "P2 现文：第一段改后之甲。" in r
    assert "P3 现文：（已标记删除，接受修订后此段消失）" in r
    assert "最新读视图如下" not in r  # 无插段=不附全视图（现文块已覆盖确认需求）

    # 格编辑:建节带表,fill 后返回格现文
    body = _json.dumps([
        {"type": "table", "header": ["项", "值"], "rows": [["工期", ""], ["人员", "5"]]},
    ], ensure_ascii=False)
    docx_section_create.invoke({"path": "body/现文表", "title": "现文表", "body": body})
    edits2 = _json.dumps([{"table": 1, "row": 2, "col": 2, "action": "fill", "text": "3 个月"}])
    r2 = docx_section_revise.invoke({"path": "body/现文表", "edits": edits2})
    assert "T1R2C2 现文：3 个月" in r2

    # 长段改动点在尾部也须现文可见（截断对齐 _VIEW_TEXT_LIMIT，2026-09-14 review 修复：
    # 原 200 字截断会把长段后半的替换结果裁掉，「把结果送到眼前」对长段失效）
    tail = "铺垫" * 130 + "结尾改这里"
    docx_section_create.invoke(
        {"path": "body/现文长段", "title": "现文长段", "paragraphs": tail + "原文。"}
    )
    r3 = docx_section_revise.invoke(
        {"path": "body/现文长段",
         "edits": _json.dumps([{"para": 2, "action": "replace", "find": "原文", "text": "改后"}],
                              ensure_ascii=False)}
    )
    assert "结尾改这里改后。" in r3  # 尾部改动点未被截断裁掉


def test_sanitize_name_leading_dot():
    """清洗去前导点：`.` 开头的文件名会被 workbench/本轮文件按隐藏文件隐掉而
    账面照算（账实不一致）——正文标题不受影响，仅文件名（2026-09-15 审计）。"""
    from app.tools.body_contract import sanitize_name

    assert sanitize_name(".NET 架构方案") == "NET 架构方案"
    assert sanitize_name("..反思与总结") == "反思与总结"
    assert sanitize_name("///") == "未命名"
    assert sanitize_name("A/B:方案") == "A B 方案"
    assert sanitize_name("技术部分") == "技术部分"  # 无前导点照旧
