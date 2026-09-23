"""类型化 Artifact：/api/contracts、/api/artifacts 列表与内容端点。"""

import os

from app import artifact_store, db, publish

KEY = "tender.directory/tender-response-docs@1"


def _content(name="技术部分"):
    return {
        "response_documents": [
            {"name": name, "scope": "", "directory": [{"目录名称": "目录", "level": 1, "children": []}]}
        ]
    }


def _task_conv():
    """§16：发布需要任务上下文；client fixture 已隔离 DATA_DIR。"""
    task = db.create_task("测试任务")
    conv = db.create_conversation(task["id"])
    return task, conv


def _pub(content, conv=None, **kw):
    """发布到任务；不传 conv 时新建任务+会话（覆盖语义需要同一任务）。"""
    if conv is None:
        _, conv = _task_conv()
    return publish.publish_artifact(KEY, content, task_id=conv["task_id"], conversation_id=conv["id"], **kw)


def test_contracts_endpoint(client):
    r = client.get("/api/contracts")
    assert r.status_code == 200
    contracts = r.json()["contracts"]
    keys = [c["key"] for c in contracts]
    assert KEY in keys
    dir_c = next(c for c in contracts if c["key"] == KEY)
    assert dir_c["kind"] == "tender.directory"
    assert dir_c["cardinality"] == "task-single"
    assert dir_c["llm_write_mode"] == "suggest"
    assert dir_c["editable"] is True
    assert dir_c["default_display_name"] == "投标目录"


def test_artifacts_list_and_content(client):
    task, conv = _task_conv()
    m = publish.publish_artifact(
        KEY, _content(), task_id=task["id"], conversation_id=conv["id"],
        source={"skill": "demo-skill", "thread_id": conv["id"], "run_id": "r_1"},
    )
    aid = m["artifact_id"]

    r = client.get("/api/artifacts")
    assert r.status_code == 200
    arts = r.json()["artifacts"]
    assert len(arts) == 1
    a = arts[0]
    assert a["artifact_id"] == aid
    assert a["display_name"] == "投标目录"
    assert a["kind"] == "tender.directory"
    assert a["schema_id"] == "tender-response-docs"
    assert a["schema_version"] == 1
    assert a["editable"] is True
    assert a["source"] == {"thread_id": conv["id"], "run_id": "r_1"}
    # 文件归任务：包路径在 <task>/work/artifacts/ 下；分隔符随平台（os.sep，
    # 写死 "/" 在 Windows 上假红——v0.2.1 win 首跑实证）
    assert a["path"].endswith(
        f"{task['id']}{os.sep}work{os.sep}artifacts{os.sep}{aid}{os.sep}content.json"
    )

    r = client.get(f"/api/artifacts/{aid}/content")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/json")
    assert r.json()["response_documents"][0]["name"] == "技术部分"

    assert client.get("/api/artifacts/nope/content").status_code == 404


def test_meta_endpoint(client):
    """轻量探测：只回版本号；保存后版本号递增。"""
    task, conv = _task_conv()
    aid = _pub(_content(), conv)["artifact_id"]

    r = client.get(f"/api/artifacts/{aid}/meta")
    assert r.status_code == 200
    assert r.json() == {"artifact_id": aid, "content_seq": 1}

    updated = _content(name="用户改后")
    rr = client.put(f"/api/artifacts/{aid}/content", json={"content": updated, "base_content_seq": 1})
    assert rr.status_code == 200
    assert client.get(f"/api/artifacts/{aid}/meta").json()["content_seq"] == 2

    assert client.get("/api/artifacts/nope/meta").status_code == 404


def test_put_content_flow_no_lease(client):
    """编辑保存无租约（文件夹语义）：正常保存、409 探测、force 用户裁决覆盖（留恢复点）、422。"""
    m = _pub(_content())
    aid = m["artifact_id"]

    # 正常保存：无需任何租约
    updated = _content()
    updated["response_documents"][0]["name"] = "用户改后"
    r = client.put(f"/api/artifacts/{aid}/content", json={"content": updated, "base_content_seq": 1})
    assert r.status_code == 200
    assert r.json()["content_seq"] == 2
    assert "用户改后" in client.get(f"/api/artifacts/{aid}/content").text

    # 正常（非 force）保存同样留恢复点（2026-09-23 A4：此前只有 force 留底，
    # 纯编辑流里 restore_available 恒 false、「恢复上一版」永不出现）
    assert client.get("/api/artifacts").json()["artifacts"][0]["restore_available"] is True
    rr = client.post(f"/api/artifacts/{aid}/restore")
    assert rr.status_code == 200
    assert "用户改后" not in client.get(f"/api/artifacts/{aid}/content").text  # 回到发布版
    # 恢复本身也留底：再恢复一次回到「用户改后」（seq 随每次写入递增，用 meta 对齐基线）
    assert client.post(f"/api/artifacts/{aid}/restore").status_code == 200
    assert "用户改后" in client.get(f"/api/artifacts/{aid}/content").text
    seq = client.get(f"/api/artifacts/{aid}/meta").json()["content_seq"]

    # 再保存一次（seq → seq+1）
    r = client.put(f"/api/artifacts/{aid}/content", json={"content": updated, "base_content_seq": seq})
    assert r.status_code == 200
    seq = r.json()["content_seq"]

    # 基于旧 seq 保存 → 409 探测信号（客户端据此弹「拉取最新/保留我的」）
    r = client.put(f"/api/artifacts/{aid}/content", json={"content": updated, "base_content_seq": seq - 1})
    assert r.status_code == 409

    # force：用户裁决保留自己的版本，无条件覆盖，被顶掉的版本留恢复点
    mine = _content()
    mine["response_documents"][0]["name"] = "我的版本"
    r = client.put(
        f"/api/artifacts/{aid}/content",
        json={"content": mine, "base_content_seq": 1, "force": True},
    )
    assert r.status_code == 200
    assert "我的版本" in client.get(f"/api/artifacts/{aid}/content").text
    final_seq = r.json()["content_seq"]

    # 列表带 content_seq 与 restore_available
    a = client.get("/api/artifacts").json()["artifacts"][0]
    assert a["content_seq"] == final_seq
    assert a["restore_available"] is True

    # schema 破坏始终拒绝（force 也不豁免）
    r = client.put(
        f"/api/artifacts/{aid}/content",
        json={"content": {"foo": 1}, "base_content_seq": final_seq, "force": True},
    )
    assert r.status_code == 422


def test_restore_roundtrip(client):
    """恢复点安全网：发布覆盖留底 → 恢复上一版 → 恢复本身可再撤销。"""
    _, conv = _task_conv()
    _pub(_content("初版"), conv=conv, source={"skill": "t"})
    m2 = _pub(_content("新版"), conv=conv, source={"skill": "t"})
    aid = m2["artifact_id"]

    r = client.post(f"/api/artifacts/{aid}/restore")
    assert r.status_code == 200
    assert "初版" in client.get(f"/api/artifacts/{aid}/content").text

    # 再次恢复 → 回到「新版」（恢复前也留了底）
    r = client.post(f"/api/artifacts/{aid}/restore")
    assert r.status_code == 200
    assert "新版" in client.get(f"/api/artifacts/{aid}/content").text


def test_put_force_preserves_fresh_publish_state(client):
    """跨线程竞态防线：PUT 读行后 worker 线程的发布落地（emitted 复位 0、seq+1），
    编辑保存持锁重读，不得用陈旧整行回写抹掉 emitted/last_run_id——否则该次
    发布的 artifact.created 永不发出，产物面板失去刷新触发。"""
    _, conv = _task_conv()
    m = _pub(_content("v1"), conv=conv, source={"skill": "t", "thread_id": conv["id"], "run_id": "run_a"})
    aid = m["artifact_id"]
    db.mark_emitted(aid)  # 模拟 run 边界已发出 artifact.created（emitted=1）

    stale = db.get_artifact_index(aid)  # PUT 已过检查点的读行快照
    _pub(_content("v2"), conv=conv, source={"skill": "t", "thread_id": conv["id"], "run_id": "run_b"})

    mine = _content("我的版本")
    r = client.put(
        f"/api/artifacts/{aid}/content",
        json={"content": mine, "base_content_seq": stale["content_seq"], "force": True},
    )
    assert r.status_code == 200

    row = db.get_artifact_index(aid)
    assert row["emitted"] == 0  # 未被陈旧行（emitted=1）抹掉
    assert row["last_run_id"] == "run_b"
    assert db.pending_emit("run_b"), "run_b 的 artifact.created 账本不得被 PUT 清掉"
    assert "我的版本" in client.get(f"/api/artifacts/{aid}/content").text
    assert row["content_seq"] == stale["content_seq"] + 2  # publish +1、PUT +1


def test_read_content_resolved_blocks_escape(client, tmp_path):
    """containment 读侧（编辑留底/恢复/转正/发布留底/read_artifact 的读取入口）：
    包内 content.json 被换成指向 workspace 外的 symlink 时返回 None，不外泄。"""
    _, conv = _task_conv()
    m = _pub(_content(), conv=conv)
    aid = m["artifact_id"]
    assert "技术部分" in artifact_store.read_content_resolved(aid, m)

    outside = tmp_path / "secret.json"
    outside.write_text('{"stolen": true}', encoding="utf-8")
    cp = artifact_store.content_path(aid, m)
    cp.unlink()
    cp.symlink_to(outside)
    assert artifact_store.read_content_resolved(aid, m) is None
    # 非 resolved 版本仍跟随 symlink（对照：防护确实来自 resolve + is_relative_to）
    assert artifact_store.read_content(aid, m) is not None


def test_put_unknown_artifact(client):
    r = client.put(
        "/api/artifacts/art_0000000000ff/content", json={"content": {}, "base_content_seq": 1}
    )
    assert r.status_code == 404
    assert client.post("/api/artifacts/art_0000000000ff/restore").status_code == 404


def test_artifacts_list_filters_missing_package(client):
    """索引在、磁盘包被删：列表过滤、内容端点 410（与工具侧 containment 同标准）。"""
    m = _pub(_content())
    aid = m["artifact_id"]

    import shutil

    shutil.rmtree(artifact_store.artifact_dir(aid, m))
    assert client.get("/api/artifacts").json()["artifacts"] == []
    assert client.get(f"/api/artifacts/{aid}/content").status_code == 410


# ---------- 文件型产物（tender.volume=整本标书 docx，2026-09-13） ----------

VOLUME_KEY = "tender.volume/tender-volume-docx@1"


def test_volume_artifact_file_endpoint(client):
    """GET /artifacts/{aid}/file：包内唯一 docx 流式下发（docx MIME + 附件名）；
    列表行的 content_type 走契约注册表、editable=False。"""
    import hashlib

    from docx import Document

    task, conv = _task_conv()
    body_dir = artifact_store.work_dir(task["id"]) / "body"
    body_dir.mkdir(parents=True, exist_ok=True)
    src = body_dir / "整本-技术册.docx"
    doc = Document()
    doc.add_paragraph("整本正文")
    doc.save(src)

    m = publish.publish_file_artifact(
        VOLUME_KEY,
        file_path=src,
        content_meta={
            "filename": src.name,
            "book": "技术册",
            "size": src.stat().st_size,
            "sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
            "merged_sections": 1,
        },
        display_name="技术册",
        source={"skill": "tender-body", "thread_id": conv["id"], "run_id": "r_1"},
        task_id=task["id"],
        conversation_id=conv["id"],
    )
    aid = m["artifact_id"]

    r = client.get(f"/api/artifacts/{aid}/file")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert r.content == src.read_bytes()
    assert "attachment" in r.headers.get("content-disposition", "")

    a = next(a for a in client.get("/api/artifacts").json()["artifacts"] if a["artifact_id"] == aid)
    assert a["content_type"].startswith("application/vnd.openxmlformats")
    assert a["editable"] is False
    assert a["restore_available"] is False  # 文件型产物明确不做恢复点

    # editable=False 纵深：PUT content / restore 均拒绝（此前无校验，直调 API 可改机器元信息）
    r = client.put(f"/api/artifacts/{aid}/content", json={"content": {}, "base_content_seq": 1})
    assert r.status_code == 409
    assert client.post(f"/api/artifacts/{aid}/restore").status_code == 409

    assert client.get("/api/artifacts/nope/file").status_code == 404

    # 契约目录端点含新 key（客户端启动对账能看到）
    keys = [c["key"] for c in client.get("/api/contracts").json()["contracts"]]
    assert VOLUME_KEY in keys


def test_volume_artifact_file_prefers_content_filename(client):
    """污染包择包防呆：包内残留多个 docx 时优先下发 content.json 登记的
    filename（旧版清理只走覆盖路径，短路分支不清理，历史包可能多 docx），
    盲取字典序第一会把旧册发给用户。"""
    import hashlib

    from docx import Document

    task, conv = _task_conv()
    body_dir = artifact_store.work_dir(task["id"]) / "body"
    body_dir.mkdir(parents=True, exist_ok=True)
    src = body_dir / "整本-技术册.docx"
    Document().save(src)  # 空文档即可，只比对字节

    m = publish.publish_file_artifact(
        VOLUME_KEY,
        file_path=src,
        content_meta={
            "filename": src.name,
            "book": "技术册",
            "size": src.stat().st_size,
            "sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
            "merged_sections": 1,
        },
        display_name="技术册",
        source={"skill": "tender-body", "thread_id": conv["id"], "run_id": "r_1"},
        task_id=task["id"],
        conversation_id=conv["id"],
    )
    aid = m["artifact_id"]

    # 塞一个字典序更早的旧册 docx（"aaa" < "技"），模拟历史污染包
    decoy = artifact_store.artifact_dir(aid, m) / "整本-aaa旧册.docx"
    decoy.write_bytes(b"stale-bytes")

    r = client.get(f"/api/artifacts/{aid}/file")
    assert r.status_code == 200
    assert r.content == src.read_bytes()
