"""知识库入库管线（v3 事实层）：三合一抽取+说明段、失败降级、图片落盘（仅供查看）。
章节拆分/写作素材在独立素材库（test_materials_lib.py），知识库无章节块。"""

import json

from app import db
from app.knowledge import store
from app.knowledge.ingest import run_ingest
from tests.util import init_env


class _FakeResp:
    def __init__(self, content):
        self.content = content


def _extract_json(statement=None, doc_type="qualification_certificate", fields=None):
    return json.dumps(
        {"doc_type": doc_type, "confidence": 0.9, "statement": statement or "",
         "fields": fields or {}},
        ensure_ascii=False,
    )


class _SmartLLM:
    """抽取工序 fake：按文件名判类型（标书类→past_proposal，默认→证书）。"""

    def __init__(self, **kwargs):
        pass

    def invoke(self, messages):
        prompt = messages[-1][1]
        if "文件名：历史标书" in prompt or "文件名：大标书" in prompt or "文件名：范文" in prompt:
            return _FakeResp(_extract_json(
                doc_type="past_proposal",
                statement="技术标共 8 章。",
                fields={"project_name": {"value": "某园区项目", "source": "L1"}},
            ))
        return _FakeResp(_extract_json(
            statement="ISO9001 质量管理体系认证，证书编号 CN-001（第1页），有效期至 2028-06-30（第1页）。",
            fields={"valid_until": {"value": "2028-06-30", "source": "第1页"}},
        ))


class _BoomLLM:
    def __init__(self, **kwargs):
        pass

    def invoke(self, messages):
        raise RuntimeError("llm down")


def _setup(tmp_path, monkeypatch, llm=None):
    init_env(tmp_path, monkeypatch)
    store.ensure_dirs()
    if llm is not None:
        monkeypatch.setenv("LLM_API_KEY", "test-key")
        import langchain_deepseek

        monkeypatch.setattr(langchain_deepseek, "ChatDeepSeek", llm)


def _upload(name: str, content: bytes) -> str:
    import hashlib

    src = store.kb_files_dir() / name
    src.write_bytes(content)
    item = db.kb_insert_item(file_name=name, file_hash=hashlib.sha256(content).hexdigest(),
                             title=src.stem, ext=src.suffix)
    return item["id"]


def test_extract_statement_and_doc_type(tmp_path, monkeypatch):
    """三合一抽取：类型+说明+时间字段落库；说明段可检索。"""
    _setup(tmp_path, monkeypatch, llm=_SmartLLM)
    kid = _upload("iso证书.txt", ("# 证书\n\n" + "ISO9001 质量管理体系认证证书正文内容，" * 4).encode())
    item = run_ingest(kid)
    assert item["parse_status"] == "ready"
    assert item["extract_status"] == "done"
    assert item["doc_type"] == "qualification_certificate"
    sug = json.loads(item["suggested_metadata"])
    assert "ISO9001" in sug["statement"] and "第1页" in sug["statement"]

    from app.knowledge import fts

    hits = db.kb_search_segments(fts.build_match_expr("CN-001"), limit=5)
    assert any(h.get("section_path") == "§statement" for h in hits)
    assert item.get("progress") is None


def test_fact_gate_no_blocks(tmp_path, monkeypatch):
    """素材分离闸门：知识库任何类型都不建章节块/不产块行（写作素材在素材库）。"""
    _setup(tmp_path, monkeypatch, llm=_SmartLLM)
    kid = _upload("历史标书.txt", ("# 一、方案\n\n" + "智慧园区总体架构内容文本。" * 30).encode())
    item = run_ingest(kid)
    assert item["doc_type"] == "past_proposal"
    # 知识库侧无块表/无素材端点；materials.json 不产出
    _, _, _, mp = store.kb_parse_paths("历史标书.txt")
    assert not (mp.parent / "blocks.json").exists()
    # 检索段仍就绪（outline 段+说明段）
    from app.knowledge import fts

    assert db.kb_search_segments(fts.build_match_expr("智慧园区"), limit=3)


def test_extract_failure_degrades(tmp_path, monkeypatch):
    """抽取失败：可检索、字段空、extract failed——独立降级。"""
    _setup(tmp_path, monkeypatch, llm=_BoomLLM)
    kid = _upload("资料.txt", ("# 资料\n\n" + "公司拥有信息安全服务资质证书与专业团队，服务过多家大型客户并获得好评。" * 2).encode())
    item = run_ingest(kid)
    assert item["parse_status"] == "ready"
    assert item["extract_status"] == "failed"
    from app.knowledge import fts

    assert db.kb_search_segments(fts.build_match_expr("资质"), limit=3)


def test_no_key_skips_extract(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)  # 无 LLM_API_KEY
    kid = _upload("c.txt", b"# x\n\n" + "内容".encode() * 100)
    item = run_ingest(kid)
    assert item["extract_status"] == "skipped"


def test_docx_image_count_in_meta(tmp_path, monkeypatch):
    """docx 内嵌图片计数进解析概况（不再静默丢弃）。"""
    _setup(tmp_path, monkeypatch)
    from app.parse import convert as parse_convert

    real = parse_convert

    class _FakeResult:
        md = "# 证\n\n内容"
        info = {"conversion": "docx-native", "tables": 0, "image_count": 37}

    monkeypatch.setattr("app.knowledge.ingest.parse_convert", lambda src: _FakeResult())
    _ = real  # 保留引用
    kid = _upload("带图.docx", b"fake-docx-bytes")
    run_ingest(kid)
    _, _, meta_path, _ = store.kb_parse_paths("带图.docx")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    assert meta["image_count"] == 37


def test_extract_validation(tmp_path, monkeypatch):
    """extract.validate_suggested：类型兜底/锚点键分流/写法类短说明/事实类过短说明无效。"""
    from app.knowledge.extract import validate_suggested

    out = validate_suggested(
        {"doc_type": "bogus", "statement": "x" * 10,
         "fields": {"valid_until": {"value": "2027-01-01", "source": "L1"},
                    "自由键": {"value": "v", "source": "L2"}}},
        "随便.docx",
    )
    assert out["doc_type"] == "other"  # 无 hint 兜底
    assert "valid_until" in out["fields"] and "自由键" not in out["fields"]
    assert "自由键" in out["extra"]
    assert out.get("statement") is None or out["statement"] == ""  # 事实类过短说明无效

    out2 = validate_suggested(
        {"doc_type": "past_proposal", "statement": "技术标共 8 章，重点：园区架构与运维。"},
        "标书.docx",
    )
    assert "共 8 章" in out2["statement"]  # 写法类保留结构概览


def test_extract_questions_validation(tmp_path, monkeypatch):
    """检索问题清洗：去空/去重/超长丢弃/上限 8 条；写法类整体丢弃。"""
    from app.knowledge.extract import validate_suggested

    qs = ["公司规模多大？", "  公司规模多大？  ", "", "持有哪些软件著作权？", "超" * 31,
          None, 123, "做过哪些医疗行业项目？"]
    qs += [f"问题{i}？" for i in range(8)]  # 凑过上限
    out = validate_suggested(
        {"doc_type": "contract_case",
         "statement": "XX 医院智慧后勤平台合同，金额 380 万元（第1页）。",
         "questions": qs},
        "合同.docx",
    )
    got = out.get("questions")
    assert got == ["公司规模多大？", "持有哪些软件著作权？", "做过哪些医疗行业项目？"] + [
        f"问题{i}？" for i in range(8)
    ][:8 - 3]  # 去空/去重/超长/非字符串丢弃后封顶 8 条

    # 写法类不生成检索问题（检索消费方是素材库）
    out2 = validate_suggested(
        {"doc_type": "past_proposal", "statement": "技术标共 8 章。",
         "questions": ["怎么写技术方案？"]},
        "标书.docx",
    )
    assert "questions" not in out2

    # 全无效时不落键（不产生空 questions）
    out3 = validate_suggested(
        {"doc_type": "financial_report", "statement": "净资产 8.2 亿元（第7页）。", "questions": ["", "  "]},
        "财报.docx",
    )
    assert "questions" not in out3


# ---------- 图片抽取（确定性零 LLM，仅供内容页查看） ----------

def _make_docx_with_images(tmp_path, n_images=3) -> bytes:
    """构造带嵌入图片的 docx（python-docx add_picture 真实图片）。"""
    import io

    import fitz
    from docx import Document

    doc = Document()
    doc.add_heading("产品介绍", 0)
    doc.add_paragraph("本产品手册介绍公司核心产品线的总体架构与关键能力特性说明。")
    for i in range(40):
        doc.add_paragraph(f"产品能力说明第{i + 1}条：平台核心功能模块的详细能力描述文本，"
                          "覆盖架构分层、部署形态与运维保障等多个维度的内容补充。")
    for i in range(n_images):
        # 每张图不同（docx 对相同图片去重成一个 media part）
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 600 + i, 400))
        pix.set_rect(pix.irect, (200 + i * 10, 120, 40))
        doc.add_paragraph(f"图{i + 1}说明：系统架构与部署拓扑的示意说明文本。")
        doc.add_picture(io.BytesIO(pix.tobytes("png")))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_docx_images_land_on_disk_only(tmp_path, monkeypatch):
    """图片抽取降级：落盘供内容页查看（清单可列），与素材彻底无关。"""
    _setup(tmp_path, monkeypatch, llm=_SmartLLM)

    class _ProfileLLM(_SmartLLM):
        def invoke(self, messages):
            prompt = messages[-1][1]
            if "文件名：产品介绍" in prompt:
                return _FakeResp(_extract_json(doc_type="technical_doc", statement="产品手册，含架构图。"))
            return super().invoke(messages)

    import langchain_deepseek

    monkeypatch.setattr(langchain_deepseek, "ChatDeepSeek", _ProfileLLM)
    kid = _upload("产品介绍.docx", _make_docx_with_images(tmp_path))
    item = run_ingest(kid)
    assert item["parse_status"] == "ready"
    assert item["doc_type"] == "technical_doc"
    # 图片落盘 + 清单可列（内容页折叠区数据源）
    from app.knowledge import images as images_mod
    from app.knowledge import store as kstore

    img_dir = kstore.kb_images_dir("产品介绍.docx")
    assert sum(1 for p in img_dir.iterdir() if p.is_file()) == 3  # 三张 600×400 图全过过滤
    assert len(images_mod.list_images("产品介绍.docx")) == 3


def test_small_images_filtered(tmp_path, monkeypatch):
    """装饰图过滤：短边 <200px 的图被丢弃。"""
    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    import io

    import fitz
    from docx import Document

    doc = Document()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 80, 40))  # 80×40 小图
    pix.set_rect(pix.irect, (10, 10, 10))
    doc.add_picture(io.BytesIO(pix.tobytes("png")))
    out = io.BytesIO()
    doc.save(out)
    src = store.kb_files_dir() / "装饰.docx"
    src.write_bytes(out.getvalue())
    written, skipped = images_mod.extract_images(src, "装饰.docx")
    assert written == 0 and skipped == 1  # 装饰图计入跳过


def test_unfriendly_image_converted_to_png():
    """TIFF/BMP（浏览器原生不解码）转 PNG；友好格式原样返回。"""
    from app.knowledge import images as images_mod

    # 最小 1×1 24bit BMP（fitz 不能写 BMP/TIFF，手拼字节）
    bmp = (
        b"BM" + (58).to_bytes(4, "little") + b"\x00\x00\x00\x00" + (54).to_bytes(4, "little")
        + (40).to_bytes(4, "little") + (1).to_bytes(4, "little") + (1).to_bytes(4, "little")
        + (1).to_bytes(2, "little") + (24).to_bytes(2, "little") + (0).to_bytes(4, "little")
        + (4).to_bytes(4, "little") + b"\x00" * 16
        + b"\xff\x00\x00\x00"
    )
    data, ext = images_mod.as_browser_friendly(bmp, ".bmp")
    assert ext == ".png" and data.startswith(b"\x89PNG")
    # 转码失败兜底：返回原始字节
    data2, ext2 = images_mod.as_browser_friendly(b"not-an-image", ".tif")
    assert data2 == b"not-an-image" and ext2 == ".tif"
    # 友好格式原样透传
    data3, ext3 = images_mod.as_browser_friendly(b"\x89PNG-rest", ".png")
    assert data3 == b"\x89PNG-rest" and ext3 == ".png"


def test_extract_docx_zip_guards(tmp_path, monkeypatch):
    """zip 解压护栏（2026-09-10 review）：单条声明超限跳过、累计声明超限早停、
    张数上限提前到读阶段——高压缩比 media 条目可在小压缩包里声明数 GB 解压量，
    必须读前用 file_size 拦，否则 OOM。阈值缩小驱动逻辑，真实阈值 50MB/200MB。"""
    import zipfile

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    src = store.kb_files_dir() / "护栏.docx"
    src.write_bytes(_make_docx_with_images(tmp_path))
    member, total, _count = (
        images_mod._MAX_MEMBER_BYTES,
        images_mod._MAX_TOTAL_BYTES,
        images_mod._MAX_IMAGES,
    )
    with zipfile.ZipFile(src) as z:
        sizes = [i.file_size for i in z.infolist() if i.filename.startswith("word/media/")]
    assert len(sizes) == 3

    # 单条上限：真实 PNG 全部「声明超限」→ 全跳过，空结果不打崩
    monkeypatch.setattr(images_mod, "_MAX_MEMBER_BYTES", 16)
    assert images_mod._extract_docx(src) == []

    # 累计上限：首条放行、第二条累计超限早停 → 恰好 1 条
    monkeypatch.setattr(images_mod, "_MAX_MEMBER_BYTES", member)
    monkeypatch.setattr(images_mod, "_MAX_TOTAL_BYTES", sizes[0] + 1)
    assert len(images_mod._extract_docx(src)) == 1

    # 张数上限在读阶段生效（不再先解压全部再截断）
    monkeypatch.setattr(images_mod, "_MAX_TOTAL_BYTES", total)
    monkeypatch.setattr(images_mod, "_MAX_IMAGES", 2)
    assert len(images_mod._extract_docx(src)) == 2


# ---------- PDF 整页渲染（2026-09-13：贴资料的单位是「那一页的复印件」） ----------

_PAGE_W, _PAGE_H = 400.0, 600.0


def _expected_px(pts: float) -> float:
    from app.knowledge import images as images_mod

    return pts * images_mod._RENDER_DPI / 72.0


def test_pdf_renders_whole_page_not_embedded_image(tmp_path, monkeypatch):
    """电子版证书（底图+文字浮层）：按嵌入图抽只剩空底框，整页渲染才拿到带字的页。

    构造一张覆盖整页的底图（原生仅 100×100）+ 叠在上面的文字层——旧逐图抽取会
    因短边 <200 被滤掉（0 张）；新路径渲染整页，尺寸=页面 150 DPI，文字在画里。
    """
    import fitz

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    doc = fitz.open()
    page = doc.new_page(width=_PAGE_W, height=_PAGE_H)
    base = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100))  # 小底图铺满整页
    page.insert_image(page.rect, stream=base.tobytes("png"))
    page.insert_text((60, 300), "ISO9001 CERTIFICATE 0350324Q30696R1M", fontsize=14)
    src = store.kb_files_dir() / "iso证书.pdf"
    doc.save(str(src))
    doc.close()

    written, skipped = images_mod.extract_images(src, "iso证书.pdf")
    assert written == 1 and skipped == 0
    img = store.kb_images_dir("iso证书.pdf") / "img_001.png"
    out = fitz.Pixmap(str(img))
    assert abs(out.width - _expected_px(_PAGE_W)) <= 2
    assert abs(out.height - _expected_px(_PAGE_H)) <= 2


def test_pdf_vector_only_page_still_renders(tmp_path, monkeypatch):
    """纯矢量/文字页（无任何嵌入图）——旧路径 0 张，渲染路径照出 1 张。"""
    import fitz

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    doc = fitz.open()
    page = doc.new_page(width=_PAGE_W, height=_PAGE_H)
    page.insert_text((40, 80), "VECTOR ONLY CERTIFICATE", fontsize=18)
    src = store.kb_files_dir() / "vector.pdf"
    doc.save(str(src))
    doc.close()

    written, skipped = images_mod.extract_images(src, "vector.pdf")
    assert written == 1 and skipped == 0
    assert len(images_mod.list_images("vector.pdf")) == 1


def test_pdf_blank_page_skipped_and_page_numbering(tmp_path, monkeypatch):
    """空白页跳过、文件名即页码（编号留空洞）——模型据页锚点可推回页图。"""
    import fitz

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    doc = fitz.open()
    p1 = doc.new_page(width=_PAGE_W, height=_PAGE_H)
    p1.insert_text((40, 80), "PAGE 1", fontsize=18)
    doc.new_page(width=_PAGE_W, height=_PAGE_H)  # 第 2 页空白
    p3 = doc.new_page(width=_PAGE_W, height=_PAGE_H)
    p3.insert_text((40, 80), "PAGE 3", fontsize=18)
    src = store.kb_files_dir() / "空白页.pdf"
    doc.save(str(src))
    doc.close()

    written, skipped = images_mod.extract_images(src, "空白页.pdf")
    assert written == 2 and skipped == 1
    names = [p.name for p in store.kb_images_dir("空白页.pdf").iterdir() if p.is_file()]
    assert sorted(names) == ["img_001.png", "img_003.png"]  # 无 img_002


def test_pdf_render_respects_max_images(tmp_path, monkeypatch):
    """张数上限：写满即停，长文档不白渲染后续页。"""
    import fitz

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page(width=_PAGE_W, height=_PAGE_H)
        page.insert_text((40, 80), f"PAGE {i + 1}", fontsize=18)
    src = store.kb_files_dir() / "四页.pdf"
    doc.save(str(src))
    doc.close()

    monkeypatch.setattr(images_mod, "_MAX_IMAGES", 2)
    written, _ = images_mod.extract_images(src, "四页.pdf")
    assert written == 2


def test_docx_skips_emf_and_zip_dir_entries(tmp_path, monkeypatch):
    """docx 侧两个坑：word/media/ 目录占位条目（0 字节）、EMF/WMF 矢量图（浏览器
    不认）——都不落盘，PNG 照旧。"""
    import zipfile

    from app.knowledge import images as images_mod

    _setup(tmp_path, monkeypatch)
    src = store.kb_files_dir() / "混合.docx"
    src.write_bytes(_make_docx_with_images(tmp_path))
    # 追加一个 EMF 媒体条目 + 一个目录占位条目（重写 zip）
    tmp_zip = src.with_suffix(".tmp.docx")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp_zip, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr("word/media/", b"")                     # 目录占位
        zout.writestr("word/media/vector1.emf", b"\x01\x00\x00\x00EMF")  # 矢量图
    tmp_zip.replace(src)

    raw = images_mod._extract_docx(src)
    assert all(not n.endswith(".emf") for _d, n in raw)  # EMF 未被抽出
    written, _ = images_mod.extract_images(src, "混合.docx")
    assert written == 3  # 三张真实 PNG 照旧落盘


# ---------- 锚点回文核对 + 自动确认（两主人模型） ----------

_CERT_MD = (
    "# 证书\n\nXX 建设银行数据中心项目 ISO9001 质量管理体系认证，证书编号 CN-001。\n"
    "有效期自 2023-05-01 至 2028-06-30。\n" + "体系覆盖软件开发与信息系统集成服务。" * 6
)


def _seq_llm(*responses):
    """按调用次序返回预设抽取 JSON（模拟重抽结果变化）；ChatDeepSeek 替身。
    队列在多次实例化间共享（run_extract 每次调用新建模型实例）。"""
    queue = list(responses)

    class _SeqModel:
        def __init__(self, **kwargs):
            pass

        def invoke(self, messages):
            return _FakeResp(queue.pop(0))

    return _SeqModel


def test_autoconfirm_on_check_pass(tmp_path, monkeypatch):
    """核对全过：自动确认 + business 副本带 auto 标记 + 确认版字段段可检索。"""
    _setup(tmp_path, monkeypatch, llm=_seq_llm(_extract_json(
        statement="ISO9001 质量管理体系认证，证书编号 CN-001（第1页），有效期至 2028-06-30（第1页）。",
        fields={"valid_until": {"value": "2028-06-30", "source": "第1页"},
                "client": {"value": "建设银行", "source": "第1页"}},
    )))
    kid = _upload("好证书.txt", _CERT_MD.encode())
    item = run_ingest(kid)
    assert item["review_status"] == "confirmed"
    biz = json.loads(item["business_metadata"])
    assert biz["confirmed_by"] == "auto"
    assert biz["fields"]["valid_until"]["value"] == "2028-06-30"
    assert json.loads(item["check_result"])["status"] == "pass"

    from app.knowledge import fts

    hits = db.kb_search_segments(fts.build_match_expr("2028"), limit=5)
    assert any(h.get("section_path") == "条目信息（人工确认）" and h.get("item_id") == kid for h in hits)


def test_check_fail_stays_pending(tmp_path, monkeypatch):
    """核对有败（日期不在原文）：留待确认 + 无 business + check_result 点名原因。"""
    _setup(tmp_path, monkeypatch, llm=_seq_llm(_extract_json(
        statement="ISO9001 质量管理体系认证证书。",
        fields={"valid_until": {"value": "2099-01-01", "source": "第1页"}},
    )))
    kid = _upload("坏证书.txt", _CERT_MD.encode())
    item = run_ingest(kid)
    assert item["review_status"] == "pending_review"
    assert not item["business_metadata"]
    results = json.loads(item["check_result"])["results"]
    assert any(not r["ok"] and "未在原文找到该日期" in r["detail"] for r in results)


def test_auto_demote_on_recheck(tmp_path, monkeypatch):
    """自动确认后重抽变坏：回落待确认并清掉旧自动副本（防陈旧盖章）。"""
    good = _extract_json(
        statement="ISO9001 质量管理体系认证证书，证书编号 CN-001（第1页）。",
        fields={"valid_until": {"value": "2028-06-30", "source": "第1页"}},
    )
    bad = _extract_json(
        statement="ISO9001 质量管理体系认证证书。",
        fields={"valid_until": {"value": "2099-01-01", "source": "第1页"}},
    )
    _setup(tmp_path, monkeypatch, llm=_seq_llm(good, bad))
    kid = _upload("复检证书.txt", _CERT_MD.encode())
    first = run_ingest(kid)
    assert first["review_status"] == "confirmed"
    second = run_ingest(kid)
    assert second["review_status"] == "pending_review"
    assert not second["business_metadata"]
    assert json.loads(second["check_result"])["status"] == "fail"


def test_human_confirm_immune_to_recheck(tmp_path, monkeypatch):
    """人工确认过的条目：重抽只更新 suggested/check_result，business 与确认状态不动。"""
    good = _extract_json(
        statement="ISO9001 质量管理体系认证证书，证书编号 CN-001（第1页）。",
        fields={"valid_until": {"value": "2028-06-30", "source": "第1页"}},
    )
    bad = _extract_json(
        statement="ISO9001 质量管理体系认证证书。",
        fields={"valid_until": {"value": "2099-01-01", "source": "第1页"}},
    )
    _setup(tmp_path, monkeypatch, llm=_seq_llm(good, bad))
    kid = _upload("人工证书.txt", _CERT_MD.encode())
    assert run_ingest(kid)["review_status"] == "confirmed"
    # 模拟人工确认（PUT /metadata 的落库形状：business 无 auto 标记）
    human_biz = {"doc_type": "qualification_certificate",
                 "fields": {"valid_until": {"value": "2030-12-31"}}}
    db.kb_update_item(kid, review_status="confirmed",
                      business_metadata=json.dumps(human_biz, ensure_ascii=False))
    item = run_ingest(kid)
    assert item["review_status"] == "confirmed"
    assert json.loads(item["business_metadata"]) == human_biz  # 人工值原样
    assert json.loads(item["check_result"])["status"] == "fail"  # 核对仍更新展示
