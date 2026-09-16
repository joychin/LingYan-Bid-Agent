"""docx 正文工具族（tender-body docx 直出路线，2026-09-06 定案）。

正文节文件 = 任务 work/body/ 下每节一个 .docx；模型不碰二进制——本工具族把
docx 翻译成文本世界（读视图/编号寻址），写入全由程序机械完成：

- docx_section_create：建节文件（标题 + 可选初始内容；已存在不覆盖，
  重写传 replace=true——旧版自动入恢复点栈）。初始内容两条通道：paragraphs
  纯文字换行分段（旧），body 块序列 JSON 段落+表格混排一次成形（2026-09-14
  表格通道批，推荐——人员配置/里程碑/对比类内容用表格不用流水句）
- docx_diagram_insert：程序生成图示（表格拼装，2026-09-14 批）——layered
  分层/组织架构、gantt 甘特、radial 中心辐射；模型只给语义拓扑，灰阶版式
  全程序机械生成；图示承载承诺数据故走表格不走图片（可编辑、validate 可读）
- docx_section_read：序列化视图（P1..Pn 段落+样式+图片标记，T1..Tm 表格逐行
  展开 R行 C列= 格文本——合并格标「同左/同上」，格坐标即修订寻址）
- docx_material_inject：素材块元素级注入——从素材 docx 原件把块区间对应的
  元素（段落/表格连合并单元格、图片连关系）整体拷进目标节文件，零转写保真；
  寻址 = 素材解析时落盘的 element_map（md 行号区间 → 原件 body 子元素）；
  元素引用的样式与自动编号定义随迁（含 basedOn/link 链；编号定义按内容判定
  沿用、id 冲突时重映射并改写引用——numbering id 只是文档内部门牌号）
- docx_source_inject：招标原件拷贝——任务 sources/ 的 docx 按「整文件或 md
  行号区间」元素级拷进正文节；现场跑解析注册表拿 element_lines 定位（确定性、
  不依赖任务侧解析落盘）；格式跟随与格式件（投标函/一览表等模板填充类）场景，
  与素材注入共用同一套迁移引擎
- docx_image_insert：单图插入——证书复印件/扫描件/截图等独立图片（docx 直出
  管线的放图通道；素材块/招标件内的图走注入、不经此工具）。图源=知识库图片
  （PDF 整页渲染/ docx 内嵌图）/任务 sources 图片/PDF 原件按页现场渲染；全宽居中、段落标记+内容双插入修订
  （与插段同构——拒绝修订=整段含图消失）
- docx_section_revise：定向修订，全部落成 Word 原生修订标记（w:ins/w:del，
  author=灵燕智能）——用户在 Word 审阅界面逐条接受/拒绝；正文段落按
  P 序号、表格单元格按 table/row/col（replace 改旧值、fill 填空格）；落盘前
  程序做「拒绝全部修订后文本与修订前逐字一致」的自校验（标记写坏的机械防线，
  覆盖正文与表格单元格段落）
- docx_assemble_volume：整本合册——按投标目录树序把各节 docx 合并成每册
  一个整本文件（容器节点发章标题——按树序自动编号（第X章/1.1，格式取目录
  产物 numbering 字段）、一级章前分页、页脚页码；模板填充类叶子
  产出节文件即按树序并入、未产出按附件对待不占整本位；「目录」节点机械
  生成目录页（Word 目录域+缓存清单，2026-09-14 批）且目录前前置页不占章号、
  空容器章标题抑制；**整本=交付态**：
  并入时按「接受全部修订」压平（节文件保留修订供审阅，改内容回节级改再
  重合册）、拷入的树外标题段摘出大纲层级（导航窗格只剩章节骨架）、
  手写目录页与实收章节机械对账不符点名；整本是派生产物，内容真值在节文件，
  重新合册覆盖）

寻址纪律：段落序号以 docx_section_read 视图为准（body 直属段落，不含表格内
段落）；revise 提交 (序号, 期望原文 find) 双重校验防漂移，多条段落编辑按序号
从大到小应用（防插入/删除位移；表格格编辑不改变段落序号，与段落批互不影响，
另批处理）。表格的插行/删行与嵌套表格不支持（在 Word 中改）。

设计参考：document-skills 插件（Proprietary——只借思想不搬代码：修订标记
形态、组合寻址、图片三件套、生成后机械校验）；图片关系迁移走
get_or_add_image（按内容去重、自动避让 partname）。
"""

from __future__ import annotations

import functools
import hashlib
import json
import re
import tempfile
import threading
import unicodedata
import uuid
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt
from langchain_core.tools import tool
from lxml import etree

from .. import db, publish, render_queue, runctx
from ..artifact_store import RESTORE_KEEP, sources_dir, work_dir
from ..config import workspace_dir
from ..knowledge import materials_lib
from ..parse import convert as parse_convert
from ..parse import pdfium_kit
from ..parse.pdf import render_page_png
from . import body_contract

_AUTHOR = "灵燕智能"
_COMMENT_AUTHOR = "灵燕智能"  # 批注作者（Word 审阅侧栏可见；待办批注与修订标记分属两套体系）
_VIEW_TEXT_LIMIT = 800  # 视图单段截断（修订需精确文本，超长段提示去 Word 处理）
_CELL_TEXT_LIMIT = 60  # 视图单元格截断（填空场景格文本短；长格内容去 Word 看）

# 表格通道批（2026-09-14）：程序生成表格/图示的灰阶与上限——视觉约定吸收自
# 头部产品样例（域头/表头深灰、成员/数据格浅灰、留白格无底纹），打印黑白安全。
_DIAG_HEAD_FILL = "D9D9D9"
_DIAG_ITEM_FILL = "F2F2F2"
_CAPTION_STYLE_NAME = "Tender Caption"  # 图注/表题样式（基准模板提供；缺失回落正文样式）
_MAX_TABLE_COLS = 8
_MAX_TABLE_ROWS = 60
_MAX_CELL_CHARS = 200


def _task_id() -> str | None:
    ctx = runctx.current_run()
    if ctx is None:
        return None
    conv = db.get_conversation(ctx.conversation_id) if ctx.conversation_id else None
    return (conv or {}).get("task_id")


def _dest_path(task_id: str, rel: str, *, must_exist: bool) -> tuple[Path, str]:
    """正文 docx 路径解析：只放行 work/body/ 下、自动补 .docx、防目录穿越。

    返回 (绝对路径, 相对 work/ 的展示路径)；违规抛 ValueError（工具层转人话）。
    """
    rel = (rel or "").strip().replace("\\", "/").lstrip("/")
    if not rel:
        raise ValueError("未指定文件路径")
    if not rel.endswith(".docx"):
        rel += ".docx"
    if not rel.startswith("body/"):
        raise ValueError("docx 正文只允许落在任务 work/body/ 目录下")
    root = work_dir(task_id).resolve()
    p = (root / rel).resolve()
    body_root = root / "body"
    # resolve 后校验 body 目录在祖先链——词法前缀检查挡不住 body/../x.docx
    # 这类单级穿越（折叠后落 work/ 根），也不能漏掉 symlink 指向外部的情况
    if body_root != p and body_root not in p.parents:
        raise ValueError("路径越界（docx 正文只允许落在任务 work/body/ 目录下）")
    if must_exist and not p.is_file():
        raise ValueError(f"文件不存在：work/{rel}（先用 docx_section_create 创建）")
    return p, rel


def _rotate_restore_point(dst: Path) -> None:
    """旧文件入恢复点栈（<文件名>.restorepoints/NNNN.bak、留 RESTORE_KEEP 个轮换
    ——工作台编辑恢复点同款目录形态；.bak 不匹配 rglob("*.docx")，不进面板列表与
    「本轮文件」）。栈深单一真值=artifact_store.RESTORE_KEEP。"""
    d = dst.with_name(dst.name + ".restorepoints")
    d.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d%H%M%S%f")
    dst.rename(d / f"{ts}.bak")
    points = sorted(p for p in d.iterdir() if p.is_file() and p.suffix == ".bak")
    for old in points[:-RESTORE_KEEP]:
        old.unlink()


# ---------- 同文件并发写防线（2026-09-13 事故批） ----------
#
# 动因：ToolNode 对同一消息的多个工具调用真并行（max_concurrency=8），模型一turn
# 连发 N 条 docx_comment_add/revise 打同一节文件时，N 个线程各自
# Document(开)→改→save(path) 原地覆盖——zip 写坏（BadZipFile）或后来者整存覆盖
# 丢更新。事故链：8 连批注打坏「资质证书.docx」→ 写手 read_file 读证书 PNG「看
# 一眼」→ base64 内联 156 万字符 → 下一次模型调用撞 1M 上下文上限（读侧拦截见
# fs_guard 同批）。两层各治一半，均为用户不可见 plumbing（原子落盘同款，不违
# 「跨 run 无锁」铁则——那只管后台协调，此处是同进程文件写完整性）：
# - 按路径 threading.Lock 串行化同文件「开→改→存」整段：并行调用不丢更新
#   （8 条批注一条不丢）；锁条目只增不删（回收有竞态；量级=节文件数，可忽略）；
# - _atomic_save：uuid 后缀 tmp 同目录落盘 + os.replace（artifact_store 同款，
#   固定 .tmp 名并发互踩）——读者（section_read/validate_body/合册读节）永远
#   看到完整旧版或新版，不再有撕裂 zip。恢复点 rename 语义不变。
_PATH_LOCKS: dict[str, threading.Lock] = {}
_PATH_LOCKS_GUARD = threading.Lock()


@contextmanager
def _docx_path_lock(path: Path):
    """同目标文件的改写互斥（按解析后绝对路径）。"""
    with _PATH_LOCKS_GUARD:
        lock = _PATH_LOCKS.setdefault(str(path), threading.Lock())
    lock.acquire()
    try:
        yield
    finally:
        lock.release()


def _atomic_save(doc, path: Path) -> None:
    """docx 落盘原子替换：uuid 后缀 tmp + os.replace，异常清理残件。"""
    tmp = path.with_name(f"{path.name}.{uuid.uuid4().hex[:8]}.tmp")
    try:
        doc.save(tmp)
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


# source_inject 的元素映射缓存：同原件一个 run 拷多个格式件（投标函/授权书/
# 一览表…）不重复整本解析——几百页招标文件每次解析秒级；(路径, mtime) 为键，
# 原件被替换（mtime 变）即自动失效重解析；FIFO 限 8 条防膨胀
_ELEMENT_LINES_CACHE: dict[tuple[str, float], list] = {}
_ELEMENT_LINES_CACHE_MAX = 8


def _source_element_lines(src_path: Path) -> list:
    if not src_path.is_file():
        return []
    key = (str(src_path), src_path.stat().st_mtime)
    hit = _ELEMENT_LINES_CACHE.get(key)
    if hit is not None:
        return hit
    res = parse_convert(src_path)
    lines = list(res.info.get("element_lines") or [])
    if len(_ELEMENT_LINES_CACHE) >= _ELEMENT_LINES_CACHE_MAX:
        _ELEMENT_LINES_CACHE.pop(next(iter(_ELEMENT_LINES_CACHE)))
    _ELEMENT_LINES_CACHE[key] = lines
    return lines


# ---------- 修订标记（tracked changes）核心 ----------

def _now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _next_rev_id(doc: Document) -> int:
    ids = [
        int(e.get(qn("w:id")))
        for e in doc.element.body.iter()
        if e.tag in (qn("w:ins"), qn("w:del")) and (e.get(qn("w:id")) or "").isdigit()
    ]
    return (max(ids) + 1) if ids else 1


def _has_tracked(p_el) -> bool:
    # iter 覆盖全部后代：段内 w:ins/w:del 与段落标记修订（pPr/rPr 内）一并检出
    return any(e.tag in (qn("w:ins"), qn("w:del")) for e in p_el.iter())


def _text_runs(p_el) -> list:
    """参与文本定位的 run 流：直属 w:r + w:ins 内 w:r（顺序），跳过 w:del 子树。

    同段多轮修订需要穿过本轮已落的修订标记定位（嵌套 ins>del 合法，Word
    显示为对插入内容的再修订）；定位基准 = 接受视角文本。
    """
    out: list = []

    def walk(node):
        for ch in node:
            if ch.tag == qn("w:del"):
                continue
            if ch.tag == qn("w:r"):
                out.append(ch)
                continue
            walk(ch)

    walk(p_el)
    return out


def _run_text(r) -> str:
    return "".join(t.text or "" for t in r.findall(qn("w:t")))


def _set_run_text(r, text: str) -> None:
    """run 文本重设：保留首个 w:t（带 preserve），其余 w:t 清空。"""
    ts = r.findall(qn("w:t"))
    if not ts:
        t = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve">{escape(text)}</w:t>')
        r.append(t)
        return
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for extra in ts[1:]:
        extra.text = ""


def _split_run_head(run, keep: int):
    """run 截为前 keep 字符，剩余文本生成新 run 插在其后（拷贝全部属性）。"""
    text = _run_text(run)
    tail = text[keep:]
    _set_run_text(run, text[:keep])
    if not tail:
        return run
    new_r = deepcopy(run)
    _set_run_text(new_r, tail)
    run.addnext(new_r)
    return new_r


def _tracked_replace(para, old: str, new: str, rev_id: int) -> None:
    """段内定向替换（接受视角定位，可穿过已有修订标记）：每个匹配 run 在原
    位置包独立 w:del（w:t→w:delText；ins 内 run 变 ins>del 合法嵌套），新文本
    以首匹配 run 的格式包 w:ins 插在末个 del 之后。"""
    p_el = para._p
    runs = _text_runs(p_el)
    full = "".join(_run_text(r) for r in runs)
    pos = full.find(old)
    if pos < 0:
        raise ValueError(f"段落内未找到期望文本「{old[:60]}」")
    end = pos + len(old)
    # 字符偏移 → run 序列，首尾拆分对齐匹配边界
    off = 0
    match_runs: list = []
    for r in runs:
        rlen = len(_run_text(r))
        if rlen == 0 or off + rlen <= pos or off >= end:
            off += rlen
            continue
        if off < pos:
            r = _split_run_head(r, pos - off)
            off = pos
        run_end = off + len(_run_text(r))
        if run_end > end:
            _split_run_head(r, end - off)
        match_runs.append(r)
        off = run_end
    if not match_runs:
        raise ValueError("匹配区间为空")
    date = _now_iso()
    # per-run 原位包 del：匹配可能横跨「已有 w:ins 内容 + 原文」（二轮修订常态），
    # run 的父级不一——每个 run 在自己原位置包独立 w:del（ins 内变 ins>del 合法
    # 嵌套、正文留 p>del），拒绝视角原文完整；若整批共用一个 del 挂在首 run
    # 位置，会把正文 run 挪进 ins 子树、拒绝修订后原文凭空消失（自校验必拒）。
    last_del = None
    for r in match_runs:
        for t in r.findall(qn("w:t")):
            t.tag = qn("w:delText")
        del_el = parse_xml(f'<w:del {nsdecls("w")} w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}"/>')
        r.addprevious(del_el)
        del_el.append(r)  # move 出原位置、原父级原顺序
        last_del = del_el
    rpr = match_runs[0].find(qn("w:rPr"))
    rpr_xml = rpr.xml if rpr is not None else ""
    ins_el = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="{rev_id + 1}" w:author="{_AUTHOR}" w:date="{date}">'
        f'<w:r>{rpr_xml}<w:t xml:space="preserve">{escape(new)}</w:t></w:r></w:ins>'
    )
    last_del.addnext(ins_el)


def _tracked_insert_after(anchor_p_el, text: str, rev_id: int, style_id: str | None = None):
    """新段落插在锚段（lxml 元素）之后，段落标记与内容均为插入修订（拒绝修订=
    整段消失）。返回新段元素——同段连续 insert_after 以返回值为下一次的锚
    （否则 addnext 紧跟定位段，同批多条会倒序）。style_id 非空时挂该样式
    （w:pStyle 挂 styleId：正文段=Tender Body，封面行=Tender Cover/Sub，
    由调用方按样式名解析，随锚段所在文档的模板有无自适应）。"""
    date = _now_iso()
    pstyle = f'<w:pStyle w:val="{style_id}"/>' if style_id else ""
    new_p = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr>{pstyle}<w:rPr><w:ins w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}"/></w:rPr></w:pPr>'
        f'<w:ins w:id="{rev_id + 1}" w:author="{_AUTHOR}" w:date="{date}">'
        f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:ins></w:p>'
    )
    anchor_p_el.addnext(new_p)
    return new_p


def _tracked_fill(para, text: str, rev_id: int) -> None:
    """空格子填入：文本以 w:ins 追加进段落（拒绝修订=回到空格子）。

    只用于接受视角为空的表格单元格——段落本身存在（python-docx 建格自带空段），
    追加插入修订即可；沿用格子既有段落属性（对齐等）。
    """
    date = _now_iso()
    ins_el = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}">'
        f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r></w:ins>'
    )
    para._p.append(ins_el)


def _tracked_delete(para, rev_id: int) -> None:
    """段落删除修订：全部 run（含 w:ins 内，嵌套=删除先前的插入）文本转
    删除文本，段落标记本身也标删除（接受修订时段落真正消失；pPr/rPr 的
    w:del 置于最前——CT_ParaRPr 顺序）。"""
    p_el = para._p
    date = _now_iso()
    for r in list(p_el.iter(qn("w:r"))):  # list 快照：边遍历边移动
        del_el = parse_xml(f'<w:del {nsdecls("w")} w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}"/>')
        r.addprevious(del_el)
        del_el.append(r)
        for t in r.findall(qn("w:t")):
            t.tag = qn("w:delText")
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
        p_el.insert(0, ppr)
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
        ppr.insert(0, rpr)
    rpr.insert(0, parse_xml(
        f'<w:del {nsdecls("w")} w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}"/>'
    ))


def _collect_rejected(node, parts: list) -> None:
    """「拒绝全部修订」视角收文本：跳过 w:ins 子树；w:t 与 w:delText 均计入。"""
    for ch in node:
        if ch.tag == qn("w:ins"):
            continue
        if ch.tag in (qn("w:t"), qn("w:delText")):
            parts.append(ch.text or "")
        _collect_rejected(ch, parts)


def _flatten_rejected(doc: Document) -> str:
    """「拒绝全部修订」视角收全文（正文段落 + 表格单元格段落，顺序稳定）。

    表格部分按 表格序→行序→格序 展开——格编辑写坏修订标记（拒绝视角文本
    变化）与正文段落一样触发落盘前自校验拒绝。合并格在 row.cells 里重复
    出现，重复收集对称存在于修订前后两侧，不影响相等性比较。
    """
    def _para_rejected(p_el) -> str | None:
        ppr = p_el.find(qn("w:pPr"))
        rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
        if rpr is not None and rpr.find(qn("w:ins")) is not None:
            return None  # 段落标记为插入——拒绝修订后整段不存在
        parts: list[str] = []
        _collect_rejected(p_el, parts)
        return "".join(parts)

    out: list[str] = []
    for p in doc.paragraphs:
        text = _para_rejected(p._p)
        if text is not None:
            out.append(text)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    text = _para_rejected(p._p)
                    if text is not None:
                        out.append(text)
    return "\n".join(out)


def _collect_accepted(node, parts: list) -> None:
    """「接受全部修订」视角收文本：w:ins 内容计入、w:del 子树跳过、w:delText 不计。"""
    for ch in node:
        if ch.tag == qn("w:del"):
            continue
        if ch.tag == qn("w:t"):
            parts.append(ch.text or "")
        _collect_accepted(ch, parts)


def _accepted_text(p_el) -> str:
    parts: list[str] = []
    _collect_accepted(p_el, parts)
    return "".join(parts)


# ---------- 合册交付态：接受修订压平 + 树外标题摘大纲 + 目录页对账 ----------
# 2026-09-12 目录乱象批：实测整本带 600+ 处修订标记——未接受修订前新旧标题
# 成对交错（用户看到「目录乱」的主诉之一）；131 个节内小标题/素材自带标题
# 用 Heading 样式进大纲，导航窗格一锅粥；目录页静态条目与实收章节脱节。
# 修法全在合册侧（整本是派生产物，节文件层的修订审阅入口不动）。


def _accept_revisions_inplace(el) -> bool:
    """元素级「接受全部修订」压平：w:del 子树丢弃（w:delText 随之消失）、
    w:ins 剥壳内容按原序上提、段落标记修订（rPr/trPr 内的 ins/del）与格式
    变更记录（*Change）清除、表格行删除标记丢弃整行。

    返回 False=元素整体消失（段落标记删除修订且内容已删净——_tracked_delete
    形态，接受修订后整段不存在）。批注锚点（commentRangeStart/End/
    commentReference）不是修订，原样保留。
    """
    # 表格行删除修订：接受=整行消失（先于段落压平，行内段落不再处理）
    for tr in list(el.findall(".//" + qn("w:tr"))):
        trpr = tr.find(qn("w:trPr"))
        if trpr is not None and trpr.find(qn("w:del")) is not None:
            tr.getparent().remove(tr)
    # 先记录段落标记删除的段（rPr/w:del 会在下面的清除轮里被摘掉）
    paras = ([el] if el.tag == qn("w:p") else []) + el.findall(".//" + qn("w:p"))
    mark_deleted = [
        p for p in paras
        if (ppr := p.find(qn("w:pPr"))) is not None
        and (rpr := ppr.find(qn("w:rPr"))) is not None
        and rpr.find(qn("w:del")) is not None
    ]
    # w:del 子树丢弃：内容级删除（含「删除先前插入」的 ins>del 嵌套，整树带走）
    for node in list(el.iter(qn("w:del"))):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
    # w:ins 剥壳：子元素上提到父原位（run 级插入；rPr 里的空段落标记壳同样摘除）
    for node in list(el.iter(qn("w:ins"))):
        parent = node.getparent()
        if parent is None:
            continue
        index = list(parent).index(node)
        for ch in list(node):
            parent.insert(index, ch)
            index += 1
        parent.remove(node)
    # 残余标记：行级插入/删除标记（trPr）与格式变更记录——保留当前格式即接受
    for tag in ("w:ins", "w:del", "w:rPrChange", "w:pPrChange",
                "w:tblPrChange", "w:trPrChange", "w:sectPrChange"):
        for node in list(el.iter(qn(tag))):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    # 段落标记删除且内容已删净 → 整段消失；仍有内容=罕见形态，保守保留（近似）
    for p in mark_deleted:
        if _accepted_text(p).strip():
            continue
        if p is el:
            return False  # 顶层元素自身消失，调用方丢弃
        parent = p.getparent()
        if parent is not None:
            parent.remove(p)
            # OOXML 硬要求 w:tc 至少一个块级子元素：格内唯一段被删净时留空段，
            # 否则病态素材会让整本被 Word 判损坏（正常 Word 编辑产生不了此形态）
            if parent.tag == qn("w:tc") and not any(
                ch.tag in (qn("w:p"), qn("w:tbl")) for ch in parent
            ):
                parent.append(parse_xml(f"<w:p {nsdecls('w')}/>"))
    return True


def _heading_style_ids(doc: Document) -> set[str]:
    """样式表里大纲标题样式的 styleId 集：段落样式名含 heading/标题、或样式
    定义自带 outlineLvl（素材迁入的自定义标题样式多带）。Word 导航窗格与
    自动目录按样式的大纲层级取条目。"""
    ids: set[str] = set()
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH or not style.style_id:
            continue
        el = style.element
        name_el = el.find(qn("w:name"))
        name = (name_el.get(qn("w:val")) or "").lower() if name_el is not None else ""
        ppr = el.find(qn("w:pPr"))
        has_outline = ppr is not None and ppr.find(qn("w:outlineLvl")) is not None
        if "heading" in name or "标题" in name or has_outline:
            ids.add(style.style_id)
    return ids


def _demote_extra_headings(doc: Document, elements) -> int:
    """树外标题摘出大纲：拷入元素里会进大纲的段落显式设 outlineLvl=9（正文级）
    ——整本导航窗格/自动目录只剩合册器按树发的章节骨架（发标题的树对账已定，
    拷入面全是树外内容：节内小标题、素材自带章标题等）。
    判据两形态（2026-09-13 批扩）：命中标题样式的段落（pStyle∈标题样式集），与
    直接挂 outlineLvl<9 的段落（无标题样式、Word 里手点大纲级别的形态——招标
    格式件常见；`_strip_copy_residue` 已在拷贝入口剥，此处兜底存量/其他路径）。
    只改段落大纲层级，样式不动——视觉（字体字号缩进）零变化。"""
    if not elements:
        return 0
    heading_ids = _heading_style_ids(doc)
    if not heading_ids:
        return 0
    demoted = 0
    for el in elements:
        for p in ([el] if el.tag == qn("w:p") else []) + el.findall(".//" + qn("w:p")):
            ppr = p.find(qn("w:pPr"))
            pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
            style_heading = (
                pstyle is not None and (pstyle.get(qn("w:val")) or "") in heading_ids
            )
            direct = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
            direct_lvl = 9
            if direct is not None:
                try:
                    direct_lvl = int(direct.get(qn("w:val")) or "9")
                except ValueError:
                    direct_lvl = 9
            if not style_heading and direct_lvl >= 9:
                continue
            if ppr is None:
                ppr = parse_xml(f"<w:pPr {nsdecls('w')}/>")
                p.insert(0, ppr)
            for old in ppr.findall(qn("w:outlineLvl")):
                ppr.remove(old)
            node = parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="9"/>')
            anchor = None
            for tail in (qn("w:rPr"), qn("w:sectPr"), qn("w:pPrChange")):
                anchor = ppr.find(tail)
                if anchor is not None:
                    break
            if anchor is not None:  # CT_PPr 顺序：outlineLvl 在 rPr/sectPr 之前
                anchor.addprevious(node)
            else:
                ppr.append(node)
            demoted += 1
    return demoted


_TOC_TAIL_PAGE = re.compile(r"[.·．…]{2,}\s*\d+\s*$")
_TOC_NUMBER_PREFIXES = (
    r"^第[一二三四五六七八九十百\d]+章\s*",
    r"^[一二三四五六七八九十]+、",
    r"^（[一二三四五六七八九十]+）",
    r"^\d{1,2}(?:\.\d+)*\s*",
)


def _toc_normalize(text: str) -> str:
    """目录页条目/章节标题的比对键：剥编号前缀（与 _SELF_NUMBERED 同形态）与
    点线页码尾巴、NFKC 归一（全半角/大小写）、去空白——条目写法差异不误报。"""
    s = _TOC_TAIL_PAGE.sub("", text.strip())
    for pat in _TOC_NUMBER_PREFIXES:
        s = re.sub(pat, "", s)
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s)).lower()


def _toc_match(key: str, keys: list[str]) -> bool:
    """目录页条目与章节标题的宽松匹配：归一键相等、或双向前缀（短方 ≥4 字，
    容忍条目简写/带尾巴的写法差异）。"""
    return any(
        key == k or (min(len(key), len(k)) >= 4 and (key.startswith(k) or k.startswith(key)))
        for k in keys
    )


# ---------- 视图与图片迁移 ----------

def _comment_texts(doc: Document) -> dict[int, str]:
    """comment id → 批注文本（无批注部件返回空——读视图/清点不因批注缺失报错）。

    python-docx 的 Document.comments 在无部件时懒创建空部件（不落盘无副作用）。
    """
    try:
        return {c.comment_id: (c.text or "").strip() for c in doc.comments if (c.text or "").strip()}
    except Exception:
        return {}


def view_lines(doc: Document) -> list[str]:
    view = []
    comment_texts = _comment_texts(doc)
    paras = doc.paragraphs
    for i, para in enumerate(paras, 1):
        style = para.style.name if para.style is not None else "?"
        imgs = len(para._p.findall(".//" + qn("a:blip")))
        bits = f"[P{i}]（{style}）"
        if imgs:
            bits += f"〔图×{imgs}〕"
        ppr = para._p.find(qn("w:pPr"))
        rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
        para_deleted = rpr is not None and rpr.find(qn("w:del")) is not None
        if para_deleted:
            bits += "〔删除修订：接受后此段消失〕"
        elif _has_tracked(para._p):
            bits += "〔已含修订标记〕"
        # 本段锚定的待办批注（commentRangeStart 是段落直接子元素；id 对不上批注部件=陈旧引用跳过）
        notes = [
            t for m in para._p.findall(qn("w:commentRangeStart"))
            if (raw := m.get(qn("w:id")) or "").isdigit() and (t := comment_texts.get(int(raw)))
        ]
        if notes:
            shown = "；".join(n[:40] + ("…" if len(n) > 40 else "") for n in notes)
            bits += f"〔批注：{shown}〕"
        # python-docx 的 .text 不含修订标记内容——视图按「接受全部修订后」文本展示
        text = _accepted_text(para._p) or "（空段落）"
        if len(text) > _VIEW_TEXT_LIMIT:
            text = text[:_VIEW_TEXT_LIMIT] + f"…（截断，全段 {len(text)} 字，请在 Word 中查看）"
        view.append(bits + text)
    for ti, t in enumerate(doc.tables, 1):
        view.append(
            f"[T{ti}] 表格 {len(t.rows)}行×{len(t.columns)}列"
            "（单元格按 R行C列 寻址：docx_section_revise 传 table/row/col；合并格标「同左/同上」）"
        )
        prev_cells: list = []
        for ri, row in enumerate(t.rows, 1):
            bits: list[str] = []
            last_tc = None
            for ci, cell in enumerate(row.cells, 1):
                tc = cell._tc
                above = prev_cells[ci - 1] if ci <= len(prev_cells) else None
                if tc is last_tc:
                    bits.append(f"C{ci}=（同左合并格）")
                elif above is not None and tc is above._tc:
                    bits.append(f"C{ci}=（同上合并格）")
                else:
                    text = " / ".join(
                        s for p in cell.paragraphs if (s := _accepted_text(p._p).strip())
                    ) or "（空）"
                    if len(text) > _CELL_TEXT_LIMIT:
                        text = text[:_CELL_TEXT_LIMIT] + "…（截断）"
                    bits.append(f"C{ci}={text}")
                last_tc = tc
            view.append(f"[T{ti}] R{ri}：" + " ".join(bits))
            prev_cells = list(row.cells)
    head = f"共 {len(paras)} 个段落、{len(doc.tables)} 个表格（段落按 P 编号，表格按 T 编号）"
    if comment_texts:
        head += f"，待办批注 {len(comment_texts)} 条"
    view.insert(0, head)
    return view


def _migrate_images(src_doc: Document, dst_doc: Document, element) -> int:
    """把元素引用的源文档图片登记进目标文档并改写引用（按内容去重、partname 自动避让）。"""
    moved = 0
    for blip in element.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            image_part = src_doc.part.related_parts[rid]
            blob = image_part.blob
        except KeyError:
            continue
        try:
            new_rid, _image = dst_doc.part.get_or_add_image(BytesIO(blob))
        except Exception:
            continue  # 图片部件不可读：跳过（保留旧引用=裂图）——跨包挂外部 part 会损坏序列化
        blip.set(qn("r:embed"), new_rid)
        moved += 1
    return moved


def _sig_shingles(text: str) -> set[str]:
    """强归一化（只留中文与字母数字）的 10 字 shingle 集合：重复注入探测用，
    格式符号（表格竖线/视图标签/列表前缀）不参与比对。"""
    s = re.sub(r"[^\u4e00-\u9fffa-zA-Z0-9]", "", text)
    return {s[i : i + 10] for i in range(max(0, len(s) - 9))}


def _strip_inner_sectpr(el) -> None:
    """剥段落 pPr 内的分节符：素材/节文件自带的分节属性（纸向/页边距/页码
    重起）随元素拷贝会中途生效、突变节文件版式（2026-09-08 实证：横向页设置
    进入节文件）。文档级分节只归节文件自身的文末 sectPr（不在拷贝面）。"""
    ppr = el.find(qn("w:pPr"))
    if ppr is None:
        return
    for sect in ppr.findall(qn("w:sectPr")):
        ppr.remove(sect)


def _strip_copy_residue(el) -> None:
    """拷入卫生（2026-09-13 目录乱号批）：剥段落直接挂的大纲级别与 Word 内部书签。

    大纲级别是导航元数据不是版式——「保真拷贝」不含它：招标格式件/素材段落带着
    直挂 outlineLvl 进来，轻则导航出现树外条目，重则写手在该段里改写文本后整段
    正文混进大纲（实测 P469 形态）。节内小标题要走 Heading 样式（视图可见）。
    `_` 前缀书签（_Toc/_Ref…）指向源文档自己的目录/交叉引用域，我们的文档没有
    这些域——死引用留着无意义，同段拷进多个节还会造成书签 id 重复（Word 严格
    校验可能弹修复）；普通书签保留。"""
    paras = [el] if el.tag == qn("w:p") else []
    paras += el.findall(".//" + qn("w:p"))
    for p in paras:
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            continue
        for ol in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(ol)
    starts = [
        b for b in el.iter(qn("w:bookmarkStart"))
        if (b.get(qn("w:name")) or "").startswith("_")
    ]
    if not starts:
        return
    dead_ids = {b.get(qn("w:id")) for b in starts}
    for b in starts:
        parent = b.getparent()
        if parent is not None:
            parent.remove(b)
    for b in list(el.iter(qn("w:bookmarkEnd"))):
        # 区间跨元素时 End 可能不在本元素内：留在原处成孤儿 End，Word 忽略无
        # 起点的终点；范围内的都清掉，不留半个书签
        if b.get(qn("w:id")) in dead_ids:
            parent = b.getparent()
            if parent is not None:
                parent.remove(b)


_BASE_TEMPLATE = Path(__file__).resolve().parent.parent / "resources" / "tender_base_template.docx"
_BODY_STYLE_NAME = "Tender Body"  # 版式文件自定义正文样式（1.5 倍行距+首行缩进 2 字符）
_COVER_NODE_NAME = "封面"  # 树首封面节点的约定名（合册按清洗后标题识别，tender-outline 定下）
_TOC_NODE_NAME = "目录"  # 目录页节点的约定名（同为 tender-outline 结构约定；内容由合册机械生成）
TEMPLATE_SETTING_KEY = "docx_template"  # app_settings 键：默认版式的用户文件名（空=内置基准）
BUILTIN_TEMPLATE_KEY = "__builtin__"  # 内置版式寻址键（版式库 API 与 LLM 工具共用）
BUILTIN_TEMPLATE_NAME = "内置标书基准版式"  # 用户可见名（2026-09-09「模板库」改名「版式库」随改）


def templates_dir() -> Path:
    """用户版式文件目录（data/templates/，版式库 API 与工具共用）。"""
    from .. import config

    d = config.data_dir() / "templates"
    d.mkdir(parents=True, exist_ok=True)
    return d


def active_template_name() -> str | None:
    """默认版式的用户文件名；None=内置基准生效。"""
    from .. import db

    return (db.get_setting(TEMPLATE_SETTING_KEY) or "").strip() or None


def list_templates_info() -> list[dict]:
    """版式库清单（读侧真值，HTTP API 与 list_templates 工具同源消费）：
    内置恒首位 + 用户版式 mtime 降序；active=命中默认位（未设默认时内置即
    默认）。stat 竞态（glob 到 stat 之间文件被删）跳行不抛。"""
    active = active_template_name()

    def _row(p: Path, *, key: str, name: str, builtin: bool) -> dict:
        st = p.stat()
        return {
            "name": name,
            "key": key,
            "builtin": builtin,
            "active": key == active or (builtin and active is None),
            "size": st.st_size,
            "mtime": st.st_mtime,
        }

    out = [_row(_BASE_TEMPLATE, key=BUILTIN_TEMPLATE_KEY, name=BUILTIN_TEMPLATE_NAME, builtin=True)]
    items: list[tuple[float, Path]] = []
    for p in templates_dir().glob("*.docx"):
        try:
            items.append((p.stat().st_mtime, p))
        except OSError:
            continue
    for _, p in sorted(items, reverse=True):
        try:
            out.append(_row(p, key=p.name, name=p.stem, builtin=False))
        except OSError:
            continue
    return out


def _active_template_path() -> Path:
    """默认版式：app_settings 命中的用户版式文件（data/templates/）优先，
    否则内置基准。用户版式文件被误删等异常静默回落内置（探测不报错——
    版式缺失不该打断建节）。"""
    try:
        name = active_template_name()
        if name:
            p = templates_dir() / name
            if p.is_file():
                return p
    except Exception:
        pass
    return _BASE_TEMPLATE


def _blank_from(template: Path) -> Document:
    """从指定版式文件起建空白文档（剥离自带的样式示例段，只留 sectPr）。"""
    doc = Document(str(template))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc


def _new_document() -> Document:
    """从默认版式起建（格式与内容分离：版式全部活在版式文件 styles.xml；
    默认位=app_settings 的用户版式，缺省内置基准，由 scripts/
    make_base_template.py 生成维护；代码只挂样式名）。

    版式文件 body 带样式示例段（打开可直观预览/改版式），起建时整段剥离
    （body 只留 sectPr=版面/页脚），示例永不进入节文件与合册。版式文件缺失
    时回落 python-docx 默认模板（英文版式，仅防打包漏带资源，不作为常态）。"""
    tpl = _active_template_path()
    if tpl.is_file():
        return _blank_from(tpl)
    return Document()


def _body_style(doc: Document):
    """Tender Body 样式对象，版式文件不带时返回 None（正文段回落 Normal，
    建节不因换版式缺样式名而失败）。"""
    try:
        return doc.styles[_BODY_STYLE_NAME]
    except KeyError:
        return None


def _style_id_by_name(doc: Document, name: str) -> str | None:
    """按样式名解析 styleId（w:pStyle 的 w:val 挂的是 styleId 不是样式名），
    无此样式返回 None——调用方回落默认样式，换版式缺样式名不失败。"""
    try:
        return doc.styles[name].element.get(qn("w:styleId"))
    except KeyError:
        return None


def _ensure_body_style(p_el, style_id: str) -> None:
    """无样式引用的素材正文段挂版式正文样式（素材拷贝归顺版式，2026-09-08
    用户拍板：素材归顺、招标格式件保真）：获得标书正文缩进/行距/对齐；
    已带样式引用（内置标题自动吃宿主定义、自定义样式走保真迁移）与表格整表
    保持原样——某段自定义素材该不该归顺是语义判断，归写作流程的改写适配
    步骤，机械层不做。"""
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
        p_el.insert(0, ppr)
    if ppr.find(qn("w:pStyle")) is None:
        ppr.insert(0, parse_xml(f'<w:pStyle {nsdecls("w")} w:val="{style_id}"/>'))


# ---------- 样式与自动编号定义迁移 ----------

def _find_by_id(root, tag: str, attr: str, val: str):
    for e in root.findall(qn(tag)):
        if e.get(qn(attr)) == val:
            return e
    return None


def _style_name(style_el) -> str:
    ne = style_el.find(qn("w:name"))
    return ((ne.get(qn("w:val")) if ne is not None else "") or "").strip()


def _norm_style_xml(style_el) -> bytes:
    """样式定义归一化比对键：剥 styleId 属性、rsid 编辑痕迹与 default 标记
    （文档内部门牌/编辑元数据，非定义内容）——撞 id 时「定义是否等价」的判据。"""
    e = deepcopy(style_el)
    e.attrib.pop(qn("w:styleId"), None)
    e.attrib.pop(qn("w:default"), None)
    for k in list(e.attrib):
        if k.split("}")[-1].lower().startswith("rsid"):
            e.attrib.pop(k)
    return etree.tostring(e)


def _free_style_id(taken: set[str]) -> str:
    """分配未用的数字样式 id（目标既有 id 混有非数字名，数字递增最稳）。"""
    n = 1
    while str(n) in taken:
        n += 1
    return str(n)


# 编号定义不允许绑定的目标样式名（小写）：内建标题——章节编号是合册按树序写进
# 标题文本的，「不走样式绑定自动编号」是设计铁则（样式绑定会把拷入标题段一起
# 计数打乱章序）；Normal 是全文档兜底样式，绑了就全篇计数。迁入编号的 pStyle
# 链接解析到这些样式时剥除（2026-09-13 目录乱号批；素材自有样式的编号迁在新
# id/新名下照常工作，不受影响）。
_NUM_STYLE_GUARD_NAMES = (
    {f"heading {i}" for i in range(1, 10)}
    | {f"标题 {i}" for i in range(1, 10)}
    | {"normal", "正文"}
)


def _sanitize_num_style_links(abs_el, dst_styles_root, id_remap: dict[str, str]) -> None:
    """迁入编号定义的样式绑定（lvl/pStyle）改写与防线：值按样式 id 重映射表
    改写（样式撞 id 换新 id 后链接不改写=绑到目标里同 id 的无关样式）；解析到
    目标内建标题/Normal 样式的链接剥除（见 _NUM_STYLE_GUARD_NAMES）。"""
    for lv in abs_el.findall(qn("w:lvl")):
        ps = lv.find(qn("w:pStyle"))
        if ps is None:
            continue
        v = ps.get(qn("w:val")) or ""
        if v in id_remap:
            v = id_remap[v]
            ps.set(qn("w:val"), v)
        target = _find_by_id(dst_styles_root, "w:style", "w:styleId", v)
        if target is not None and _style_name(target).lower() in _NUM_STYLE_GUARD_NAMES:
            lv.remove(ps)


def _merge_missing_numbering(
    src_doc: Document, dst_doc: Document, elements, style_id_remap: dict[str, str] | None = None
) -> int:
    """编号定义迁移：按定义内容判定沿用，id 冲突时重映射——拷贝元素的编号引用
    解析到与源文档相同的定义，并改写 elements 上的引用。

    numbering id 只是文档内部门牌号：建节模板自带 numId 1-9/abstractNumId
    0-8，素材（历史标书）的 id 同样从小数字分配——同 id 不代表同定义，沿用
    目标定义会让中文编号静默变圆点、多级编号错位（2026-09-08 review 实证）。
    迁移规则：
    - dst 缺该 numId：原 id 迁入（abstractNumId 同缺则原 id 齐迁；abstractNumId
      相撞且定义不同则 abstractNum 换新 id 迁入，num 改指向新 id）；
    - dst 已有同 numId：定义内容等价才沿用；不同则分配未用的新 numId 迁入
      源定义并改写引用；
    - 同一次调用内同 numId 共享映射（块内列表编号连续）；跨调用不复用——
      不同素材块的同定义列表保持独立序列（各自从 1 开始，贴近原文语义）；
    - w:numStyleLink（num 链编号样式）不重映射，沿用目标定义——编号样式
      多为内建、语义跨文档一致。
    style_id_remap（随 `_merge_missing_styles` 返回）：迁入 abstractNum 的
    lvl/pStyle 样式绑定按表改写；解析到目标内建标题/Normal 的绑定剥除
    （`_sanitize_num_style_links`，2026-09-13 批——素材多级编号定义自带
    「heading N」绑定，不改写会绑到别人的样式甚至骨架样式上）。
    """
    try:
        src_root = src_doc.part.numbering_part.element
        dst_root = dst_doc.part.numbering_part.element
    except Exception:
        return 0  # 任一侧无 numbering 部件 = 无编号可迁
    used = sorted({
        n.get(qn("w:val"))
        for el in elements
        for n in el.findall(".//" + qn("w:numId"))
        if n.get(qn("w:val"))
    })
    if not used:
        return 0

    def norm_num(num_el) -> bytes:
        # 剥 numId 属性与 abstractNumId 引用（文档内部门牌，非定义内容）
        e = deepcopy(num_el)
        e.attrib.pop(qn("w:numId"), None)
        ref = e.find(qn("w:abstractNumId"))
        if ref is not None:
            e.remove(ref)
        return etree.tostring(e)

    def norm_abs(abs_el) -> bytes:
        # 剥 id 属性与 nsid/tmpl（Word 各文档独立生成的随机 GUID，比对噪音）
        e = deepcopy(abs_el)
        e.attrib.pop(qn("w:abstractNumId"), None)
        for t in ("w:nsid", "w:tmpl"):
            for x in e.findall(qn(t)):
                e.remove(x)
        return etree.tostring(e)

    def def_key(root, nid: str) -> tuple[bytes, bytes] | None:
        num = _find_by_id(root, "w:num", "w:numId", nid)
        if num is None:
            return None
        ref = num.find(qn("w:abstractNumId"))
        abs_e = (
            _find_by_id(root, "w:abstractNum", "w:abstractNumId", ref.get(qn("w:val")))
            if ref is not None else None
        )
        return (norm_num(num), norm_abs(abs_e) if abs_e is not None else b"")

    dst_nums = {n.get(qn("w:numId")): n for n in dst_root.findall(qn("w:num"))}
    dst_abs = {a.get(qn("w:abstractNumId")): a for a in dst_root.findall(qn("w:abstractNum"))}

    def _ints(ids) -> set[int]:
        out = set()
        for k in ids:
            try:
                out.add(int(k))
            except (TypeError, ValueError):
                continue
        return out

    free_num = max(_ints(dst_nums), default=0) + 1
    free_abs = max(_ints(dst_abs), default=0) + 1

    def add_abs(src_aid: str, new_aid: str | None = None) -> str:
        """迁 abstractNum（new_aid 缺省保留原 id）；abstractNum 必须排在全部 num 之前。
        迁入即过样式绑定防线（pStyle 改写+内建保护）。"""
        a = deepcopy(_find_by_id(src_root, "w:abstractNum", "w:abstractNumId", src_aid))
        aid = new_aid or src_aid
        a.set(qn("w:abstractNumId"), aid)
        _sanitize_num_style_links(a, dst_doc.styles.element, style_id_remap or {})
        first_num = dst_root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            dst_root.append(a)
        dst_abs[aid] = a
        return aid

    remap: dict[str, str] = {}
    moved = 0
    for nid in used:
        key = def_key(src_root, nid)
        if key is None:
            continue  # 源侧定义缺失（悬挂引用）——无从迁
        src_num = _find_by_id(src_root, "w:num", "w:numId", nid)
        if src_num.find(qn("w:numStyleLink")) is not None:
            continue  # 编号样式链场景不重映射（见 docstring）
        if nid in dst_nums and def_key(dst_root, nid) == key:
            remap[nid] = nid  # id 相撞但定义等价：沿用目标
            continue
        ref = src_num.find(qn("w:abstractNumId"))
        src_aid = ref.get(qn("w:val")) if ref is not None else None
        new_aid = src_aid
        if src_aid:
            existing = dst_abs.get(src_aid)
            if existing is None:
                add_abs(src_aid)
            elif norm_abs(existing) != key[1]:
                new_aid = add_abs(src_aid, str(free_abs))  # 相撞且定义不同 → 新 id 迁入
                free_abs += 1
            # 相撞但定义等价 → 复用目标既有 abstractNum
        new_num = deepcopy(src_num)
        new_nid = nid
        if nid in dst_nums:
            while str(free_num) in dst_nums:
                free_num += 1
            new_nid = str(free_num)
            free_num += 1
        new_num.set(qn("w:numId"), new_nid)
        if src_aid and new_aid != src_aid:
            new_num.find(qn("w:abstractNumId")).set(qn("w:val"), new_aid)
        dst_root.append(new_num)
        dst_nums[new_nid] = new_num
        remap[nid] = new_nid
        moved += 1

    # 改写拷贝元素上的引用（含迁入的样式定义里的 numPr——调用方传副本进来）
    for el in elements:
        for n in el.findall(".//" + qn("w:numId")):
            v = n.get(qn("w:val"))
            if v in remap and remap[v] != v:
                n.set(qn("w:val"), remap[v])
    return moved


def _merge_missing_styles(src_doc: Document, dst_doc: Document, elements) -> tuple[int, dict[str, str]]:
    """把元素引用而目标缺定义的样式（段落/字符/表格样式）从源拷过来。

    元素级拷贝沿用源文档样式引用——目标缺定义时 Word 回退默认样式，自定义样式
    （带编号的标题、特殊字体段落）会丢观感。依赖链（basedOn/link/next 指向的
    样式、样式自身的编号引用）递归补齐；源也没有的（如内建样式 id 差异）不硬造。

    去重双维（2026-09-13 目录乱号批，实测素材自带「heading 2 同名+挂编号」样式
    迁入后 WPS/LibreOffice 按名解析把整本 98 行正文套上连续序号）：
    - **同 styleId**：与目标同 id 的样式**同名**（内建标题/通用样式跨文档同款，
      素材标题段沿用宿主定义——2026-09-08「素材归顺、招标件保真」拍板）或定义
      等价（`_norm_style_xml`）→ 沿用宿主定义（原行为）；同 id **异名**（不同
      素材生成器的数字 id 空间撞车，id 先到先得随合并序漂移）→ 分配未用新 id
      迁入并改写引用（元素上的 pStyle/rStyle/tblStyle 与迁入簇内部的
      basedOn/link/next）——此前静默跳过会把拷贝段绑到目标里同 id 的无关样式。
    - **同名**（大小写不敏感，含内建 heading N/normal 等）：迁入副本改名
      （"name 2"、"name 3"…直到唯一）——Word/WPS/LibreOffice 对样式存在按名
      解析路径，同名并存会让素材的编号/格式挂到全部同名段落；改名只断名字
      合并路径，styleId 绑定与观感不变。
    - 迁入副本剥 `w:default`（源文档的默认样式标记不顶掉目标默认）。
    返回 (迁入数, 源styleId→目标styleId 重映射表)——调用方传给
    `_merge_missing_numbering` 改写 abstractNum 的 pStyle 绑定。
    """
    src_root = src_doc.styles.element
    dst_root = dst_doc.styles.element
    have_ids = {s.get(qn("w:styleId")) for s in dst_root.findall(qn("w:style"))}
    have_names = {
        _style_name(s).lower()
        for s in dst_root.findall(qn("w:style"))
        if _style_name(s)
    }
    pending = {
        r.get(qn("w:val"))
        for el in elements
        for tag in ("w:pStyle", "w:rStyle", "w:tblStyle")
        for r in el.findall(".//" + qn(tag))
        if r.get(qn("w:val"))
    }
    id_remap: dict[str, str] = {}
    migrated: list = []
    visited: set[str] = set()
    moved = 0
    while pending:
        sid = pending.pop()
        if not sid or sid in visited:
            continue
        visited.add(sid)
        src_style = _find_by_id(src_root, "w:style", "w:styleId", sid)
        if src_style is None:
            continue
        new_sid = sid
        if sid in have_ids:
            existing = _find_by_id(dst_root, "w:style", "w:styleId", sid)
            if existing is not None and (
                _style_name(existing).lower() == _style_name(src_style).lower()
                or _norm_style_xml(existing) == _norm_style_xml(src_style)
            ):
                continue  # 同款样式（同 id 同名=内建/通用样式，或定义等价）：沿用宿主定义
            new_sid = _free_style_id(have_ids)
        new_style = deepcopy(src_style)
        new_style.set(qn("w:styleId"), new_sid)
        new_style.attrib.pop(qn("w:default"), None)
        name = _style_name(new_style)
        if name:
            if name.lower() in have_names:
                k = 2
                while f"{name} {k}".lower() in have_names:
                    k += 1
                name = f"{name} {k}"
                new_style.find(qn("w:name")).set(qn("w:val"), name)
            have_names.add(name.lower())
        dst_root.append(new_style)
        have_ids.add(new_sid)
        if new_sid != sid:
            id_remap[sid] = new_sid
        migrated.append(new_style)
        moved += 1
        for tag in ("w:basedOn", "w:link", "w:next"):
            ref = src_style.find(qn(tag))
            val = ref.get(qn("w:val")) if ref is not None else None
            if val and val not in visited:
                pending.add(val)
    if id_remap:
        for el in list(elements) + migrated:
            for tag in ("w:pStyle", "w:rStyle", "w:tblStyle", "w:basedOn", "w:link", "w:next"):
                for r in el.findall(".//" + qn(tag)):
                    v = r.get(qn("w:val"))
                    if v in id_remap:
                        r.set(qn("w:val"), id_remap[v])
    if migrated:
        # 簇齐后带 remap 一次迁编号：样式 numPr 引用的定义 + abstractNum 的
        # pStyle 绑定改写（样式换 id 后绑定不改写=绑到别人样式）
        _merge_missing_numbering(src_doc, dst_doc, migrated, style_id_remap=id_remap)
    return moved, id_remap


def _merge_missing_comments(src_doc: Document, dst_doc: Document, elements) -> int:
    """把拷贝元素引用的源文档批注迁入目标文档（id 重映射，与编号迁移同构）。

    随段落 deepcopy 拷入的 commentRangeStart/End/commentReference 引用的是源文档
    comments part 的 id——不迁移则批注悬空丢失。只迁拷贝元素实际引用的条目（跳过
    被丢弃的节标题段上的锚），源里查不到的陈旧标记就地摘除防悬空。
    """
    markers = (qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:commentReference"))
    nodes: list = []
    refs: set[str] = set()
    for el in elements:
        for m in el.iter(*markers):
            nodes.append(m)
            refs.add(m.get(qn("w:id")) or "")
    if not nodes:
        return 0
    id_map: dict[str, str] = {}
    for cid in sorted(refs):
        if not cid.isdigit():
            continue
        src_c = src_doc.comments.get(int(cid))
        if src_c is None:
            continue
        new_c = dst_doc.comments.add_comment(
            text=src_c.text or "", author=src_c.author or _COMMENT_AUTHOR, initials=src_c.initials or "AI"
        )
        id_map[cid] = str(new_c.comment_id)
    for m in nodes:
        old = m.get(qn("w:id")) or ""
        if old in id_map:
            m.set(qn("w:id"), id_map[old])
        else:
            m.getparent().remove(m)  # 陈旧/悬空标记：摘除防引用错乱
    return len(id_map)


# ---------- 工具 ----------


def _tool_guard(label: str):
    """工具体兜底：意外异常（损坏 docx、IO 失败等）转失败文案返回，不打崩 run。

    工具铁律：外部输入（用户上传的招标/素材原件）半截或损坏是现实输入，
    Document()/save 抛出的非 ValueError 异常必须转字符串；具体的人话报错
    （ValueError 分支）优先于本兜底。
    """

    def deco(fn):
        @functools.wraps(fn)
        def wrapper(*args, **kwargs):
            try:
                return fn(*args, **kwargs)
            except Exception as e:
                return f"[{label}失败] {type(e).__name__}: {e}"

        return wrapper

    return deco


def _shade_cell(cell, fill: str) -> None:
    """单元格底纹（w:shd 直接写 tcPr，不依赖样式表——换版式不失败）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def _apply_table_borders(tbl) -> None:
    """细灰边框直接写 tblPr——不依赖 Table Grid 样式名（用户版式可能没有该名，
    无边框的表格观感即残；生成表格全部自带边框，样式名零依赖）。"""
    edges = "".join(
        f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="999999"/>'
        for e in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    tbl._tbl.tblPr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>{edges}</w:tblBorders>'))


def _caption_style(doc: Document):
    """图注/表题样式：Tender Caption（合册按样式名识别做全局重编号）；版式缺该
    样式时回落正文样式——建表不因换版式失败，编号侧缺样式=跳过重编号不误伤。"""
    try:
        return doc.styles[_CAPTION_STYLE_NAME]
    except KeyError:
        return _body_style(doc)


def _add_caption(doc: Document, text: str) -> object:
    """图注/表题居中段（中文文档惯例：表题在表格上方、图注在图示下方——调用方
    控制与表格的先后顺序即落对位置）。"""
    para = doc.add_paragraph(text.strip(), style=_caption_style(doc))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def _avail_width(doc: Document):
    """版心宽度（EMU）；取不到回落 15cm——生成表格的列宽按它比例分配。"""
    try:
        sec = doc.sections[-1]
        w = sec.page_width - sec.left_margin - sec.right_margin
        if isinstance(w, int) and w > 0:
            return w
    except (IndexError, TypeError):
        pass
    return Cm(15)


def _set_col_widths(tbl, ratios: list[float], avail) -> None:
    """按比例分配列宽（禁 autofit + gridCol/逐格 tcW 双写——单写 gridCol 在部分
    Word 版本会被内容自配覆盖）。须在 cell.merge 之前调用。"""
    tbl.autofit = False
    total = sum(ratios) or 1.0
    widths = [int(avail * r / total) for r in ratios]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
        for row in tbl.rows:
            row.cells[i].width = w


def _grid_table(doc: Document, grid: list[list[dict]], ratios: list[float]):
    """格子规格网格 → docx 表格（纯机械渲染后端，W1 数据表/W2 图示共用）：
    逐格落 text/fill/bold/center，自带细灰边框，列宽按 ratios 比例分版心。
    格子 dict 键均可省：text（默认空）、fill（默认无底纹）、bold、center。"""
    tbl = doc.add_table(rows=len(grid), cols=len(grid[0]))
    _apply_table_borders(tbl)
    _set_col_widths(tbl, ratios, _avail_width(doc))
    for r, row in enumerate(grid):
        for c, spec in enumerate(row):
            cell = tbl.cell(r, c)
            para = cell.paragraphs[0]
            text = str(spec.get("text", ""))
            if text:
                run = para.add_run(text)
                if spec.get("bold"):
                    run.bold = True
            if spec.get("center"):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fill = spec.get("fill")
            if fill:
                _shade_cell(cell, fill)
    return tbl


def _pad_row(cells: list, ncols: int) -> list:
    """行补齐/截断到列数（模型少给多给一格都不至于建废表）。"""
    return (list(cells) + [""] * ncols)[:ncols]


def _cell_str(v) -> str:
    return str(v) if v is not None else ""


def _parse_body_blocks(body: str) -> tuple[list[dict] | None, str | None]:
    """建节 body 块序列参数解析与校验：[{"type":"p","text":…},
    {"type":"table","caption":…,"header":[…],"rows":[[…],…]}]。
    返回 (规范化块列表, 错误文案)；错误文案带修复写法，不猜不吞。"""
    try:
        blocks = json.loads(body)
    except ValueError:
        return None, (
            "body 不是合法 JSON——须为块数组："
            '[{"type":"p","text":"段落文字"},{"type":"table","caption":"表题",'
            '"header":["列名"],"rows":[["值", …], …]}, …]'
        )
    if not isinstance(blocks, list) or not blocks:
        return None, "body 须为非空 JSON 数组（每块 {\"type\":\"p\"|\"table\", …}）"
    norm: list[dict] = []
    for i, blk in enumerate(blocks):
        if not isinstance(blk, dict):
            return None, f"body 第 {i + 1} 块不是对象"
        btype = blk.get("type")
        if btype == "p":
            text = _cell_str(blk.get("text")).strip()
            if not text:
                return None, f"body 第 {i + 1} 块 p 缺 text"
            norm.append({"type": "p", "text": text})
        elif btype == "table":
            header = blk.get("header") or []
            rows = blk.get("rows")
            if not isinstance(header, list) or any(not isinstance(h, str) for h in header):
                return None, f"body 第 {i + 1} 块 table 的 header 须为字符串数组（可省略）"
            if len(header) > _MAX_TABLE_COLS:
                return None, f"body 第 {i + 1} 块 table 列数 {len(header)} 超上限（≤{_MAX_TABLE_COLS}，列多请拆表或换横向表述）"
            if not isinstance(rows, list) or not rows:
                return None, f"body 第 {i + 1} 块 table 缺 rows（二维字符串数组）"
            if len(rows) > _MAX_TABLE_ROWS:
                return None, f"body 第 {i + 1} 块 table 行数 {len(rows)} 超上限（≤{_MAX_TABLE_ROWS}，超长表请拆节或精简）"
            for r, row in enumerate(rows, 1):
                if not isinstance(row, list) or any(not isinstance(v, (str, int, float)) for v in row):
                    return None, f"body 第 {i + 1} 块 table 第 {r} 行不是字符串数组"
                # 行宽超 header 报错不截断：_pad_row 会把超宽行静默裁到 header
                # 列数（多出的格丢失）；少列的宽容保留（补空串）
                if header and len(row) > len(header):
                    return None, (
                        f"body 第 {i + 1} 块 table 第 {r} 行 {len(row)} 列超过 header 的 "
                        f"{len(header)} 列——多出的列并入合适列，或补全 header；列不足可省略（自动留空）"
                    )
            ncols = len(header) or max(len(r) for r in rows)
            if ncols > _MAX_TABLE_COLS:
                return None, f"body 第 {i + 1} 块 table 列数 {ncols} 超上限（≤{_MAX_TABLE_COLS}）"
            rows_n = [_pad_row([_cell_str(v) for v in row], ncols) for row in rows]
            for r, row in enumerate(rows_n, 1):
                for c, v in enumerate(row, 1):
                    if len(v) > _MAX_CELL_CHARS:
                        return None, f"body 第 {i + 1} 块 table 第 {r} 行第 {c} 格 {len(v)} 字超上限（≤{_MAX_CELL_CHARS}，长文放段落不放格）"
            norm.append(
                {
                    "type": "table",
                    "caption": _cell_str(blk.get("caption")).strip(),
                    "header": [h.strip() for h in header],
                    "rows": rows_n,
                }
            )
        else:
            return None, f"body 第 {i + 1} 块 type 非法（{btype}，须 p 或 table）"
    return norm, None


def _add_data_table(doc: Document, header: list[str], rows: list[list[str]]):
    """普通数据表：表头行加粗+浅灰底纹居中，数据行不加修饰（可编辑、validate
    可读格文本——承诺数据落格不落图）。"""
    ncols = len(header) if header else max(len(r) for r in rows)
    grid: list[list[dict]] = []
    if header:
        grid.append(
            [{"text": h, "bold": True, "fill": _DIAG_ITEM_FILL, "center": True} for h in header]
        )
    for row in rows:
        grid.append([{"text": v} for v in _pad_row(row, ncols)])
    return _grid_table(doc, grid, [1.0] * ncols)


@tool
@_tool_guard("创建")
def docx_section_create(path: str, title: str, paragraphs: str = "", replace: bool = False, body: str = "") -> str:
    """创建正文节 docx 文件（work/body/ 下，每节一文件）。

    用途：tender-body 正文阶段为一节建档（标题 + 可选初始内容）。已存在默认不
    覆盖（防误清已写内容）；**重写本节传 replace=true**（用户裁决的重写范围），
    旧版自动入恢复点栈（保留最近 3 版，文件名不变旁挂 .restorepoints/ 目录）。
    Args:
        path: 相对任务 work/ 的路径，如 body/技术部分/3.1 需求分析.docx（自动补 .docx）
        title: 节标题（写入文档一级标题样式，进目录域）
        paragraphs: 可选初始段落，换行分隔（纯文字成节——人员配置/里程碑/对比类
               内容**不要用这个**，用 body 块序列混排表格）
        body: 可选块序列 JSON（推荐：段落+表格混排一次成形）——
              [{"type":"p","text":"段落文字"},
               {"type":"table","caption":"表题（可选，表上方居中）",
                "header":["岗位","人数","职责"],"rows":[["项目经理","1","…"],…]},
               …]；表头行自动加粗+浅灰，数据格可编辑可校验；约束 列≤8/行≤60/
              单格≤200 字、行宽不得超过 header 列数，超限报错给修复写法。
              body 存在时 paragraphs 被忽略
        replace: 已存在的节重写时传 true；默认 false 不覆盖
    """
    task_id = _task_id()
    if not task_id:
        return "[创建失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, path, must_exist=False)
    except ValueError as e:
        return f"[创建失败] {e}"
    blocks: list[dict] | None = None
    if (body or "").strip():
        blocks, err = _parse_body_blocks(body)
        if err:
            return f"[创建失败] {err}"
    replaced = False
    with _docx_path_lock(dst):
        if dst.exists():
            if not replace:
                return (
                    f"[创建失败] work/{rel} 已存在（不覆盖已写内容）——"
                    "确属重写范围请传 replace=true（旧版自动存恢复点）"
                )
            _rotate_restore_point(dst)
            replaced = True
        doc = _new_document()
        doc.add_heading(title.strip() or "未命名", 1)
        body_style = _body_style(doc)
        n_init = 0
        n_tables = 0
        if blocks is not None:
            # 块序列混排：段/表按序落（表题先于表格 add 即落其上方），新建内容
            # 不带修订标记（与初始段落同口径——修订标记只属于对既有文本的改写）
            for blk in blocks:
                if blk["type"] == "p":
                    doc.add_paragraph(blk["text"], style=body_style.name if body_style else None)
                    n_init += 1
                else:
                    if blk["caption"]:
                        _add_caption(doc, blk["caption"])
                    _add_data_table(doc, blk["header"], blk["rows"])
                    n_tables += 1
        else:
            for text in paragraphs.split("\n"):
                if text.strip():
                    doc.add_paragraph(text.strip(), style=body_style.name if body_style else None)
                    n_init += 1
        dst.parent.mkdir(parents=True, exist_ok=True)
        _atomic_save(doc, dst)
    head = "[已重建] " if replaced else "[已创建] "
    stat = f"正文 {n_init} 段" + (f"+表格 {n_tables} 张" if n_tables else "")
    return (
        f"{head}work/{rel}（标题「{title.strip()}」，{stat}）\n"
        "下一步：贴素材底稿用 docx_material_inject；直接写内容/修订用 docx_section_revise。"
    )


@tool
@_tool_guard("读取")
def docx_section_read(path: str) -> str:
    """读取正文节 docx 的文本视图（段落编号 P1..Pn + 样式 + 图片/修订标记，表格 T1..Tm）。

    用途：写或修订前先读视图拿段落序号；docx_section_revise 的 para 序号以本视图为准。
    """
    task_id = _task_id()
    if not task_id:
        return "[读取失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, path, must_exist=True)
    except ValueError as e:
        return f"[读取失败] {e}"
    doc = Document(str(dst))
    return "\n".join(view_lines(doc))


@tool
@_tool_guard("注入")
def docx_material_inject(block_id: str, dest: str) -> str:
    """把写作素材块原样注入正文节 docx（元素级拷贝：段落/表格含合并单元格、图片零转写）。

    用途：素材修订模式的底稿落地——检索（search_references）命中块后，用其块 id
    （blk_ 开头）把历史章节整体拷进目标节文件，格式与内容保真，后续用
    docx_section_revise 定向修订。注入内容不带修订标记（干净底稿，修订才标记）。
    仅支持 docx 原件的素材（历史标书大头）；pdf 等原件无可注入元素。
    Args:
        block_id: 素材块 id（search_references 命中的引用键 mt:<fid>:b:<块id> 中的块 id）
        dest: 目标节文件（须已用 docx_section_create 创建）
    """
    task_id = _task_id()
    if not task_id:
        return "[注入失败] 当前会话未归属任务"
    b = db.mt_get_block(block_id.strip())
    if not b:
        return "[注入失败] 素材块不存在（可能已被删除，请重新检索选用）"
    f = db.mt_get_file(b["file_id"])
    if not f:
        return "[注入失败] 素材文件记录缺失"
    src_path = materials_lib.mt_files_dir() / f["file_name"]
    if src_path.suffix.lower() != ".docx":
        return f"[注入失败] 素材「{f['file_name']}」原件不是 docx（{src_path.suffix or '无后缀'}），没有可保真注入的元素；请以文本方式参考该素材自行撰写"
    try:
        dst, rel = _dest_path(task_id, dest, must_exist=True)
    except ValueError as e:
        return f"[注入失败] {e}"
    if not src_path.is_file():
        return "[注入失败] 素材原件文件缺失（可能已被删除）"
    element_map = materials_lib.read_element_map(f["file_name"])
    if element_map is None:
        # 旧解析未产映射：幂等重跑补产。补跑前后比对 md——「同原件确定性解析、
        # 锚定不变」只在解析代码不变时成立，解析升级会使行号漂移、旧块勾选区间
        # 错位（拷错元素而非报错），变了拒绝注入、点名重勾（探测+用户裁决）
        md_path = materials_lib.mt_parse_paths(f["file_name"])[0]
        old_md = md_path.read_text(encoding="utf-8") if md_path.is_file() else None
        materials_lib.run_parse(f["id"])
        element_map = materials_lib.read_element_map(f["file_name"])
        if (
            element_map is not None
            and old_md is not None
            and old_md != md_path.read_text(encoding="utf-8")
        ):
            return (
                "[注入失败] 素材重新解析后内容行号已变化，现有素材块的勾选区间会错位"
                "——请在写作素材库重新确认该文件的勾选（删除重建相关块）后再注入"
            )
    if element_map is None:
        return "[注入失败] 素材元素映射不可用（解析未产出），请删除后重新上传该素材"
    # 同块重复注入探测：注入是 append 语义，重复=内容翻倍。块特征 shingle 已大量
    # 出现在节文件接受视角文本中=已注入过（硬拦；概率防线，注入后大幅改写会漏拦
    # ——validate_body 的素材使用率另有兜底）
    block_sections = (materials_lib.block_content(b["id"]) or {}).get("sections") or []
    bs = _sig_shingles("\n".join(sec.get("text") or "" for sec in block_sections))
    if len(bs) >= 5:
        cur = "\n".join(section_text_lines(Document(str(dst))))
        if len(bs & _sig_shingles(cur)) / len(bs) >= 0.6:
            return (
                f"[注入失败] 素材块《{b.get('title') or b['id']}》的内容已大量出现在 "
                f"work/{rel}（已注入过）——注入是追加语义，重复注入会翻倍内容；"
                "确需重新贴底稿请先重建节文件（docx_section_create 传 replace=true）"
            )
    ranges = b.get("ranges") or []  # db 层已反序列化为 [[s,e],…]
    idx = sorted({
        el for el, s, e in element_map
        if isinstance(el, int) and any(
            isinstance(r, (list, tuple)) and len(r) == 2 and r[0] <= e and r[1] >= s
            for r in ranges
        )
    })
    if not idx:
        return "[注入失败] 素材块区间未映射到任何 docx 元素（可能只勾选了空段落）"
    with _docx_path_lock(dst):
        src = Document(str(src_path))
        dst_doc = Document(str(dst))
        body_style = _body_style(dst_doc)
        body_style_id = body_style.element.get(qn("w:styleId")) if body_style else None
        children = list(src.element.body.iterchildren())
        sect = dst_doc.element.body.find(qn("w:sectPr"))
        n_para = n_tbl = n_img = 0
        copied: list = []
        for i in idx:
            if i >= len(children):
                continue
            el = children[i]
            if el.tag.split("}")[-1] not in ("p", "tbl"):
                continue
            new_el = deepcopy(el)
            _strip_inner_sectpr(new_el)
            _strip_copy_residue(new_el)
            if new_el.tag == qn("w:p") and body_style_id:
                _ensure_body_style(new_el, body_style_id)
            n_img += _migrate_images(src, dst_doc, new_el)
            if sect is not None:
                sect.addprevious(new_el)
            else:
                dst_doc.element.body.append(new_el)
            copied.append(new_el)
            if el.tag.split("}")[-1] == "tbl":
                n_tbl += 1
            else:
                n_para += 1
        if n_para + n_tbl == 0:
            return "[注入失败] 映射区间内没有可注入的段落/表格元素"
        n_style, style_remap = _merge_missing_styles(src, dst_doc, copied)
        n_num = _merge_missing_numbering(src, dst_doc, copied, style_id_remap=style_remap)
        _atomic_save(dst_doc, dst)
    # AI 引用打点：注入落盘成功才算一次引用（容错——打点失败不影响注入结果）
    try:
        db.mt_touch_blocks([b["id"]])
    except Exception:
        pass
    return (
        f"[已注入] 《{b.get('title') or '素材块'}》→ work/{rel}：段落 {n_para} 个、"
        f"表格 {n_tbl} 张、图片 {n_img} 张（内容与格式逐字节取自素材原件，样式引用沿用源文档）"
        + (f"，随迁样式定义 {n_style} 个、编号定义 {n_num} 组" if n_style or n_num else "")
        + "\n下一步：docx_section_read 拿段落序号 → docx_section_revise 定向修订"
        "（旧项目名/公司名/参数按本次承诺清单改）→ check_name_residue 扫残留。"
    )


@tool
@_tool_guard("注入")
def docx_source_inject(source: str, dest: str, lines: str = "") -> str:
    """把任务 sources/ 里招标原件的 docx 内容原样拷进正文节（元素级：段落/表格/图片零转写）。

    用途：格式跟随与格式件节点（投标函/授权书/投标一览表等）——招标规定的格式
    原样拷入目标节文件（格式逐字节保真，不重排不走样），再用 docx_section_revise
    填空适配（表格空格子 fill、旧值 replace；承诺值出自清单、未定值【待补：…】）。
    仅支持 docx 原件；pdf 等无可拷元素，按该来源的解析文本自行成形。注入内容
    不带修订标记（干净底稿，填空改值才标记）。
    Args:
        source: 任务 sources/ 里的文件名（如 招标文件.docx、格式附件-投标函.docx；
                目录部分会被剥掉，只在 sources/ 内寻址）
        dest: 目标节文件（须已用 docx_section_create 创建）
        lines: 可选行号区间 "起-止"（该来源解析 md 的行号闭区间，取自
               work/parse/<文件名>/<文件名>.outline.json 的 start_line/end_line 或
               analysis 表出处列，可带 L 前缀如 L412-L430）；留空 = 整份文件
               （独立格式附件的常态用法）
    """
    task_id = _task_id()
    if not task_id:
        return "[注入失败] 当前会话未归属任务"
    name = (source or "").strip().replace("\\", "/").rsplit("/", 1)[-1]
    if not name:
        return "[注入失败] 未指定来源文件"
    sroot = sources_dir(task_id).resolve()
    src_path = (sroot / name).resolve()
    if sroot != src_path and sroot not in src_path.parents:
        return f"[注入失败] 路径越界：{source}"
    if not src_path.is_file():
        return (
            f"[注入失败] sources/ 下没有文件：{name}"
            "（文件名以 check_pipeline_state 的 [sources]/[candidates] 段为准）"
        )
    if src_path.suffix.lower() != ".docx":
        return (
            f"[注入失败] 「{name}」原件不是 docx（{src_path.suffix or '无后缀'}），"
            "没有可保真拷贝的元素；请按该来源的解析文本（work/parse）自行成形本节"
        )
    lo = hi = None
    if (lines or "").strip():
        try:
            a, _, b = (lines or "").strip().partition("-")
            lo, hi = int(a.strip().lstrip("Ll")), int(b.strip().lstrip("Ll"))
        except ValueError:
            return "[注入失败] lines 须为「起-止」行号（如 412-430 或 L412-L430），取自 outline.json 或 analysis 出处列"
        if lo > hi:
            lo, hi = hi, lo
    try:
        dst, rel = _dest_path(task_id, dest, must_exist=True)
    except ValueError as e:
        return f"[注入失败] {e}"
    # 现场跑解析注册表拿元素映射：确定性幂等（同原件同映射），不依赖任务侧
    # 解析是否落盘——独立格式附件可能没走过 parse_document（走缓存）
    idx = sorted({
        el for el, s, e in _source_element_lines(src_path)
        if isinstance(el, int) and (lo is None or (s <= hi and e >= lo))
    })
    if not idx:
        return "[注入失败] 行号区间未映射到任何 docx 元素（核对 outline.json 的行号区间）"
    with _docx_path_lock(dst):
        src = Document(str(src_path))
        dst_doc = Document(str(dst))
        children = list(src.element.body.iterchildren())
        # 同区间重复注入拦截（2026-09-10 review，与素材侧同口径）：注入是 append
        # 语义，同区间重拷=格式件整段翻倍（写手遗忘已注入/工具超时重试/HITL 续跑
        # 重放同一调用都可能触发），且下游 validate_body 节间查重/合册对同节重复
        # 均无防线、可直达交付合册。概率防线：大幅改写后漏拦可接受
        incoming = _sig_shingles(
            "".join("".join(children[i].itertext()) for i in idx if i < len(children))
        )
        if len(incoming) >= 5:
            # 目标侧同口径 itertext（section_text_lines 的表格行/单元格有截断，
            # 重叠率被稀释漏拦——两侧同为元素全量文本，重复注入时 ≈100% 命中）
            cur = _sig_shingles("".join(dst_doc.element.body.itertext()))
            if len(incoming & cur) / len(incoming) >= 0.6:
                scope_hit = "整份文件" if lo is None else f"行号区间 {lo}-{hi}"
                return (
                    f"[注入失败] 《{name}》（{scope_hit}）的内容已大量出现在 work/{rel}"
                    "（已注入过）——注入是追加语义，重复拷贝会翻倍格式件；确需重拷请先重建"
                    "节文件（docx_section_create 传 replace=true）"
                )
        sect = dst_doc.element.body.find(qn("w:sectPr"))
        n_para = n_tbl = n_img = 0
        copied: list = []
        for i in idx:
            if i >= len(children):
                continue
            el = children[i]
            if el.tag.split("}")[-1] not in ("p", "tbl"):
                continue
            new_el = deepcopy(el)
            _strip_inner_sectpr(new_el)
            _strip_copy_residue(new_el)
            n_img += _migrate_images(src, dst_doc, new_el)
            if sect is not None:
                sect.addprevious(new_el)
            else:
                dst_doc.element.body.append(new_el)
            copied.append(new_el)
            if el.tag.split("}")[-1] == "tbl":
                n_tbl += 1
            else:
                n_para += 1
        if n_para + n_tbl == 0:
            return "[注入失败] 区间内没有可拷贝的段落/表格元素"
        n_style, style_remap = _merge_missing_styles(src, dst_doc, copied)
        n_num = _merge_missing_numbering(src, dst_doc, copied, style_id_remap=style_remap)
        _atomic_save(dst_doc, dst)
    scope = "整份文件" if lo is None else f"行号区间 {lo}-{hi}"
    return (
        f"[已注入] 《{name}》（{scope}）→ work/{rel}：段落 {n_para} 个、"
        f"表格 {n_tbl} 张、图片 {n_img} 张（内容与格式逐字节取自招标原件）"
        + (f"，随迁样式定义 {n_style} 个、编号定义 {n_num} 组" if n_style or n_num else "")
        + "\n下一步：docx_section_read 拿段落序号与表格格坐标 → docx_section_revise 填空"
        "（表格空格子 fill、旧值 replace；招标方信息保留、我方信息按承诺清单填，未定值【待补：…】）"
        "→ validate_body 自查。"
    )


# 图片显示高度上限（2026-09-14 实测修复）：mermaid TD 长链流程图非常高（实测
# 显示高 34.7~62.5cm，A4 版心可用高仅 ~24.6cm），按版心宽等比缩放后一图占一页
# 还溢出。封顶 18cm：超高的图按高度反缩（变窄居中），另配 mermaid 紧凑排版
# （前端 nodeSpacing/rankSpacing 收紧）从源头降低自然高度。
_IMG_MAX_HEIGHT = Cm(18)


def _tracked_image_paragraph(doc: Document, data: bytes):
    """全宽居中 + 段落标记/内容双插入修订的图片段（拒绝修订=整段含图消失）。
    docx_image_insert 与 docx_html_figure/docx_diagram_insert(kind=flow) 共用；
    图片无法解析抛 ValueError（调用方转人话）。图片段不挂 Tender Body（正文样式
    带首行缩进会把图推偏）。全宽等比缩放后超高（>18cm）的按高度反缩居中。"""
    sec = doc.sections[-1] if doc.sections else None
    try:
        avail = sec.page_width - sec.left_margin - sec.right_margin if sec else None
    except TypeError:
        avail = None
    width = avail if isinstance(avail, int) and avail > 0 else Cm(15)
    # drawing XML 让 python-docx 生成（先临时挂在节末），修订标记再手包+搬位
    tmp_p = doc.add_paragraph()
    run = tmp_p.add_run()
    try:
        shape = run.add_picture(BytesIO(data), width=width)
        if shape.height > _IMG_MAX_HEIGHT:  # 超高图按高度反缩（等比，宽度随之变窄）
            ratio = _IMG_MAX_HEIGHT / shape.height
            shape.height = int(_IMG_MAX_HEIGHT)
            shape.width = int(shape.width * ratio)
    except Exception as e:
        raise ValueError(f"图片无法解析（{type(e).__name__}）") from e
    date = _now_iso()
    rev_id = _next_rev_id(doc)
    ins_el = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="{rev_id + 1}" w:author="{_AUTHOR}" w:date="{date}"/>'
    )
    r_el = run._r
    r_el.addprevious(ins_el)
    ins_el.append(r_el)
    p_el = tmp_p._p
    # 缺段落标记修订会让拒绝视角多出空行（自校验拦截）
    p_el.insert(0, parse_xml(
        f'<w:pPr {nsdecls("w")}>'
        f'<w:jc w:val="center"/>'
        f'<w:rPr><w:ins w:id="{rev_id}" w:author="{_AUTHOR}" w:date="{date}"/></w:rPr>'
        f'</w:pPr>'
    ))
    return p_el


def _place_after_anchor(doc: Document, p_el, after: str, label: str) -> tuple[str | None, str]:
    """段元素放位：after=P 序号锚后（可带 P 前缀），空=文末（sectPr 前）。
    返回 (错误文案|None, 位置描述)——错误文案带调用方 label 前缀。"""
    if (after or "").strip():
        try:
            idx = int((after or "").strip().lstrip("Pp"))
        except ValueError:
            return f"{label} after 须为段落号（如 5 或 P5，以 docx_section_read 视图为准）", ""
        paras = doc.paragraphs
        if not 1 <= idx <= len(paras):
            return f"{label} 段落号超范围：after={idx}，本节视图共 {len(paras)} 段", ""
        paras[idx - 1]._p.addnext(p_el)
        return None, f"P{idx} 之后"
    sect = doc.element.body.find(qn("w:sectPr"))
    if sect is not None:
        sect.addprevious(p_el)
    else:
        doc.element.body.append(p_el)
    return None, "节末"


@tool
@_tool_guard("插图")
def docx_image_insert(dest: str, image: str, after: str = "", page: int = 1) -> str:
    """把一张图片插入正文节 docx（证书复印件/扫描件/系统截图——独立图片进正文的唯一
    通道，自写文字无法带图）。

    图片作为新段落落在指定位置：全宽适配（页宽减边距）、居中、带插入修订标记
    （Word 审阅可逐张接受/拒绝）。多张图=多次调用。
    Args:
        dest: 目标节文件（须已用 docx_section_create 创建）
        image: 图片路径（工作区相对，三类来源）：①知识库图片
               knowledge/parse/<文件stem>/images/img_001.png——search_company_assets
               命中行会给出路径；PDF 是整页渲染图、文件名即页码（img_007.png=第 7 页，
               空白页跳过），docx 是文档内嵌图；
               ②当前任务上传图 sources/<文件名>（可带任务前缀）；
               ③PDF 原件（knowledge/files/xxx.pdf 或 sources/xxx.pdf）——传页号
               现场渲染那一页为图。支持 png/jpg/bmp/gif/webp（webp 自动转 png
               插入）；docx 原件不支持（换 PDF 或图片版）。
        after: 插在该段落号之后（P 序号以 docx_section_read 视图为准，可带 P 前缀）；
               留空=追加到节末尾
        page: image 指向 PDF 时的页号（1 起；图片文件忽略此参数）
    """
    task_id = _task_id()
    if not task_id:
        return "[插图失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, dest, must_exist=True)
    except ValueError as e:
        return f"[插图失败] {e}"
    # 图源解析：sources/ 短前缀（含任务前缀形态）归一到当前任务 sources/，其余按
    # workspace 相对（knowledge/…）；resolve 后 containment——图源只认工作区内文件
    segs = [s for s in (image or "").strip().replace("\\", "/").split("/") if s]
    if segs and segs[0] == task_id:
        segs = segs[1:]
    if not segs:
        return "[插图失败] 未指定图片路径"
    if segs[0] == "sources":
        if len(segs) < 2:
            return "[插图失败] sources/ 后未指定文件名"
        base = sources_dir(task_id).resolve()
        src = (base / "/".join(segs[1:])).resolve()
        if base != src and base not in src.parents:
            return f"[插图失败] 路径越界：{image}"
    else:
        root = workspace_dir().resolve()
        src = (root / "/".join(segs)).resolve()
        if root != src and root not in src.parents:
            return (
                f"[插图失败] 路径越界或不在工作区内：{image}"
                "（图源=知识库图片 knowledge/parse/…/images/ 或任务 sources/ 下的文件）"
            )
    if not src.is_file():
        return f"[插图失败] 图片文件不存在：{image}（以检索命中行给出的路径或任务 sources/ 清单为准）"

    suffix = src.suffix.lower()
    if suffix == ".pdf":
        n_pages = pdfium_kit.page_count(src)
        if not 1 <= int(page) <= n_pages:
            return f"[插图失败] 页号超范围：该 PDF 共 {n_pages} 页（page 须 1-{n_pages}）"
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tf:
            tmp_png = Path(tf.name)
        try:
            render_page_png(src, int(page), tmp_png)
            data = tmp_png.read_bytes()
        finally:
            tmp_png.unlink(missing_ok=True)
        src_label = f"{image} 第 {page} 页（按页渲染）"
    elif suffix in (".png", ".jpg", ".jpeg", ".bmp", ".gif"):
        data = src.read_bytes()
        src_label = image
    elif suffix == ".webp":
        # python-docx 不认 webp（docx.image 只有 png/jpg/bmp/gif/tiff 解析器），
        # Pillow 解码后转 PNG 再插——知识库上传白名单收 webp（IMAGE_EXTS），此处
        # 不接会把「收得进、插不进」的矛盾留给用户（2026-09-14 复核批）。
        # 坏文件（半截/伪造后缀）按人话报错，不裸抛类型名（_tool_guard 兜底是英文）
        from PIL import Image

        try:
            with Image.open(src) as im:
                buf = BytesIO()
                im.convert("RGBA").save(buf, format="PNG")
        except Exception:
            return (
                f"[插图失败] {image} 不是有效的 webp 图片（无法解码）——"
                "请确认文件完好，或换 png/jpg 上传"
            )
        data = buf.getvalue()
        src_label = f"{image}（webp 已转 png）"
    else:
        return (
            f"[插图失败] 「{image}」不是可插入的图片或 PDF（{suffix or '无后缀'}）"
            "——docx 原件不作图源，请传 PDF 或图片版"
        )

    with _docx_path_lock(dst):
        doc = Document(str(dst))
        before = _flatten_rejected(doc)
        try:
            p_el = _tracked_image_paragraph(doc, data)
        except ValueError as e:
            return f"[插图失败] {e}——请确认文件是完好的 png/jpg 图片"
        err, where = _place_after_anchor(doc, p_el, after, "[插图失败]")
        if err:
            return err
        if _flatten_rejected(doc) != before:
            return "[插图失败] 修订标记自校验未通过（未保存，文件未变）——请重试或换图源"
        _atomic_save(doc, dst)
    return (
        f"[已插图] {src_label} → work/{rel}（{where}）：图片 1 张，全宽居中、带插入修订标记"
        "\n下一步：docx_section_read 确认位置（图片段显示〔图×1〕）；需要图注可在该段前后"
        "用 docx_section_revise 插文字段。"
    )


# ---- 界面原型（webview 光栅化，2026-09-14 批二）----
# server 端 HTML 清洗：剥脚本块/iframe/on* 事件属性——静默剥不报错（CSP 在前端
# wrapper 页再兜一层；两道锁都过不了的部分本来就不该出现在原型里）
_HTML_MAX_BYTES = 100 * 1024
_SCRIPT_BLOCK_RE = re.compile(r"<script\b[^>]*>.*?</script\s*>|<script\b[^>]*/\s*>", re.I | re.S)
_IFRAME_TAG_RE = re.compile(r"</?iframe\b[^>]*/?\s*>", re.I)
_ON_ATTR_RE = re.compile(r"\son[a-z]+\s*=\s*(?:\"[^\"]*\"|'[^']*'|[^\s>]+)", re.I)


def _sanitize_html(html: str) -> str:
    return _ON_ATTR_RE.sub("", _IFRAME_TAG_RE.sub("", _SCRIPT_BLOCK_RE.sub("", html)))


@tool
@_tool_guard("原型")
def docx_html_figure(dest: str, html: str, after: str = "", caption: str = "") -> str:
    """把界面原型渲染成图片插入正文节 docx（webview 光栅化，2026-09-14 批二）。

    用途：系统界面/功能页的原型图——模型只写 HTML，渲染由应用自带的前端 webview
    完成（隐藏 iframe + html2canvas，本机离线、零外部依赖）。适用于「方案节要展示
    系统长什么样」的场景；不适用于真实系统截图（那是 docx_image_insert 的事）。
    按写作指引「图示」列的原型计划调用，计划外不配图。
    Args:
        dest: 目标节文件（相对任务 work/，须已用 docx_section_create 创建）
        html: 原型页 HTML——**全部样式内联**（<style> 或 style 属性），禁外链
               资源/字体/图片（外链会被内容安全策略拦掉，页面会缺件）、禁脚本
               （会被剥除）；≤100KB、建议 ≤150 行；写法示例见 section-writing.md
               「界面原型」
        after: 插在该段落号之后（P 序号以 docx_section_read 视图为准，可带 P 前缀）；
               留空=追加到节末尾
        caption: 图片下方居中图注（如「审批模块界面原型」——不带编号，整本合册时
               按树序全局编号）；空=不加图注

    渲染产物自动带「界面原型 · 示意图」角标（wrapper 固定注入，防止原型被误当
    真实截图）；图片全宽居中、带插入修订标记；渲染源 HTML 落盘
    work/body/assets/ 留底。前端渲染服务不可用/超时（应用不在前台等）时返回
    失败提示——改用文字描述界面或批注记待补，勿反复重试。
    """
    task_id = _task_id()
    if not task_id:
        return "[原型失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, dest, must_exist=True)
    except ValueError as e:
        return f"[原型失败] {e}"
    raw = (html or "").strip()
    if not raw:
        return "[原型失败] html 为空——界面原型须给出完整 HTML（全部样式内联）"
    n_bytes = len(raw.encode("utf-8"))
    if n_bytes > _HTML_MAX_BYTES:
        return f"[原型失败] html {n_bytes // 1024}KB 超上限（≤100KB）——精简页面或拆成多张原型"
    clean = _sanitize_html(raw)
    # 源留底：工作台只列 .md/.docx（assets/ 不可见），纯磁盘资产——将来「改 HTML
    # 重渲染」的入口，run_files 的 .md/.docx 收录口径也不受扰
    assets_dir = dst.parent / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    (assets_dir / f"{uuid.uuid4().hex}.html").write_text(clean, encoding="utf-8")
    return _render_and_insert(
        task_id, dst, rel, after, caption,
        {"kind": "html", "html": clean}, "[原型失败]",
        f"[已插原型] work/{rel}：界面原型 1 张（webview 渲染，带角标）",
    )


def _insert_rendered_png(
    dst: Path, rel: str, after: str, caption: str, png: bytes, label: str
) -> tuple[str | None, str | None]:
    """渲染产物 PNG 落进节文件（docx_html_figure 与 docx_diagram_insert kind=flow
    共用）：全宽居中+插入修订的图片段、锚定放位、自校验后加图注（图注在图下方，
    无修订新建段——自校验须在加图注之前做，否则拒绝视角比对误判）。
    返回 (错误文案|None, 成功文案|None)。label 用于错误前缀（[原型失败]/[图示失败]）。"""
    with _docx_path_lock(dst):
        doc = Document(str(dst))
        before = _flatten_rejected(doc)
        try:
            p_el = _tracked_image_paragraph(doc, png)
        except ValueError as e:
            return f"{label} 渲染产物异常（{e}）——请重试一次，再败则降级处理", None
        err, where = _place_after_anchor(doc, p_el, after, label)
        if err:
            return err, None
        if _flatten_rejected(doc) != before:
            return f"{label} 修订标记自校验未通过（未保存，文件未变）——请重试", None
        cap = (caption or "").strip()
        if cap:
            p_el.addnext(_add_caption(doc, cap)._p)
        _atomic_save(doc, dst)
    return None, where


def _render_and_insert(
    task_id: str, dst: Path, rel: str, after: str, caption: str, payload: dict, label: str, stat: str
) -> str:
    """经 webview 光栅化队列渲染并插入（html 原型/mermaid 流程共用）：
    登记渲染请求 → 阻塞等待回执 → 插 PNG；超时/失败给降级文案（不死等不卡 run）。
    label 为错误前缀（[原型失败]/[图示失败]），降级前缀由它派生（[原型渲染失败]）。"""
    rid = render_queue.register(task_id, rel, after, caption, payload)
    png = render_queue.wait(rid)
    if png is None:
        return (
            f"{label.replace('失败]', '渲染失败]')} 前端渲染服务未响应（应用可能不在前台或"
            "渲染超时）——改用文字描述该内容，或 docx_comment_add 批注记待补后继续；"
            "不要原地反复重试"
        )
    err, where = _insert_rendered_png(dst, rel, after, caption, png, label)
    if err:
        return err
    bits = f"{stat}（{where}）"
    if (caption or "").strip():
        bits += (
            "，图注 1 段（图注段使其后段落序号 +1）\n（最新读视图如下——后续修订按此序号"
            "定位即可，无需再调 docx_section_read）\n" + "\n".join(view_lines(Document(str(dst))))
        )
    else:
        bits += (
            "。图片段显示〔图×1〕；图注须在插入时经 caption 参数给出——漏了可接受"
            "无图注，或 docx_comment_add 批注（锚定图片附近段落）请用户在 Word 补；"
            "不要再调本工具补图注（会重复出图）"
        )
    return bits


@tool
@_tool_guard("批注")
def docx_comment_add(path: str, text: str, after: str = "") -> str:
    """给正文节 docx 加一条 Word 批注——缺料/待澄清/待核验等待办的唯一落点。

    批注锚定在指定段落上：Word 审阅侧栏可见（作者 灵燕智能）、打印与 PDF
    导出默认不带；docx_section_read 视图该段行尾显示〔批注：…〕，validate_body
    清点待办批注，合册时批注原样迁移进整本。**正文里禁止写【待补】【待澄清】
    占位文字**（会进交付稿）——缺什么、要用户裁决什么、哪个数字待核验，都走本工具。
    Args:
        path: 目标节文件（相对任务 work/，须已用 docx_section_create 创建）
        text: 批注内容——写清缺什么/要澄清什么（如「项目经理 PMP 证书编号缺失，
              需向用户确认后回填」）
        after: 锚定段落号（P 序号以 docx_section_read 视图为准，可带 P 前缀）；
               留空=节末最后一段
    """
    task_id = _task_id()
    if not task_id:
        return "[批注失败] 当前会话未归属任务"
    if not (text or "").strip():
        return "[批注失败] 批注内容为空——写清缺什么/要澄清什么"
    try:
        dst, rel = _dest_path(task_id, path, must_exist=True)
    except ValueError as e:
        return f"[批注失败] {e}"
    with _docx_path_lock(dst):
        doc = Document(str(dst))
        paras = doc.paragraphs
        if not paras:
            return "[批注失败] 本节没有段落可锚定"
        if (after or "").strip():
            try:
                idx = int((after or "").strip().lstrip("Pp"))
            except ValueError:
                return "[批注失败] after 须为段落号（如 5 或 P5，以 docx_section_read 视图为准）"
            if not 1 <= idx <= len(paras):
                return f"[批注失败] 段落号超范围：after={idx}，本节视图共 {len(paras)} 段"
            para, where = paras[idx - 1], f"P{idx}"
        else:
            para, where = paras[-1], f"P{len(paras)}"
        # 段落直接子 run 为空（全空段/内容全裹在修订标记里）时补一个零宽锚点 run——
        # 批注锚在段尾，不往修订子树里插标记元素
        runs = para.runs or [para.add_run("")]
        doc.add_comment(runs, text=text.strip(), author=_COMMENT_AUTHOR, initials="AI")
        _atomic_save(doc, dst)
        total = len(_comment_texts(doc))
    snippet = text.strip()[:30] + ("…" if len(text.strip()) > 30 else "")
    return (
        f"[已加批注] work/{rel} {where}：「{snippet}」（本节待办批注共 {total} 条）"
        "\n待办批注收尾必须逐条向用户点名（validate_body 也会清点）；正文里不要写占位文字。"
    )


def _grid_layered(layers: list[dict]) -> tuple[list[list[dict]], list[float]]:
    """分层/组织架构网格：域头列+四成员列，层间 ↓ 箭头行，层内 >4 项折行。"""
    grid: list[list[dict]] = []
    for i, layer in enumerate(layers):
        chunks = [layer["items"][j : j + 4] for j in range(0, len(layer["items"]), 4)] or [[]]
        for k, chunk in enumerate(chunks):
            row: list[dict] = [
                {
                    "text": layer["title"] if k == 0 else "",
                    "bold": True,
                    "fill": _DIAG_HEAD_FILL,
                    "center": True,
                }
            ]
            for c in range(4):
                row.append(
                    {"text": chunk[c], "fill": _DIAG_ITEM_FILL, "center": True}
                    if c < len(chunk)
                    else {}
                )
            grid.append(row)
        if i < len(layers) - 1:
            grid.append([{}, {}, {"text": "↓", "center": True}, {}, {}])
    return grid, [0.18, 0.205, 0.205, 0.205, 0.205]


def _grid_gantt(tasks: list[dict], weeks: int, milestones: list[dict]):
    """甘特网格：任务行×周列，区间格填色即时间条（空文本+底纹），里程碑行 ◆。
    周列窄格是形态本身——列数上限不适用，由 weeks ≤36 上限管横向溢出。"""
    head_cell = {"bold": True, "fill": _DIAG_HEAD_FILL, "center": True}
    grid: list[list[dict]] = [
        [{"text": "任务", **head_cell}]
        + [{"text": f"W{w}", **head_cell} for w in range(1, weeks + 1)]
    ]
    for t in tasks:
        row = [{"text": t["name"]}]
        for w in range(1, weeks + 1):
            row.append({"fill": _DIAG_HEAD_FILL} if t["start"] <= w <= t["end"] else {})
        grid.append(row)
    for m in milestones:
        row = [{"text": m["label"]}]
        for w in range(1, weeks + 1):
            row.append({"text": "◆", "center": True} if w == m["week"] else {})
        grid.append(row)
    return grid, [0.22] + [(1 - 0.22) / weeks] * weeks


def _grid_radial(center: str, left: list[str], right: list[str], below: list[str]):
    """中心辐射网格：中列（建表后纵向合并）+左右翼+↓下方行——返回 (网格, 列宽比,
    中列行数) 供 merge。"""
    mid = max(len(left), len(right), 1)
    grid: list[list[dict]] = []
    for i in range(mid):
        lv = {"text": left[i], "fill": _DIAG_ITEM_FILL, "center": True} if i < len(left) else {}
        rv = {"text": right[i], "fill": _DIAG_ITEM_FILL, "center": True} if i < len(right) else {}
        grid.append(
            [
                lv,
                {"text": "→" if i < len(left) else "", "center": True},
                {
                    "text": center if i == 0 else "",
                    "bold": True,
                    "fill": _DIAG_HEAD_FILL,
                    "center": True,
                },
                {"text": "←" if i < len(right) else "", "center": True},
                rv,
            ]
        )
    if below:
        grid.append([{}, {}, {"text": "↓", "center": True}, {}, {}])
        for item in below:
            grid.append([{}, {}, {"text": item, "fill": _DIAG_ITEM_FILL, "center": True}, {}, {}])
    return grid, [0.26, 0.08, 0.32, 0.08, 0.26], mid


def _parse_diagram_spec(kind: str, spec: str) -> tuple[dict | list | None, str | None]:
    """图示 spec 解析与校验（各 kind 逐字段校验、超限给修复写法）。"""
    try:
        data = json.loads(spec)
    except ValueError:
        return None, "spec 不是合法 JSON（须为对象，各 kind 字段见工具说明）"
    if not isinstance(data, dict):
        return None, "spec 须为 JSON 对象"

    def _strs(v: list | None, label: str, where: str) -> tuple[list[str] | None, str | None]:
        if v is None:
            return [], None
        if not isinstance(v, list):
            return None, f"{where} 的 {label} 须为字符串数组"
        out = [_cell_str(x).strip() for x in v]
        if any(not x for x in out):
            return None, f"{where} 的 {label} 含空项"
        if any(len(x) > _MAX_CELL_CHARS for x in out):
            return None, f"{where} 的 {label} 有条目超 {_MAX_CELL_CHARS} 字上限"
        return out, None

    if kind == "layered":
        layers = data.get("layers")
        if not isinstance(layers, list) or not layers:
            return None, (
                "layered 须带 layers：[{\"title\":\"主数据域\",\"items\":[\"公司\",\"组织\",\"职位\",\"人员\"]}, …]"
            )
        norm: list[dict] = []
        for i, ly in enumerate(layers, 1):
            if not isinstance(ly, dict):
                return None, f"layers 第 {i} 项不是对象"
            title = _cell_str(ly.get("title")).strip()
            items, err = _strs(ly.get("items"), "items", f"layers 第 {i} 项")
            if err:
                return None, err
            if not title or not items:
                return None, f"layers 第 {i} 项缺 title 或 items（均必填、items 非空）"
            if len(title) > _MAX_CELL_CHARS:
                return None, f"layers 第 {i} 项 title 超 {_MAX_CELL_CHARS} 字上限"
            norm.append({"title": title, "items": items})
        return norm, None

    if kind == "gantt":
        tasks = data.get("tasks")
        if not isinstance(tasks, list) or not tasks:
            return None, (
                "gantt 须带 tasks：[{\"name\":\"需求调研\",\"start\":1,\"end\":4}, …]（周号从 1 起）"
            )
        try:
            weeks = int(data.get("weeks") or 0)
        except (TypeError, ValueError):
            return None, "weeks 须为整数（周数，可省略——按任务最晚结束周推定）"
        norm_tasks: list[dict] = []
        for i, t in enumerate(tasks, 1):
            if not isinstance(t, dict):
                return None, f"tasks 第 {i} 项不是对象"
            name = _cell_str(t.get("name")).strip()
            if not name:
                return None, f"tasks 第 {i} 项缺 name"
            try:
                start, end = int(t.get("start")), int(t.get("end"))
            except (TypeError, ValueError):
                return None, f"tasks 第 {i} 项 start/end 须为整数（周号从 1 起）"
            if start < 1 or end < start:
                return None, f"tasks 第 {i} 项区间非法（须 1 ≤ start ≤ end）"
            norm_tasks.append({"name": name, "start": start, "end": end})
            weeks = max(weeks, end)
        if weeks > 36:
            return None, (
                f"总周数 {weeks} 超上限（≤36）——长周期请以月为单位聚合"
                '（如 {"name":"需求调研","start":1,"end":2} 的 start/end 按月计）'
            )
        milestones_raw = data.get("milestones") or []
        if not isinstance(milestones_raw, list):
            return None, "milestones 须为数组：[{\"week\":4,\"label\":\"需求评审\"}, …]（可省略）"
        norm_ms: list[dict] = []
        for i, m in enumerate(milestones_raw, 1):
            if not isinstance(m, dict):
                return None, f"milestones 第 {i} 项不是对象"
            label = _cell_str(m.get("label")).strip()
            if not label:
                return None, f"milestones 第 {i} 项缺 label"
            try:
                week = int(m.get("week"))
            except (TypeError, ValueError):
                return None, f"milestones 第 {i} 项 week 须为整数"
            if not 1 <= week <= weeks:
                return None, f"milestones 第 {i} 项 week={week} 超出总周数 {weeks}"
            norm_ms.append({"week": week, "label": label})
        return {"tasks": norm_tasks, "weeks": weeks, "milestones": norm_ms}, None

    # kind == "radial"（调用方已把关 kind 取值）
    center = _cell_str(data.get("center")).strip()
    if not center:
        return None, 'radial 须带 center："人员主记录"（中心主题）'
    left, err = _strs(data.get("left"), "left", "radial")
    if err:
        return None, err
    right, err = _strs(data.get("right"), "right", "radial")
    if err:
        return None, err
    below, err = _strs(data.get("below"), "below", "radial")
    if err:
        return None, err
    if not (left or right or below):
        return None, "radial 至少给 left/right/below 之一（中心之外要有内容）"
    return {"center": center, "left": left, "right": right, "below": below}, None


# flow（分支回环流程图，mermaid 消费方 2026-09-14 批三）：模型只给 JSON 拓扑，
# 程序翻译成 mermaid 文本（模型永不手写 mermaid——非法输入当场报错的错误形态
# 优于「解析模型半自由语法猜错画错图」）；渲染走 webview 光栅化队列出 PNG。
_FLOW_MAX_NODES = 20
_FLOW_MAX_EDGES = 30
_FLOW_LABEL_MAX = 20
_FLOW_SHAPES = {"rect", "diamond", "round"}  # 矩形(默认)/判断菱形/圆角起止


def _parse_flow_spec(data: dict) -> tuple[dict | None, str | None]:
    """flow spec 校验与规范化：{"nodes":[{"id","label","shape"?}],"edges":[[from,to,label?]]}。"""
    nodes_raw = data.get("nodes")
    edges_raw = data.get("edges")
    if not isinstance(nodes_raw, list) or not nodes_raw:
        return None, 'flow 须带 nodes：[{"id":"a","label":"提交故障"}, {"id":"b","label":"判定级别","shape":"diamond"}]'
    if not isinstance(edges_raw, list) or not edges_raw:
        return None, "flow 须带 edges：[[\"a\",\"b\"],[\"b\",\"c\",\"重大\"]]（第三元素=边标签）"
    if len(nodes_raw) > _FLOW_MAX_NODES:
        return None, f"节点数 {len(nodes_raw)} 超上限（≤{_FLOW_MAX_NODES}）——拆成两张图或精简步骤"
    if len(edges_raw) > _FLOW_MAX_EDGES:
        return None, f"连线数 {len(edges_raw)} 超上限（≤{_FLOW_MAX_EDGES}）——拆成两张图"
    nodes: list[dict] = []
    ids: set[str] = set()
    for i, n in enumerate(nodes_raw, 1):
        if not isinstance(n, dict):
            return None, f"nodes 第 {i} 项不是对象"
        nid = _cell_str(n.get("id")).strip()
        label = _cell_str(n.get("label")).strip()
        shape = _cell_str(n.get("shape") or "rect").strip() or "rect"
        if not nid or not label:
            return None, f"nodes 第 {i} 项缺 id 或 label（均必填）"
        if not re.fullmatch(r"[A-Za-z0-9_\-]{1,16}", nid):
            return None, f"节点 id「{nid}」须为 ≤16 位的字母数字/下划线/连字符（edges 用它引用）"
        if nid in ids:
            return None, f"节点 id「{nid}」重复"
        if len(label) > _FLOW_LABEL_MAX:
            return None, f"节点「{nid}」label {len(label)} 字超上限（≤{_FLOW_LABEL_MAX}）"
        if shape not in _FLOW_SHAPES:
            return None, f"节点「{nid}」shape 须为 {'/'.join(sorted(_FLOW_SHAPES))}（默认 rect）"
        ids.add(nid)
        nodes.append({"id": nid, "label": label, "shape": shape})
    edges: list[list[str]] = []
    for i, e in enumerate(edges_raw, 1):
        if not isinstance(e, list) or len(e) not in (2, 3):
            return None, f"edges 第 {i} 条须为 [from, to] 或 [from, to, \"标签\"]"
        fr, to = _cell_str(e[0]).strip(), _cell_str(e[1]).strip()
        if fr not in ids or to not in ids:
            return None, f"edges 第 {i} 条引用了不存在的节点（{fr}→{to}）"
        label = _cell_str(e[2]).strip() if len(e) == 3 else ""
        if label and len(label) > _FLOW_LABEL_MAX:
            return None, f"edges 第 {i} 条标签 {len(label)} 字超上限（≤{_FLOW_LABEL_MAX}）"
        edges.append([fr, to, label])
    return {"nodes": nodes, "edges": edges}, None


def _flow_to_mermaid(parsed: dict) -> str:
    """规范化的 flow 拓扑 → mermaid flowchart 文本（纵向下行）。label 里的引号/
    竖线转成实体/全角，防语法注入——输入已过校验，这里是保险带。"""
    lines = ["flowchart TD"]

    def esc(label: str, pipe_to_full: bool) -> str:
        out = label.replace("\\", "\\\\").replace('"', "#quot;")
        return out.replace("|", "｜") if pipe_to_full else out

    for n in parsed["nodes"]:
        if n["shape"] == "diamond":
            lines.append(f'  {n["id"]}{{"{esc(n["label"], False)}"}}')
        elif n["shape"] == "round":
            lines.append(f'  {n["id"]}("{esc(n["label"], False)}")')
        else:
            lines.append(f'  {n["id"]}["{esc(n["label"], False)}"]')
    for edge in parsed["edges"]:
        fr, to = edge[0], edge[1]
        label = edge[2] if len(edge) > 2 else ""
        if label:
            lines.append(f"  {fr} -->|{esc(label, True)}| {to}")
        else:
            lines.append(f"  {fr} --> {to}")
    return "\n".join(lines)


@tool
@_tool_guard("图示")
def docx_diagram_insert(dest: str, kind: str, spec: str, after: str = "", caption: str = "") -> str:
    """在正文节插入程序生成的图示（表格拼装 + webview 渲染，2026-09-14 表格通道批）
    ——评委按图找要点比按段落找快，方案节的详实度工具。

    适用：实施进度（kind=gantt）、项目组织/团队层级/系统与数据分层架构
    （kind=layered）、以某物为中心的关系说明（kind=radial）、**带分支回环的流程图**
    （kind=flow——故障处置/审批流转/业务流程；渲染走 webview 光栅化出 PNG，模型
    只给 JSON 拓扑，程序翻译成 mermaid 文本）。**不适用**：纯叙述内容（加图不
    加分）、承诺函/函件类格式节。普通数据表（配置清单/对比矩阵）用建节 body 的
    table 块，不用本工具。按写作指引「图示」列的计划执行，计划外插图会在收尾
    对账被点名。
    Args:
        dest: 目标节文件（相对任务 work/，须已用 docx_section_create 创建）
        kind: layered（分层/组织架构）| gantt（甘特）| radial（中心辐射）|
              flow（分支回环流程图，2026-09-14 批三）
        spec: 各 kind 的 JSON 拓扑——
              layered: {"layers":[{"title":"主数据域","items":["公司","组织","职位","人员"]}, …]}
                       （组织架构=逐层 1..n 项的 layered）
              gantt:   {"tasks":[{"name":"需求调研","start":1,"end":4}, …],
                        "weeks":24, "milestones":[{"week":4,"label":"需求评审"}]}
                       （start/end 为周号从 1 起；weeks 可省略按最晚结束周推定，≤36；
                         长周期按月聚合）
              radial:  {"center":"人员主记录","left":["公司与组织","职位与任职"],
                        "right":["任免业务","干部考察"],"below":["同步批次","数据版本"]}
              flow:    {"nodes":[{"id":"a","label":"提交故障"},
                                 {"id":"b","label":"判定级别","shape":"diamond"},
                                 {"id":"c","label":"应急处置","shape":"round"}],
                        "edges":[["a","b"],["b","c","重大"],["c","b","重判"]]}
                       （shape∈rect默认/diamond判断/round起止；边第三元素=标签
                        「是/否/重大」；分支回环随意——布局归渲染器；≤20 节点≤30 边，
                        超限拆图）
        after: 插在该段落号之后（P 序号以 docx_section_read 视图为准，可带 P 前缀）；
               留空=追加到节末尾
        caption: 图示下方居中图注（如「故障处置流程」——不带编号，整本合册时
               按树序全局编号；空=不加图注）
    """
    task_id = _task_id()
    if not task_id:
        return "[图示失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, dest, must_exist=True)
    except ValueError as e:
        return f"[图示失败] {e}"
    if kind not in ("layered", "gantt", "radial", "flow"):
        return "[图示失败] kind 须为 layered|gantt|radial|flow（线框属后续批次）"
    if kind == "flow":
        try:
            data = json.loads(spec)
        except ValueError:
            return "[图示失败] spec 不是合法 JSON（须为对象，字段见工具说明）"
        if not isinstance(data, dict):
            return "[图示失败] spec 须为 JSON 对象"
        parsed_flow, err = _parse_flow_spec(data)
        if err:
            return f"[图示失败] {err}"
        mermaid = _flow_to_mermaid(parsed_flow)
        return _render_and_insert(
            task_id, dst, rel, after, caption,
            {"kind": "flow", "mermaid": mermaid}, "[图示失败]",
            f"[已插图示] flow → work/{rel}：流程图 1 张（webview 渲染，"
            f"节点 {len(parsed_flow['nodes'])}/连线 {len(parsed_flow['edges'])}）",
        )
    parsed, err = _parse_diagram_spec(kind, spec)
    if err:
        return f"[图示失败] {err}"
    with _docx_path_lock(dst):
        doc = Document(str(dst))
        if kind == "layered":
            grid, ratios = _grid_layered(parsed)
            tbl = _grid_table(doc, grid, ratios)
            stat = f"{len(parsed)} 层"
        elif kind == "gantt":
            grid, ratios = _grid_gantt(parsed["tasks"], parsed["weeks"], parsed["milestones"])
            tbl = _grid_table(doc, grid, ratios)
            stat = f"{len(parsed['tasks'])} 项任务×{parsed['weeks']} 周"
        else:
            grid, ratios, mid = _grid_radial(parsed["center"], parsed["left"], parsed["right"], parsed["below"])
            tbl = _grid_table(doc, grid, ratios)
            if mid > 1:
                tbl.cell(0, 2).merge(tbl.cell(mid - 1, 2))
            stat = f"中心「{parsed['center']}」"
        cap_el = None
        if (caption or "").strip():
            cap_el = _add_caption(doc, caption)._p
        if (after or "").strip():
            try:
                idx = int(after.strip().lstrip("Pp"))
            except ValueError:
                return "[图示失败] after 须为段落号（如 5 或 P5，以 docx_section_read 视图为准）"
            paras = doc.paragraphs
            if not 1 <= idx <= len(paras):
                return f"[图示失败] 段落号超范围：after={idx}，本节视图共 {len(paras)} 段"
            paras[idx - 1]._p.addnext(tbl._tbl)
            if cap_el is not None:
                tbl._tbl.addnext(cap_el)  # 图注紧随表格（建表/建注都在文末，须随表格一起搬）
            where = f"P{idx} 之后"
        else:
            where = "节末"
        _atomic_save(doc, dst)
        # doc.tables 每次访问重新包装 Table 对象（identity 比对必失败），按 XML
        # 元素身份找文档序号
        t_no = [t._tbl for t in doc.tables].index(tbl._tbl) + 1
    bits = f"[已插图示] {kind} → work/{rel}（{where}，表格 T{t_no}，{stat}）"
    if cap_el is not None:
        bits += (
            "，图注 1 段（图注段使其后段落序号 +1）\n（最新读视图如下——后续修订按此序号"
            "定位即可，无需再调 docx_section_read）\n" + "\n".join(view_lines(doc))
        )
    else:
        bits += (
            "。表格不占段落序号（P 编号不变）；图注须在插入时经 caption 参数给出——"
            "漏了可接受无图注，或 docx_comment_add 批注（锚定表格附近段落）请用户在 "
            "Word 补；不要再调本工具补图注（会重复出表）"
        )
    return bits


@tool
@_tool_guard("修订")
def docx_section_revise(path: str, edits: str) -> str:
    """对正文节 docx 做定向修订，全部落成 Word 原生修订标记（用户在 Word 中审阅接受/拒绝）。

    用途：素材底稿注入后的适配修订（换公司名/项目名/参数）、招标格式件填空
    （函件空栏/表格空格子）、局部新写、删除无关段。每处修订在 Word 修订视图可见
    （作者 灵燕智能）；程序落盘前自校验「拒绝全部修订后与修订前逐字一致」
    （覆盖正文段落与表格单元格），标记损坏即拒绝保存。
    Args:
        path: 目标节文件（相对任务 work/，如 body/技术部分/3.1 需求分析.docx）
        edits: JSON 数组，两类条目（不可混在同一条里）——正文段落
              {"para": 段落序号, "action": "replace|insert_after|delete",
               "find": 期望原文片段（replace 必填，防段落漂移的双重校验）,
               "text": 新文本, "style": 样式名（仅 insert_after 可选，如
               "Tender Cover"/"Tender Cover Sub"——封面等非正文段落挂指定
               版式；样式不存在回落正文样式并在返回中注记）}；表格单元格
               {"table": T序号, "row": 行, "col": 列,
               "action": "replace|fill", "find": 格内期望原文（replace 必填，须
               完整出现在同一格内段落）, "text": 新文本}——fill 仅限空格子
               （接受视角为空，填空动作）。replace 只替换首处命中；同段多条按
               接受视角顺序定位。多条段落编辑按序号从大到小应用（防位移），
               表格格编辑不改段落数、另批处理。表格插行/删行、嵌套表格不支持
               （在 Word 中改）。
    返回：逐条结果（每处修订的定位与落点摘要）。删段/替换/填空**不改变段落
               序号**（删除段原位保留删除标记，视图仍给原序号）；插段使其后
               段落序号 +1 位移——含插段的批次返回直接附最新读视图（段落序号
               已按插段后计），后续修订按返回定位即可，**无需再调 docx_section_read
               回读确认**（逐条结果即确认）。
    """
    task_id = _task_id()
    if not task_id:
        return "[修订失败] 当前会话未归属任务"
    try:
        dst, rel = _dest_path(task_id, path, must_exist=True)
    except ValueError as e:
        return f"[修订失败] {e}"
    try:
        items = json.loads(edits)
    except ValueError:
        return "[修订失败] edits 不是合法 JSON"
    if not isinstance(items, list) or not items:
        return "[修订失败] edits 须为非空 JSON 数组"
    errs: list[str] = []
    para_items: list[dict] = []
    cell_items: list[dict] = []
    for it in items:
        if not isinstance(it, dict):
            errs.append("每项须为对象")
            continue
        has_para = isinstance(it.get("para"), int) and not isinstance(it.get("para"), bool)
        has_cell = all(
            isinstance(it.get(k), int) and not isinstance(it.get(k), bool)
            for k in ("table", "row", "col")
        )
        if has_para and has_cell:
            errs.append("para 与 table/row/col 不可同时传")
            continue
        if has_para:
            act = it.get("action")
            if act not in ("replace", "insert_after", "delete"):
                errs.append(f"P{it.get('para')}：action 须为 replace|insert_after|delete")
            elif act == "replace" and (not it.get("find") or it.get("text") is None):
                errs.append(f"P{it.get('para')}：replace 需 find 与 text")
            elif act == "insert_after" and not it.get("text"):
                errs.append(f"P{it.get('para')}：insert_after 需 text")
            else:
                para_items.append(it)
        elif has_cell:
            act = it.get("action")
            where = f"T{it.get('table')}({it.get('row')},{it.get('col')})"
            if act not in ("replace", "fill"):
                errs.append(f"{where}：action 须为 replace|fill")
            elif act == "replace" and (not it.get("find") or it.get("text") is None):
                errs.append(f"{where}：replace 需 find 与 text")
            elif act == "fill" and not it.get("text"):
                errs.append(f"{where}：fill 需 text")
            else:
                cell_items.append(it)
        else:
            errs.append("每项须含整数 para（正文段落）或 table/row/col 三件套（表格单元格）")
    if errs:
        return "[修订失败] " + "；".join(errs)
    with _docx_path_lock(dst):
        doc = Document(str(dst))
        paras = doc.paragraphs
        before = _flatten_rejected(doc)
        applied = 0
        body_style = _body_style(doc)
        body_style_id = body_style.element.get(qn("w:styleId")) if body_style else None
        # 逐条结果行（2026-09-13 回读收敛：返回自带确认与新序号，模型不再回读视图——
        # 99 次 revise 里 76 次紧跟整读、53 批含删/插段全是重锚定刚需）
        done: list[str | None] = []
        inserts_pending: list[tuple[int, object, str, int]] = []  # (原序号, 新段元素, 文本, done 下标)
        # 改动后现文（2026-09-14 回读再收敛）：replace/delete 的段落序号 + 格编辑坐标
        # ——纯改动批（无插段）返回逐对象「现文」，把回读动机从根上拿掉（实测写手
        # 5.5 次整读/节，大部分是改完不放心再看一眼）
        touched_paras: list[int] = []  # replace/delete 过的段序号
        touched_cells: list[tuple[int, int, int]] = []  # 格编辑 (T,R,C)

        def _snip(s: str, n: int = 20) -> str:
            s = s.replace("\n", " ")
            return s[:n] + ("…" if len(s) > n else "")
        # 段落从后往前：插入/删除改变段落数，先改大序号防位移；同段连续
        # insert_after 以「上次插入的新段」为锚保提交顺序（游标按段序号记）
        insert_anchors: dict[int, object] = {}
        style_notes: set[str] = set()  # 指定了但版式里不存在的样式名（回落正文样式）
        for it in sorted(para_items, key=lambda x: x["para"], reverse=True):
            i = it["para"]
            if i < 1 or i > len(paras):
                return f"[修订失败] 段落序号 {i} 超出范围（现共 {len(paras)} 段；序号以 docx_section_read 为准）"
            para = paras[i - 1]
            rev_id = _next_rev_id(doc)
            act = it["action"]
            try:
                if act == "replace":
                    _tracked_replace(para, str(it["find"]), str(it["text"]), rev_id)
                    done.append(
                        f"P{i} 替换「{_snip(str(it['find']))}」→「{_snip(str(it['text']))}」"
                    )
                    touched_paras.append(i)
                elif act == "insert_after":
                    anchor = insert_anchors.get(i, para._p)
                    style_id = body_style_id
                    style_name = it.get("style")
                    if isinstance(style_name, str) and style_name.strip():
                        resolved = _style_id_by_name(doc, style_name.strip())
                        if resolved is None:
                            style_notes.add(style_name.strip())
                        else:
                            style_id = resolved
                    new_el = _tracked_insert_after(
                        anchor, str(it["text"]), rev_id, style_id=style_id
                    )
                    insert_anchors[i] = new_el
                    inserts_pending.append((i, new_el, str(it["text"]), len(done)))
                    done.append(None)  # 新序号待段落批结束后回填（删/插交互后的最终序号）
                else:
                    _tracked_delete(para, rev_id)
                    done.append(f"P{i} 删除（原位保留删除标记，序号不变）")
                    touched_paras.append(i)
            except ValueError as e:
                return f"[修订失败] P{i}：{e}"
            applied += 1
        # 表格单元格（格编辑不改段落数，与段落批寻址互不影响；嵌套表格内容不达——
        # find 定位不到会明确报错）
        for it in cell_items:
            t_no, r, c = it["table"], it["row"], it["col"]
            if t_no < 1 or t_no > len(doc.tables):
                return f"[修订失败] 表格序号 {t_no} 超出范围（现共 {len(doc.tables)} 张；序号以 docx_section_read 为准）"
            tbl = doc.tables[t_no - 1]
            if r < 1 or r > len(tbl.rows) or c < 1 or c > len(tbl.columns):
                return (
                    f"[修订失败] T{t_no} 单元格坐标 ({r},{c}) 超出范围"
                    f"（{len(tbl.rows)}行×{len(tbl.columns)}列；坐标以视图 R行C列 为准）"
                )
            cell_paras = tbl.cell(r - 1, c - 1).paragraphs
            rev_id = _next_rev_id(doc)
            if it["action"] == "fill":
                if any(_accepted_text(p._p).strip() for p in cell_paras):
                    return (
                        f"[修订失败] T{t_no}({r},{c}) 非空格子不能用 fill——"
                        "read 视图取该格原文片段改用 replace"
                    )
                _tracked_fill(cell_paras[0], str(it["text"]), rev_id)
                done.append(f"T{t_no} R{r}C{c} 填空「{_snip(str(it['text']))}」")
                touched_cells.append((t_no, r, c))
            else:
                find = str(it["find"])
                target = next((p for p in cell_paras if find in _accepted_text(p._p)), None)
                if target is None:
                    return (
                        f"[修订失败] T{t_no}({r},{c}) 格内未找到期望文本「{find[:60]}」"
                        "（find 须完整出现在同一格内段落）"
                    )
                try:
                    _tracked_replace(target, find, str(it["text"]), rev_id)
                    done.append(
                        f"T{t_no} R{r}C{c} 替换「{_snip(find)}」→「{_snip(str(it['text']))}」"
                    )
                    touched_cells.append((t_no, r, c))
                except ValueError as e:
                    return f"[修订失败] T{t_no}({r},{c})：{e}"
            applied += 1
        # 插段最终序号回填：删段不挪元素（序号不变）、插段新增元素——段落批全部
        # 落完后的 doc.paragraphs 顺序即最终序号（表格格编辑不动段落数）
        if inserts_pending:
            final_idx = {id(p._p): n for n, p in enumerate(doc.paragraphs, 1)}
            for orig_i, el, text, di in inserts_pending:
                n = final_idx.get(id(el))
                done[di] = (
                    f"P{n} 新段（插在原 P{orig_i} 后）「{_snip(text)}」"
                    if n
                    else f"原 P{orig_i} 后插段「{_snip(text)}」"
                )
        if _flatten_rejected(doc) != before:
            return "[修订失败] 修订标记一致性自校验未通过（未保存）——请重读视图核对序号与原文后重试"
        _atomic_save(doc, dst)
    bits = (
        f"[已修订] work/{rel}：{applied} 处已落成 Word 修订标记（作者 {_AUTHOR}），"
        f"在 Word 中审阅可逐条接受/拒绝：\n  " + "\n  ".join(str(x) for x in done)
    )
    if style_notes:
        bits += f"（注：样式 {'、'.join(sorted(style_notes))} 不存在，该段已用正文样式）"
    if inserts_pending:
        # 插段使其后序号 +1：直接附最新视图，模型无需再花一轮回读重锚定
        bits += (
            "\n（本批含插段，其后段落序号已位移；最新读视图如下——后续修订按此序号"
            "定位即可，无需再调 docx_section_read）\n" + "\n".join(view_lines(doc))
        )
    else:
        # 纯改动批（无插段，序号全部不变）：逐对象附「改动后现文」（接受视角）——
        # 回读的动机是看结果，结果直接送到眼前（2026-09-14 回读再收敛；实测 155 次
        # 写手整读里约 120 次是改完确认，5.5 次/节对设计「通读一次」）
        now_parts: list[str] = []
        for i in dict.fromkeys(touched_paras):  # 去重保序
            para = doc.paragraphs[i - 1]
            if para is None:
                continue
            # 整段删除判据=段落标记删除（pPr/rPr/w:del，_tracked_delete 落的形
            # 态、与 view_lines 的 para_deleted 同款）——不能用「段内有 w:del」
            # 判：replace 的旧文也包 run 级 w:del，会误判
            ppr = para._p.find(qn("w:pPr"))
            rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
            para_deleted = rpr is not None and rpr.find(qn("w:del")) is not None
            if para_deleted:
                now_parts.append(f"P{i} 现文：（已标记删除，接受修订后此段消失）")
            else:
                txt = _accepted_text(para._p).replace("\n", " ")
                now_parts.append(f"P{i} 现文：{_snip(txt, _VIEW_TEXT_LIMIT)}")
        for t_no, r, c in dict.fromkeys(touched_cells):
            cell = doc.tables[t_no - 1].cell(r - 1, c - 1)
            txt = " / ".join(s for p in cell.paragraphs if (s := _accepted_text(p._p).strip()))
            now_parts.append(f"T{t_no}R{r}C{c} 现文：{_snip(txt or '（空）', 60)}")
        bits += (
            "\n（删段/替换/填空均不改变段落序号——删除段原位保留删除标记、视图序号"
            "不变，无需回读视图确认。修订后建议 check_name_residue 扫旧名残留。）"
        )
        if now_parts:
            bits += "\n改动后现文（接受修订视角，无需回读确认）：\n  " + "\n  ".join(now_parts)
    return bits


# ---------- 整本合册与文本抽取 ----------

def _tree_nodes(nodes: list[dict], depth: int = 1):
    """目录树先序遍历产出 (深度, 标题, 是否容器, 交付形态, 原节点)——合册的章序
    真值是树序，不是文件名序；原节点供空容器抑制按对象身份对位。"""
    for n in nodes:
        children = n.get("children") or []
        title = str(n.get("目录名称") or "").strip()
        mode = str(n.get("交付形态") or "").strip()
        if title:
            yield depth, title, bool(children), mode, n
        yield from _tree_nodes(children, depth + 1)


def _leaf_will_emit(node: dict, wroot: Path, vol_dir: str) -> bool:
    """叶子是否会产出整本内容：需正文的（含缺文件——标题照发、缺失另有点名）恒
    真；模板填充类以节文件存在为准（含格式件与附件壳节）。与合册主循环同一判定。"""
    children = node.get("children") or []
    if children:
        return _subtree_will_emit(children, wroot, vol_dir)
    title = str(node.get("目录名称") or "").strip()
    if not title:
        return False
    if body_contract.sanitize_name(title) == _TOC_NODE_NAME:
        return True  # 目录节点恒产出（无手写文件时合册机械生成目录页）
    mode = str(node.get("交付形态") or "").strip()
    if mode not in body_contract.NON_PROSE_DELIVERY:
        return True
    stem = body_contract.sanitize_name(title)
    rel = f"body/{vol_dir}/{stem}.docx" if vol_dir else f"body/{stem}.docx"
    return (wroot / rel).is_file()


def _subtree_will_emit(nodes: list[dict], wroot: Path, vol_dir: str) -> bool:
    """容器子树是否还有将产出的叶子——全无则抑制容器章标题（子节点全走「按附件
    对待」被跳过，信息由目录页「另附」行承载；2026-09-14 拍板，兜底防线）。"""
    return any(_leaf_will_emit(n, wroot, vol_dir) for n in nodes)


def _empty_containers(nodes: list[dict], wroot: Path, vol_dir: str) -> set[int]:
    """收集子树无任何将产出叶子的容器节点（id() 集，主循环按对象身份跳过）。"""
    out: set[int] = set()
    for n in nodes:
        children = n.get("children") or []
        if children:
            out |= _empty_containers(children, wroot, vol_dir)
            if not _subtree_will_emit(children, wroot, vol_dir):
                out.add(id(n))
    return out


def _body_tail(out: Document):
    """正文当前最后一个内容元素（sectPr 前）——机械目录页的插入锚点；无内容返回 None。"""
    for el in reversed(list(out.element.body.iterchildren())):
        if el.tag != qn("w:sectPr"):
            return el
    return None


def _insert_mechanical_toc(out: Document, anchor, entries: list[tuple[str, str, int]]) -> int:
    """机械目录页（2026-09-14 结构缺口批）：Word 目录域 + 缓存静态清单。

    域包裹缓存条目（fldChar separate…end 之间）——Word/WPS 更新域即得带页码的
    正式目录并整体替换缓存；不更新也能交（缓存清单即目录）。条目=实收章节
    （编号+标题，按层级缩进、无页码——页码排版后只有 Word 知道）；未产出附件
    节在条目文本上带「（另附）」标记。anchor=目录节点位置前最后的内容元素
    （None=插到正文最前）；标题「目录」刻意不用 Heading 样式——进大纲会被
    目录域/导航窗格自引用。
    """
    if not entries:
        return 0
    els = []
    h = out.add_paragraph("目录")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(12)
    r = h.runs[0]
    r.bold = True
    r.font.size = Pt(16)
    els.append(h._p)
    for i, (prefix, title, depth) in enumerate(entries):
        p = out.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74 * max(0, depth - 1))
        if i == 0:  # 域头：begin + 指令 + separate（缓存结果由此起，更新域整体替换）
            r = p.add_run()
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            r._r.append(fc)
            r = p.add_run()
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            r._r.append(it)
            r = p.add_run()
            fs = OxmlElement("w:fldChar")
            fs.set(qn("w:fldCharType"), "separate")
            r._r.append(fs)
        p.add_run((prefix + title).strip())
        if i == len(entries) - 1:
            r = p.add_run()
            fe = OxmlElement("w:fldChar")
            fe.set(qn("w:fldCharType"), "end")
            r._r.append(fe)
        els.append(p._p)
    if anchor is not None:
        for el in reversed(els):
            anchor.addnext(el)
    else:
        sect = out.element.body.find(qn("w:sectPr"))
        for el in els:
            if sect is not None:
                sect.addprevious(el)
            else:
                out.element.body.append(el)
    return len(entries)


# ---------- 章节编号（合册按树序生成；格式取目录产物 numbering 字段） ----------

_CN_DIGITS = "零一二三四五六七八九"

# 目录节点名自带编号的前缀形态（合册会再加程序编号 → 双重编号；探测点名请用户
# 裁决，提示不是门禁）。数字形态限一两位+空格，避开「2026 年度」类年份误报。
_SELF_NUMBERED = re.compile(
    r"^(第([一二三四五六七八九十百]+|\d+)章|[一二三四五六七八九十]+、"
    r"|（[一二三四五六七八九十]+）|\d{1,2}(\.\d+)*[ 　])"
)


def _cn_num(n: int) -> str:
    """1–99 → 中文数字（十一、二十一…）；超范围回落阿拉伯数字（防御，章节数到不了）。"""
    if not 1 <= n <= 99:
        return str(n)
    if n < 10:
        return _CN_DIGITS[n]
    tens, ones = divmod(n, 10)
    head = "十" if tens == 1 else _CN_DIGITS[tens] + "十"
    return head + (_CN_DIGITS[ones] if ones else "")


class _HeadingNumberer:
    """合册章节编号器：编号=树位置的纯函数（每册一个实例，各册从首章重起）。

    编号写进标题文本，不走 Word 样式绑定自动编号——素材拷入的标题段会被 Word
    一起计数打乱章序、docx-preview 对经典绑定写法不渲染、节文件单看永远「第一章」。
    chapter=第X章+1.1（一级全角空格接题名）、decimal=1+1.1、gov=一、（一）1.、
    none=不加；封面节点不调用 prefix（不占序）。
    """

    def __init__(self, scheme: str):
        self.scheme = scheme if scheme in ("chapter", "decimal", "gov", "none") else "chapter"
        self.counters = [0] * 10  # counters[depth]：各深度当前计数（下标 1 起）

    def prefix(self, depth: int) -> str:
        """发一个 depth 级标题的编号前缀（消费一个序号并重置更深计数）。"""
        if self.scheme == "none" or not 1 <= depth <= 9:
            return ""
        self.counters[depth] += 1
        for d in range(depth + 1, 10):
            self.counters[d] = 0
        if self.scheme == "gov":
            if depth == 1:
                return f"{_cn_num(self.counters[1])}、"
            if depth == 2:
                return f"（{_cn_num(self.counters[2])}）"
            return f"{self.counters[depth]}."
        if self.scheme == "chapter" and depth == 1:
            return f"第{_cn_num(self.counters[1])}章\u3000"
        return ".".join(str(self.counters[d]) for d in range(1, depth + 1)) + " "


_CAP_NUM_STRIP_RE = re.compile(r"^[图表]\s*[0-9０-９\-－—–.．]+\s*")


def _caption_style_ids(doc: Document) -> set[str]:
    """样式名 Tender Caption 的 styleId 集（合册按样式名识别图注/表题；版式缺
    该样式=建注时回落正文样式，识别不中=跳过重编号，不误伤）。"""
    ids = set()
    for st in doc.styles:
        sid = getattr(st, "style_id", None)
        if sid and st.name == _CAPTION_STYLE_NAME:
            ids.add(sid)
    return ids


def _renumber_captions(
    copied: list, cap_ids: set[str], chapter: int, counters: dict, numbered: bool
) -> int:
    """节内图注/表题重编号（合册全局序的节级增量，2026-09-14 批）：「表」=后随
    表格的图注段（表题在表格上方——建节 body 表块形态）、「图」=其余（docx_
    diagram_insert 的图注在表格下方，表格拼装但语义为图）。编号与章节标题同一
    树序原则：节文件不带编号（树位置合册才知道）；numbered=False（目录产物
    不编号档）退化为全册平铺「图 N/表 N」。返回重编号数。"""
    n = 0
    for i, el in enumerate(copied):
        if el.tag != qn("w:p"):
            continue
        ppr = el.find(qn("w:pPr"))
        pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if pstyle is None or (pstyle.get(qn("w:val")) or "") not in cap_ids:
            continue
        nxt = copied[i + 1] if i + 1 < len(copied) else None
        kind = "表" if nxt is not None and nxt.tag == qn("w:tbl") else "图"
        key = (kind, chapter if numbered else 0)
        counters[key] = counters.get(key, 0) + 1
        label = f"{kind} {chapter}-{counters[key]}" if numbered else f"{kind} {counters[key]}"
        ts = el.findall(".//" + qn("w:t"))
        if not ts:
            continue
        text = "".join(t.text or "" for t in ts)
        ts[0].text = f"{label} {_CAP_NUM_STRIP_RE.sub('', text).strip()}".strip()
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            t.text = ""
        n += 1
    return n


def _is_title_para(p_el, title: str) -> bool:
    """节文件首段是否为其自带标题段（合册时跳过——标题由合册器按树深重发）。

    判据统一为「样式非空且接受视角文本=节标题」：只看 Heading1 会把注入素材
    自带的章标题段（Heading1、文本≠节标题）静默剥掉——素材内容丢失；宁可
    保守漏剥（首段与树标题写法有出入时合册重复发标题，可见可修）。
    """
    ppr = p_el.find(qn("w:pPr"))
    pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
    style_val = (pstyle.get(qn("w:val")) or "") if pstyle is not None else ""
    return bool(style_val) and _accepted_text(p_el).strip() == title.strip()


def section_lines_labeled(doc: Document) -> list[tuple[str, str]]:
    """(定位标签, 文本) 序列：段落 P{i}、表格行 T{t}R{r}。

    行集与 section_text_lines 同源同序——标签供 validate_body 的占位/残留提示
    定位（与 docx_section_read 视图互查，表格行不再被标成越界的 P 号）。
    合并单元格的行文本含重复格文本（python-docx row.cells 口径），只影响
    可读性不影响校验。
    """
    out: list[tuple[str, str]] = []
    for i, p in enumerate(doc.paragraphs, 1):
        out.append((f"P{i}", _accepted_text(p._p)))
    for ti, t in enumerate(doc.tables, 1):
        for ri, row in enumerate(t.rows, 1):
            cells = ["".join(_accepted_text(p._p) for p in c.paragraphs) for c in row.cells]
            out.append((f"T{ti}R{ri}", " | ".join(cells)))
    return out


def section_text_lines(doc: Document) -> list[str]:
    """docx 正文文本行（**接受全部修订后的终稿视角**——评委最终看到的文本）。

    段落逐行 + 表格每行单元格拼接。行序与 docx_section_read 视图的 P 编号
    **严格一致**：段落标记被标删的段（整段删除修订）以空行占位——空行对
    占位清点/残留扫描/素材使用率/承诺比对全部无操作，而 P 段号不漂移（校验
    提示里的定位与视图可互查）。插入修订算在内（终稿存在）、删除修订的内容
    不算（终稿消失）。
    """
    return [text for _label, text in section_lines_labeled(doc)]


def section_comments_labeled(doc: Document) -> list[tuple[str, str]]:
    """(锚定段落 P 标签, 批注文本)——待办批注的定位清点（validate_body 消费）。

    只扫正文级段落（commentRangeStart 是段落直接子元素）；锚在表格格内或跨段的
    批注落在「?」兜底行——宁可位置不明也不漏清点。
    """
    texts = _comment_texts(doc)
    if not texts:
        return []
    out: list[tuple[str, str]] = []
    anchored: set[int] = set()
    for i, para in enumerate(doc.paragraphs, 1):
        for m in para._p.findall(qn("w:commentRangeStart")):
            raw = m.get(qn("w:id")) or ""
            if not raw.isdigit():
                continue
            cid = int(raw)
            anchored.add(cid)
            t = texts.get(cid)
            if t:
                out.append((f"P{i}", t))
    out.extend(("?", t) for cid, t in texts.items() if cid not in anchored)
    return out


def head_text(path: str | Path, max_chars: int = 200) -> str:
    """docx 开头文本（接受修订视角）：前几段/表格行非空文本拼到 max_chars 截断。

    派发拼装的兄弟节摘要用（dispatch_enrich）——给并发写手看兄弟节开头、
    防跨节重复表达。文件损坏由调用方决定跳过（本函数不吞异常）。
    """
    doc = Document(str(path))
    parts: list[str] = []
    used = 0
    for text in section_text_lines(doc):
        t = text.strip()
        if not t:
            continue
        parts.append(t)
        used += len(t)
        if used >= max_chars:
            break
    return "".join(parts)[:max_chars]


def _publish_volume(task_id: str, dst: Path, vol: str, merged: int, images: int, comments: int) -> str:
    """合册产出机械发布为 tender.volume 产物（聊天产物卡/面板交付入口）。

    发布失败不影响合册结果（文件仍在 work/body/，重跑合册可补发布）；发布
    走 publish.publish_file_artifact——按册名复用身份、docx 内容级去重。
    """
    try:
        ctx = runctx.current_run()
        data = dst.read_bytes()
        meta = publish.publish_file_artifact(
            "tender.volume/tender-volume-docx@1",
            file_path=dst,
            content_meta={
                "filename": dst.name,
                "book": vol,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
                "merged_sections": merged,
                "images": images,
                "comments": comments,
            },
            display_name=vol,
            source={
                "skill": "tender-body",
                "thread_id": ctx.conversation_id if ctx else None,
                "run_id": ctx.run_id if ctx else None,
            },
            task_id=task_id,
            conversation_id=ctx.conversation_id if ctx else None,
        )
        if meta.get("_unchanged"):
            return f"{vol}：整本内容与已发布版本一致，未重复发布（{meta.get('artifact_id')}）"
        return f"{vol}：已发布为整本标书成果 {meta.get('artifact_id')}（聊天产物卡可预览/下载）"
    except Exception as e:
        return f"⚠️ {vol}：整本产物发布失败（{type(e).__name__}: {e}）——文件仍在 work/body/，重新合册可补发布"


@tool
@_tool_guard("合册")
def docx_assemble_volume() -> str:
    """按投标目录树序把正文节 docx 合册成整本文件（每册一个，tender-body 收尾必调）。

    用途：全部节完成后调用，产出 work/body/整本-<册名>.docx。容器章由本工具发
    标题（层级随树深，Word 可自动生成目录）、叶子节内容元素级拷入（图片/表格
    原样保留）、一级章前分页、页脚页码；模板填充类叶子（目录标非正文）**产出
    节文件即按树序并入**（拷原件填空的格式件本就是标书组成部分）、未产出的按
    附件对待不占整本位（返回行点名）。章节编号按树序自动生成（第一章/1.1；
    格式取目录产物的 numbering 字段，缺省第X章+1.1；封面不占序，目录产物里
    可改为 1+1.1/一、（一）/不编号）；树含「目录」节点时其前的非封面节点为
    前置区不占章号（编制索引等前置页），正文从目录后第一章起编。「目录」节点
    无节文件→**目录页由本工具机械生成**（Word 目录域+缓存静态清单：预览即见
    清单，Word/WPS 更新域即得带页码正式目录；未产出附件节标「另附」）；有节
    文件→并入并对账（手写优先）。子树无任何将产出叶子的容器不发章标题（空壳
    章抑制，信息走目录页「另附」行）。
    **整本=交付态**：节内修订标记并入时按「接受全部修订」压平（未接受修订
    前不再新旧内容并存；节文件保留修订供审阅，改内容回节文件层改再重合册）；
    拷入的树外标题段（节内小标题/素材自带标题）摘出大纲层级——导航窗格与
    自动目录只剩章节骨架，视觉样式不变；「目录」节并入后与实收章节机械对账，
    不符（列了没有的/漏了实有的）点名提醒。
    树首节点为「封面」时（tender-outline 的结构约定）整本首页即封面页：跳过册名
    大标题与封面节点自身标题、开「首页不同」（封面页不带页眉页脚，页码从封面
    后一页起显示）。缺失的正文节文件逐个点名，不中断其余节。**整本是派生产物**：改内容
    回节文件层改（或让模型改）再重新调用本工具覆盖合册，直接手改整本会被下次
    合册覆盖。**合册成功即自动发布为整本标书产物**（每册一条；docx 内容未变不
    重复发布）。
    """
    task_id = _task_id()
    if not task_id:
        return "[合册失败] 当前会话未归属任务"
    _, content, _ = body_contract.load_directory(task_id)
    if content is None:
        return "[合册失败] 无投标目录产物（或内容不可读）——合册按目录树排序，请先 tender-outline 生成目录"
    docs = [d for d in content.get("response_documents") or [] if d.get("directory")]
    if not docs:
        return "[合册失败] 目录产物没有响应文件树"
    multi = body_contract.multi_volume(content)
    numberer_scheme = str(content.get("numbering") or "").strip() or "chapter"
    wroot = work_dir(task_id).resolve()
    consumed: set[str] = set()
    reports: list[str] = []
    for doc in docs:
        vol = str(doc.get("name") or "").strip() or "主册"
        vol_dir = body_contract.sanitize_name(vol) if multi else ""
        nodes = list(_tree_nodes(doc.get("directory") or []))
        numberer = _HeadingNumberer(numberer_scheme)  # 各册从首章重起
        self_numbered: list[str] = []  # 节点名自带编号（合册再编号会双重，探测点名）
        # 封面=树首一级叶子且清洗后标题为「封面」（tender-outline 的结构约定）：
        # 整本首页即封面页——跳过册名大标题（封面自带册名）与节点自身标题，
        # 开「首页不同」让封面页不带页眉页脚；无封面时行为与旧版逐字节一致。
        has_cover = (
            bool(nodes)
            and nodes[0][0] == 1
            and not nodes[0][2]
            and body_contract.sanitize_name(nodes[0][1]) == _COVER_NODE_NAME
        )
        # 目录页节点=首个名为「目录」的一级叶子（tender-outline 结构约定，2026-09-14
        # 批）：无手写节文件→循环后在此位置机械生成目录页（Word 目录域+缓存清单）；
        # 有→按普通节并入（toc_lines 对账路径，手写优先）。其前的非封面一级节点=
        # 前置区（编制索引等前置页）不占章号，正文从目录后第一章起编；树无目录
        # 节点时维持全编号（旧目录产物兼容）。
        toc_node_idx = next(
            (
                i
                for i, (d, t, c, _m, _n) in enumerate(nodes)
                if d == 1 and not c and body_contract.sanitize_name(t) == _TOC_NODE_NAME
            ),
            None,
        )
        empty_ids = _empty_containers(doc.get("directory") or [], wroot, vol_dir)
        out = _new_document()  # 模板自带 A4 版面/页边距/页脚页码，无需再手拼
        if has_cover:
            out.sections[0].different_first_page_header_footer = True
        else:
            out.add_heading(vol, 0)
        merged = n_img = n_comment = n_cap = 0
        cap_counters: dict = {}  # (表|图, 章) → 章内序——图注/表题全局重编号
        missing: list[str] = []
        unfilled: list[tuple[str, int]] = []  # 模板填充类叶子未产出节文件（按附件对待兜底）
        placeholders: list[str] = []  # 内联占位兜底扫描命中（正规落点=批注，此为防线）
        issued_titles: list[str] = []  # 本册实际发出标题的节点（目录页对账的「实有」侧）
        toc_entries: list[tuple[str, str, int]] = []  # 目录页条目 (编号前缀, 标题, 深度)——树序
        toc_pending = False  # 目录节点无手写节文件：循环后机械回填目录页
        toc_anchor = None  # 目录节点位置前最后的内容元素（插入锚点；None=正文最前）
        toc_lines: list[str] | None = None  # 「目录」节条目行（条目超长=说明文字，不当条目）
        seen_chapter = False
        for idx, (depth, title, is_container, mode, node) in enumerate(nodes):
            if id(node) in empty_ids:
                continue  # 空容器：子树无将产出叶子（壳节漏建等极端情形）——不发章标题，信息走目录页「另附」
            if _SELF_NUMBERED.match(title):
                self_numbered.append(title)
            stem = body_contract.sanitize_name(title)
            rel_src = f"body/{vol_dir}/{stem}.docx" if vol_dir else f"body/{stem}.docx"
            is_toc_node = idx == toc_node_idx
            if not is_container and mode in body_contract.NON_PROSE_DELIVERY:
                if not (wroot / rel_src).is_file():
                    if is_toc_node:
                        # 目录页节点无手写节文件：记锚点，循环后在此位置机械生成目录页
                        toc_pending = True
                        toc_anchor = _body_tail(out)
                        continue
                    unfilled.append((title, depth))
                    toc_entries.append(("", f"{title}（另附）", depth))
                    continue  # 模板填充未产出：按附件对待（壳节漏建的兜底），不占整本位
                # 已产出（拷原件+填空/附件壳节）：按树序并入整本，与正文叶子同路
            if has_cover and idx == 0:
                # 封面节点不发「封面」标题行（节文件内容即整页），但占一级位：
                # seen_chapter 置位让第一章拿到分页、封面独占首页；编号不占序
                seen_chapter = True
            else:
                # 标题带编号发（对账剥节文件标题仍用裸 title，见 _is_title_para 调用处）；
                # 前置区（目录节点前的非封面节点）与目录页自身不占章号——正文从
                # 目录后第一章起编（2026-09-14 拍板：编制索引等前置页不带编号）
                front_matter = toc_node_idx is not None and idx < toc_node_idx
                prefix = "" if front_matter or is_toc_node else numberer.prefix(depth)
                h = out.add_heading(prefix + title, min(depth, 9))
                if not is_toc_node:
                    issued_titles.append(title)
                    toc_entries.append((prefix, title, depth))
                if depth == 1:
                    if seen_chapter:
                        h.paragraph_format.page_break_before = True
                    seen_chapter = True
            if is_container:
                continue
            src_path = wroot / rel_src
            if not src_path.is_file():
                missing.append(title)
                continue
            consumed.add(rel_src)
            src = Document(str(src_path))
            # 内联占位兜底：接受视角逐行扫（整段删除修订不误报）——正规待办落点是
            # docx_comment_add 批注，正文占位文字会进交付稿，点名逼改
            for label, text in section_lines_labeled(src):
                for mark in body_contract.PLACEHOLDER_MARKS:
                    if mark in text:
                        i = text.find(mark)
                        placeholders.append(f"{title}({label})：{text[i : i + 24].strip()}")
                        break
            children = [
                c for c in src.element.body.iterchildren()
                if c.tag.split("}")[-1] in ("p", "tbl")
            ]
            if children and children[0].tag.split("}")[-1] == "p" and _is_title_para(children[0], title):
                children = children[1:]
            if toc_lines is None and _toc_normalize(title) == "目录":
                # 目录页条目行（非空、≤60 字——更长的行是说明文字不是条目）
                toc_lines = [t for t in section_text_lines(src) if t.strip() and len(t.strip()) <= 60]
            copied: list = []
            for el in children:
                new_el = deepcopy(el)
                _strip_inner_sectpr(new_el)
                _strip_copy_residue(new_el)
                if not _accept_revisions_inplace(new_el):
                    continue  # 整段删除修订：接受后不存在
                n_img += _migrate_images(src, out, new_el)
                copied.append(new_el)
                sect = out.element.body.find(qn("w:sectPr"))
                if sect is not None:
                    sect.addprevious(new_el)
                else:
                    out.element.body.append(new_el)
            _style_n, style_remap = _merge_missing_styles(src, out, copied)
            _merge_missing_numbering(src, out, copied, style_id_remap=style_remap)
            _demote_extra_headings(out, copied)  # 树外标题摘出大纲（样式迁完再判——导航只剩骨架）
            # 图注/表题全局重编号（先迁样式再改文本——识别按 pStyle=样式名解析）
            cap_ids = _caption_style_ids(src)
            if cap_ids:
                n_cap += _renumber_captions(
                    copied, cap_ids, numberer.counters[1], cap_counters,
                    numberer_scheme != "none",
                )
            n_comment += _merge_missing_comments(src, out, copied)
            merged += 1
        if merged == 0:
            detail = f"（缺失：{'、'.join(missing)}）" if missing else "（目录树无叶子节点）"
            reports.append(f"{vol}：无已写节文件{detail}，未产出整本")
            continue
        if toc_pending:
            # 机械目录页回填到目录节点位置（Word 目录域+缓存清单——更新域得页码）
            n_toc = _insert_mechanical_toc(out, toc_anchor, toc_entries)
            if n_toc:
                bits = f"{vol}：目录页已生成（{n_toc} 条；Word/WPS 中更新目录域可得带页码目录）"
                if unfilled:
                    bits += f"；未产出附件节标「另附」{len(unfilled)} 项"
                reports.append(bits)
        try:
            dst, rel = _dest_path(
                task_id, f"body/{body_contract.VOLUME_PREFIX}{body_contract.sanitize_name(vol)}.docx", must_exist=False
            )
        except ValueError as e:
            return f"[合册失败] {e}"
        with _docx_path_lock(dst):
            dst.parent.mkdir(parents=True, exist_ok=True)
            # 原子替换：整本是派生产物，重跑直接覆盖（恢复语义在节文件层）；
            # _atomic_save 的 uuid 后缀 tmp 防并发合册同册互踩（固定 .tmp 名会互相截断）
            _atomic_save(out, dst)
        # 呈现信号（产出即开）走 _publish_volume 的 tender.volume 产物发布路径
        # （publish_file_artifact 成功分支 note kind=artifact）——工作台整本行已隐藏，
        # 不再 note kind=file 防打开隐藏行；发布失败=不自动开，chips 仍是手动出口
        bits = f"{vol}：合并 {merged} 节 → work/{rel}"
        if n_img:
            bits += f"（含图片 {n_img} 张）"
        if n_comment:
            bits += f"（含批注 {n_comment} 条待处理）"
        if n_cap:
            bits += f"（图注/表题编号 {n_cap} 处）"
        if missing:
            bits += f"；缺失 {len(missing)} 节未并入：{'、'.join(missing)}"
        reports.append(bits)
        reports.append(_publish_volume(task_id, dst, vol, merged, n_img, n_comment))
        if placeholders:
            shown = "；".join(placeholders[:8]) + ("…" if len(placeholders) > 8 else "")
            reports.append(
                f"⚠️ {vol}：正文含 {len(placeholders)} 处内联占位【待补/待澄清】"
                "（占位文字会进交付稿——删除该文字，改用 docx_comment_add 加批注记待办）："
                + shown
            )
        if unfilled:
            reports.append(
                f"{vol}：模板填充类未产出 {len(unfilled)} 节（按附件对待，不占整本位）："
                + "、".join(t for t, _d in unfilled)
            )
        if toc_lines:
            # 目录页对账（探测+提示，不是门禁）：目录页条目 vs 实收章节——列了
            # 整本没有的（未产出/树外）、漏了实有的，点名请修目录节文件后重合册。
            # 「目录」「封面」两侧豁免：目录页列不列自己/封面都合理。
            toc_pairs = [(t, _toc_normalize(t)) for t in toc_lines]
            toc_keys = [k for _t, k in toc_pairs if k]
            issued_pairs = [(t, _toc_normalize(t)) for t in issued_titles]
            extra = [
                t for t, k in toc_pairs
                if k and k not in ("目录", "封面") and not _toc_match(k, [ik for _it, ik in issued_pairs if ik])
            ]
            absent = [
                t for t, k in issued_pairs
                if k and k not in ("目录", "封面") and not _toc_match(k, toc_keys)
            ]
            if extra or absent:
                bits = f"⚠️ {vol}：目录页与正文对账不符"
                if extra:
                    bits += "——目录页列了但整本没有 " + str(len(extra)) + " 项：" + "、".join(extra[:6]) + ("…" if len(extra) > 6 else "")
                if absent:
                    bits += "；整本有但目录页未列 " + str(len(absent)) + " 项：" + "、".join(absent[:6]) + ("…" if len(absent) > 6 else "")
                bits += "（修订目录节文件后重新合册；按附件对待未产出的节在目录页标注「另附」或不列）"
                reports.append(bits)
        if self_numbered:
            shown = "、".join(self_numbered[:6]) + ("…" if len(self_numbered) > 6 else "")
            reports.append(
                f"⚠️ {vol}：{len(self_numbered)} 个目录节点名自带编号（{shown}）——"
                "合册已再按树序自动编号，会出现「第一章 一、xxx」式双重编号；"
                "请在目录产物中把节点名改为不带编号的纯标题后重新合册"
            )
    orphans: list[str] = []
    body_root = wroot / "body"
    if body_root.is_dir():
        for p in sorted(body_root.rglob("*.docx")):
            rel_path = p.relative_to(wroot).as_posix()
            if not p.name.startswith(body_contract.VOLUME_PREFIX) and rel_path not in consumed:
                orphans.append(rel_path)
    lines = ["[已合册]", *reports]
    if orphans:
        shown = "、".join(orphans[:8]) + ("…" if len(orphans) > 8 else "")
        lines.append(f"⚠️ {len(orphans)} 个节文件未并入（目录无对应节点，先看 check_pipeline_state 的 [body] 段对账）：{shown}")
    lines.append("整本是派生产物：改内容请回节文件层改，再重新调用本工具覆盖合册。")
    return "\n".join(lines)
